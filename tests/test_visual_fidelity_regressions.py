from __future__ import annotations

import os
from pathlib import Path
import shutil
import subprocess
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Pt
from pptx.oxml.xmlchemy import OxmlElement

from pptx_typst.typst_to_pptx import convert as typst_to_pptx


ROOT = Path(__file__).resolve().parents[1]


def convert_forward(source: Path, project: Path) -> None:
    env = os.environ.copy()
    env["PYTHONPATH"] = str(ROOT / "src") + os.pathsep + env.get("PYTHONPATH", "")
    subprocess.run(
        [sys.executable, "-m", "pptx_typst.pptx_to_typst", str(source), "--output", str(project)],
        cwd=ROOT, env=env, check=True, capture_output=True, text=True, timeout=120,
    )


def test_table_rows_retain_source_frame_height_and_proportions(tmp_path: Path) -> None:
    source = tmp_path / "table.pptx"
    project = tmp_path / "typst-project"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(4, 1, Pt(20), Pt(30), Pt(300), Pt(80)).table
    for row in table.rows:
        row.height = Pt(10)
    slide.shapes[-1].height = Pt(80)
    deck.save(source)

    convert_forward(source, project)
    typst_source = (project / "main.typ").read_text(encoding="utf-8")
    assert "rows:(20.000pt, 20.000pt, 20.000pt, 20.000pt)" in typst_source


def test_theme_fill_reference_is_resolved_to_native_vector_color(tmp_path: Path) -> None:
    source = tmp_path / "theme-fill.pptx"
    project = tmp_path / "typst-project"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Pt(100), Pt(100), Pt(200), Pt(200))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(1, 2, 3)
    shape._element.spPr.remove(shape._element.spPr.xpath("./a:solidFill")[0])
    fill_ref = shape._element.xpath("./p:style/a:fillRef")[0]
    fill_ref.set("idx", "1")
    fill_ref[0].set("val", "accent2")
    deck.save(source)

    convert_forward(source, project)
    vector_assets = list((project / "assets" / "media").glob("shape_*.svg"))
    assert any('#C0504D' in path.read_text(encoding="utf-8") for path in vector_assets)


def test_wingdings_section_sign_bullet_maps_to_square(tmp_path: Path) -> None:
    source = tmp_path / "wingdings-bullet.pptx"
    project = tmp_path / "typst-project"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Pt(40), Pt(40), Pt(400), Pt(100))
    paragraph = text_box.text_frame.paragraphs[0]
    paragraph.add_run().text = "List item"
    ppr = paragraph._p.get_or_add_pPr()
    bullet = OxmlElement("a:buChar")
    bullet.set("char", chr(0x00A7))
    font = OxmlElement("a:buFont")
    font.set("typeface", "Wingdings")
    ppr.append(bullet)
    ppr.append(font)
    deck.save(source)

    convert_forward(source, project)
    typst_source = (project / "main.typ").read_text(encoding="utf-8")
    assert chr(0x25AA) in typst_source
    assert chr(0x00A7) not in typst_source


def test_placeholder_run_formatting_roundtrips(tmp_path: Path) -> None:
    source = tmp_path / "rich-placeholder.pptx"
    project = tmp_path / "typst-project"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    paragraph = slide.shapes.title.text_frame.paragraphs[0]
    bold = paragraph.add_run()
    bold.text = "Bold title "
    bold.font.bold = True
    italic = paragraph.add_run()
    italic.text = "italic ending"
    italic.font.italic = True
    line_break = OxmlElement("a:br")
    paragraph._p.insert(list(paragraph._p).index(bold._r) + 1, line_break)
    deck.save(source)

    convert_forward(source, project)
    typst_source = (project / "main.typ").read_text(encoding="utf-8")
    assert "PptxPlaceholderRun" in typst_source
    assert 'weight:"bold"' in typst_source
    assert "italic:true" in typst_source
    assert 'PptxPlaceholderRun(text:"\\n", font:"Calibri", size:44.000pt' in typst_source

    typst = shutil.which("typst")
    if typst:
        subprocess.run(
            [typst, "compile", "main.typ", "output.pdf"], cwd=project,
            check=True, capture_output=True, text=True, timeout=120,
        )
    output = tmp_path / "rich-roundtrip.pptx"
    typst_to_pptx(project / "main.typ", output)
    converted = Presentation(output)
    title = next(shape for shape in converted.slides[0].shapes if shape.is_placeholder and shape.text)
    runs = title.text_frame.paragraphs[0].runs
    assert [run.text for run in runs] == ["Bold title ", "italic ending"]
    assert runs[0].font.bold is True
    assert runs[1].font.italic is True
    assert len(title.text_frame.paragraphs[0]._p.xpath('./a:br')) == 1


def test_connector_uses_editable_typst_line_and_roundtrips(tmp_path: Path) -> None:
    source = tmp_path / "connector.pptx"
    project = tmp_path / "typst-project"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Pt(80), Pt(90), Pt(420), Pt(260)
    )
    connector.line.color.rgb = RGBColor(210, 25, 35)
    connector.line.width = Pt(2.25)
    deck.save(source)

    env = os.environ.copy()
    env["PYTHONPATH"] = str(ROOT / "src") + os.pathsep + env.get("PYTHONPATH", "")
    subprocess.run(
        [sys.executable, "-m", "pptx_typst.pptx_to_typst", str(source), "--output", str(project)],
        cwd=ROOT, env=env, check=True, capture_output=True, text=True, timeout=120,
    )
    typst_source = (project / "main.typ").read_text(encoding="utf-8")
    assert "PptxLine(start:" in typst_source
    assert "#D21923" in typst_source

    typst = shutil.which("typst")
    if typst:
        subprocess.run(
            [typst, "compile", "main.typ", "output.pdf"], cwd=project,
            check=True, capture_output=True, text=True, timeout=120,
        )

    output = tmp_path / "roundtrip.pptx"
    typst_to_pptx(project / "main.typ", output)
    converted = Presentation(output)
    lines = [shape for shape in converted.slides[0].shapes
             if shape.shape_type == MSO_SHAPE_TYPE.LINE]
    assert len(lines) == 1
    assert lines[0].line.color.rgb == RGBColor(210, 25, 35)
