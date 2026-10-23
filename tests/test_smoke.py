from __future__ import annotations

import hashlib
import json
import os
import posixpath
from pathlib import Path
import shutil
import subprocess
import sys
from zipfile import ZipFile

from lxml import etree

from PIL import Image
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Pt
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]


def write_pixel(path: Path) -> None:
    Image.new("RGBA", (8, 8), (40, 120, 200, 255)).save(path)


def make_source_pptx(path: Path, image_path: Path) -> None:
    deck = Presentation()
    deck.slide_width = Pt(960)
    deck.slide_height = Pt(540)
    for index in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Pt(40), Pt(40), Pt(420), Pt(60))
        box.text_frame.text = f"PPTX smoke slide {index + 1}"
        slide.shapes.add_picture(str(image_path), Pt(60), Pt(130), width=Pt(48), height=Pt(32))
    deck.save(path)


def run_forward(source: Path, output: Path) -> None:
    env = os.environ.copy()
    source_path = str(ROOT / "src")
    env["PYTHONPATH"] = source_path + (os.pathsep + env["PYTHONPATH"] if env.get("PYTHONPATH") else "")
    subprocess.run(
        [sys.executable, "-m", "pptx_typst.pptx_to_typst", str(source), "--output", str(output), "--text-layout", "v3"],
        cwd=ROOT,
        env=env,
        check=True,
        capture_output=True,
        text=True,
        timeout=120,
    )


def test_theme_background_references_are_resolved_for_master_and_layout(tmp_path: Path) -> None:
    source = tmp_path / "theme-backgrounds.pptx"
    project = tmp_path / "typst-project"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])

    def set_theme_background(owner, scheme: str) -> None:
        background = OxmlElement("p:bg")
        reference = OxmlElement("p:bgRef")
        reference.set("idx", "1001")
        color = OxmlElement("a:schemeClr")
        color.set("val", scheme)
        reference.append(color)
        background.append(reference)
        owner._element.cSld.insert(0, background)

    set_theme_background(deck.slide_masters[0], "bg2")
    set_theme_background(slide.slide_layout, "bg1")
    deck.save(source)

    with ZipFile(source) as package:
        master = etree.fromstring(package.read("ppt/slideMasters/slideMaster1.xml"))
        theme_name = master.xpath("string(./p:clrMap/@bg2)", namespaces={
            "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
        })
        theme = etree.fromstring(package.read("ppt/theme/theme1.xml"))
        color_node = theme.xpath(
            f"//*[local-name()='{theme_name}']/*[1]"
        )[0]
        expected_master = (color_node.get("val") if etree.QName(color_node).localname == "srgbClr"
                           else color_node.get("lastClr")).upper()
        light_name = master.xpath("string(./p:clrMap/@bg1)", namespaces={
            "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
        })
        light_node = theme.xpath(f"//*[local-name()='{light_name}']/*[1]")[0]
        expected_layout = (light_node.get("val") if etree.QName(light_node).localname == "srgbClr"
                           else light_node.get("lastClr")).upper()

    run_forward(source, project)
    generated = (project / "main.typ").read_text(encoding="utf-8")
    assert f'PptxMaster(fill:rgb("#{expected_master}")' in generated
    assert f'PptxLayout(master:Master01,fill:rgb("#{expected_layout}")' in generated
    assert "fill:none," in generated

    typst = shutil.which("typst")
    if typst:
        subprocess.run(
            [typst, "compile", "main.typ", "backgrounds.pdf"], cwd=project,
            check=True, capture_output=True, text=True, timeout=120,
        )
        assert (project / "backgrounds.pdf").is_file()


def test_slide_color_map_override_does_not_change_shared_master_fill(tmp_path: Path) -> None:
    source = tmp_path / "slide-color-map.pptx"
    project = tmp_path / "typst-project"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    master = deck.slide_masters[0]
    background = OxmlElement("p:bg")
    reference = OxmlElement("p:bgRef")
    reference.set("idx", "1001")
    color = OxmlElement("a:schemeClr")
    color.set("val", "bg2")
    reference.append(color)
    background.append(reference)
    master._element.cSld.insert(0, background)

    master_mapping = master._element.xpath("./p:clrMap")[0].attrib
    slide_mapping = dict(master_mapping)
    slide_mapping["bg2"] = "lt1" if master_mapping.get("bg2") != "lt1" else "dk1"
    override = OxmlElement("p:clrMapOvr")
    override_mapping = OxmlElement("a:overrideClrMapping")
    for key, value in slide_mapping.items():
        override_mapping.set(key, value)
    override.append(override_mapping)
    slide._element.append(override)
    deck.save(source)

    with ZipFile(source) as package:
        theme = etree.fromstring(package.read("ppt/theme/theme1.xml"))
        colors = {}
        for scheme in theme.xpath("//*[local-name()='clrScheme']/*"):
            node = scheme[0]
            colors[etree.QName(scheme).localname] = (
                node.get("val") if etree.QName(node).localname == "srgbClr" else node.get("lastClr")
            ).upper()
    expected_master = colors[master_mapping["bg2"]]
    expected_slide = colors[slide_mapping["bg2"]]

    run_forward(source, project)
    generated = (project / "main.typ").read_text(encoding="utf-8")
    assert f'PptxMaster(fill:rgb("#{expected_master}")' in generated
    assert f'fill:rgb("#{expected_slide}")' in generated
    assert expected_master != expected_slide


def test_theme_image_background_is_shared_and_roundtrips(tmp_path: Path) -> None:
    from pptx_typst.typst_to_pptx import convert

    source = tmp_path / "theme-image-source.pptx"
    styled = tmp_path / "theme-image-styled.pptx"
    project = tmp_path / "typst-project"
    image_path = tmp_path / "theme-background.png"
    write_pixel(image_path)
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_picture(
        str(image_path), Pt(8), Pt(8), width=Pt(24), height=Pt(24)
    )
    deck.save(source)

    a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    pkg_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    theme_rels_path = "ppt/theme/_rels/theme1.xml.rels"
    with ZipFile(source) as original, ZipFile(styled, "w") as output:
        original_names = set(original.namelist())
        rel_root = etree.fromstring(original.read(theme_rels_path)) if theme_rels_path in original_names else etree.Element(f"{{{pkg_ns}}}Relationships", nsmap={None: pkg_ns})
        rel = etree.SubElement(rel_root, f"{{{pkg_ns}}}Relationship")
        rel.set("Id", "rIdThemeBackground")
        rel.set("Type", r_ns + "/image")
        rel.set("Target", "../media/image1.png")
        rel_data = etree.tostring(rel_root, xml_declaration=True, encoding="UTF-8", standalone=True)
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename == "ppt/theme/theme1.xml":
                root = etree.fromstring(data)
                styles = root.xpath("//*[local-name()='bgFillStyleLst']")[0]
                image_fill = etree.Element(f"{{{a_ns}}}blipFill")
                blip = etree.SubElement(image_fill, f"{{{a_ns}}}blip")
                blip.set(f"{{{r_ns}}}embed", "rIdThemeBackground")
                stretch = etree.SubElement(image_fill, f"{{{a_ns}}}stretch")
                etree.SubElement(stretch, f"{{{a_ns}}}fillRect")
                styles.replace(styles[2], image_fill)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif info.filename == theme_rels_path:
                data = rel_data
            elif info.filename == "ppt/slideMasters/slideMaster1.xml":
                root = etree.fromstring(data)
                c_sld = root.find(f"{{{p_ns}}}cSld")
                background = etree.Element(f"{{{p_ns}}}bg")
                reference = etree.SubElement(background, f"{{{p_ns}}}bgRef")
                reference.set("idx", "1003")
                scheme = etree.SubElement(reference, f"{{{a_ns}}}schemeClr")
                scheme.set("val", "bg2")
                c_sld.insert(0, background)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            output.writestr(info, data)
        if theme_rels_path not in original_names:
            output.writestr(theme_rels_path, rel_data)

    run_forward(styled, project)
    generated = (project / "main.typ").read_text(encoding="utf-8")
    assert 'PptxMaster(fill:none,background:image("assets/media/' in generated
    manifest = json.loads((project / "asset_manifest.json").read_text(encoding="utf-8"))
    assert any(media_placement.get("owner") == "master" and media_placement.get("background")
               for media in manifest["media"].values()
               for media_placement in media["placements"])

    typst = shutil.which("typst")
    if typst:
        subprocess.run(
            [typst, "compile", "main.typ", "background.pdf"], cwd=project,
            check=True, capture_output=True, text=True, timeout=120,
        )
        assert (project / "background.pdf").is_file()

    output = tmp_path / "background-roundtrip.pptx"
    convert(project / "main.typ", output)
    with ZipFile(output) as package:
        master = etree.fromstring(package.read("ppt/slideMasters/slideMaster1.xml"))
        rels = etree.fromstring(package.read("ppt/slideMasters/_rels/slideMaster1.xml.rels"))
        rid = master.xpath("string(./p:cSld/p:bg/p:bgPr/a:blipFill/a:blip/@r:embed)", namespaces={
            "p": p_ns, "a": a_ns, "r": r_ns,
        })
        rel = next(item for item in rels if item.get("Id") == rid)
        target = Path("ppt/slideMasters") / rel.get("Target")
        target = Path(posixpath.normpath(target.as_posix()))
        assert rid and rel.get("Type").endswith("/image")
        assert package.read(target.as_posix()) == image_path.read_bytes()


def test_slide_text_placeholders_render_over_later_artwork(tmp_path: Path) -> None:
    source = tmp_path / "placeholder-over-artwork.pptx"
    project = tmp_path / "typst-project"
    image = tmp_path / "overlay.png"
    write_pixel(image)

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = "Visible placeholder title"
    slide.shapes.add_picture(str(image), Pt(0), Pt(0), width=Pt(720), height=Pt(405))
    deck.save(source)

    run_forward(source, project)
    generated = (project / "main.typ").read_text(encoding="utf-8")
    assert "PptxPlaceholder(" in generated
    assert "Visible placeholder title" in generated
    assert "// Editable text: Title 1" not in generated
    typst = shutil.which("typst")
    if typst:
        subprocess.run(
            [typst, "compile", "main.typ", "placeholder.pdf"], cwd=project,
            check=True, capture_output=True, text=True, timeout=120,
        )
        assert (project / "placeholder.pdf").is_file()


def test_placeholder_layout_inherits_master_placeholder_font(tmp_path: Path) -> None:
    source = tmp_path / "placeholder-theme-font.pptx"
    project = tmp_path / "typst-project"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Theme font inheritance"

    deck.save(source)

    run_forward(source, project)
    generated = (project / "main.typ").read_text(encoding="utf-8")
    layout_line = next(line for line in generated.splitlines()
                       if "PptxLayout(" in line and "PptxPlaceholderSlot(id:\"title-0\"" in line)
    assert 'font:"Calibri"' in layout_line
    slide_instance = next(line for line in generated.splitlines() if "PptxPlaceholder(id:\"title-0\"" in line)
    assert 'font:"Arial"' not in slide_instance

    typst = shutil.which("typst")
    if typst:
        subprocess.run(
            [typst, "compile", "main.typ", "placeholder-font.pdf"], cwd=project,
            check=True, capture_output=True, text=True, timeout=120,
        )
        assert (project / "placeholder-font.pdf").is_file()


def test_native_title_body_placeholder_relationships_roundtrip(tmp_path: Path) -> None:
    from pptx_typst.typst_to_pptx import convert

    project = tmp_path / "placeholder-project"
    project.mkdir()
    (project / "template.typ").write_text(
        (ROOT / "src/pptx_typst/template.typ").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    (project / "main.typ").write_text(
        '''#import "template.typ": *
#let Brand = PptxMaster(placeholders:(
  PptxPlaceholderSlot(id:"title-0",type:"title",idx:0,x:32pt,y:24pt,width:896pt,height:72pt,font:"Arial",size:36pt,fill:rgb("#123456"),weight:"bold"),
  PptxPlaceholderSlot(id:"body-1",type:"body",idx:1,x:40pt,y:120pt,width:880pt,height:360pt,font:"Arial",size:20pt,fill:rgb("#111111")),
))
#let Content = PptxLayout(master:Brand,placeholders:(
  PptxPlaceholderSlot(id:"title-0",type:"title",idx:0,x:40pt,y:50pt,width:800pt,height:60pt),
  PptxPlaceholderSlot(id:"body-1",type:"body",idx:1,x:40pt,y:130pt,width:800pt,height:320pt,bullet:"•"),
))
#let Cover = PptxLayout(master:Brand,placeholders:(
  PptxPlaceholderSlot(id:"title-0",type:"title",idx:0,x:80pt,y:100pt,width:760pt,height:100pt),
  PptxPlaceholderSlot(id:"body-1",type:"body",idx:1,x:80pt,y:240pt,width:760pt,height:220pt),
))
#PptxSlide(width:960pt,height:540pt,layout:Content,placeholders:(
  PptxPlaceholder(id:"title-0",type:"title",idx:0,text:"Slide One",size:48pt),
  PptxPlaceholder(id:"body-1",type:"body",idx:1,text:"First body\nSecond body"),
))
#pagebreak()
#PptxSlide(width:960pt,height:540pt,layout:Content,placeholders:(
  PptxPlaceholder(id:"title-0",type:"title",idx:0,text:"Slide Two"),
  PptxPlaceholder(id:"body-1",type:"body",idx:1,text:"Independent content"),
))
#pagebreak()
#PptxSlide(width:960pt,height:540pt,layout:Cover,placeholders:(
  PptxPlaceholder(id:"title-0",type:"title",idx:0,text:"Slide Three"),
  PptxPlaceholder(id:"body-1",type:"body",idx:1,text:"Cover content"),
))
#pagebreak()
#PptxSlide(width:960pt,height:540pt,layout:Content,placeholders:(
  PptxPlaceholder(id:"title-0",type:"title",idx:0,text:"Slide Four"),
  PptxPlaceholder(id:"body-1",type:"body",idx:1,text:"Fourth content"),
))
''',
        encoding="utf-8",
    )

    typst = shutil.which("typst")
    if typst:
        subprocess.run([typst, "compile", "main.typ", "before.pdf"], cwd=project,
                       check=True, capture_output=True, text=True, timeout=120)
    output = tmp_path / "native-placeholders.pptx"
    report = convert(project / "main.typ", output)
    assert report["slides"] == 4
    assert len(report["placeholders"]) == 8

    p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ns = {"p": p_ns, "a": a_ns, "r": r_ns}
    with ZipFile(output) as package:
        master = etree.fromstring(package.read("ppt/slideMasters/slideMaster1.xml"))
        master_slots = master.xpath(".//p:sp[p:nvSpPr/p:nvPr/p:ph]", namespaces=ns)
        assert {(x.xpath("string(p:nvSpPr/p:nvPr/p:ph/@type)", namespaces=ns),
                 x.xpath("string(p:nvSpPr/p:nvPr/p:ph/@idx)", namespaces=ns)) for x in master_slots} == {
            ("title", "0"), ("body", "1")
        }
        assert master.xpath('string(p:txStyles/p:titleStyle/a:lvl1pPr/a:defRPr/a:solidFill/a:srgbClr/@val)', namespaces=ns) == "123456"
        assert master.xpath('string(p:txStyles/p:titleStyle/a:lvl1pPr/a:defRPr/@sz)', namespaces=ns) == "3600"

        slide_layout_targets = []
        for number in range(1, 5):
            slide = etree.fromstring(package.read(f"ppt/slides/slide{number}.xml"))
            slots = slide.xpath(".//p:sp[p:nvSpPr/p:nvPr/p:ph]", namespaces=ns)
            assert len(slots) == 2
            assert not slide.xpath(".//p:sp[not(p:nvSpPr/p:nvPr/p:ph)]//a:t", namespaces=ns)
            assert all(not x.xpath("./p:spPr/a:xfrm", namespaces=ns) for x in slots)
            assert {x.xpath("string(p:nvSpPr/p:nvPr/p:ph/@type)", namespaces=ns) for x in slots} == {"title", "body"}
            rels = etree.fromstring(package.read(f"ppt/slides/_rels/slide{number}.xml.rels"))
            layout = next(rel for rel in rels if rel.get("Type", "").endswith("/slideLayout"))
            slide_layout_targets.append(layout.get("Target"))
        assert slide_layout_targets[0] == slide_layout_targets[1] == slide_layout_targets[3]
        assert slide_layout_targets[2] != slide_layout_targets[0]
        content_layout = etree.fromstring(package.read(
            "ppt/slideLayouts/" + slide_layout_targets[0].rsplit("/", 1)[-1]
        ))
        inherited_body_bullet = content_layout.xpath(
            './/p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]//a:lvl1pPr/a:buChar/@char',
            namespaces=ns,
        )
        assert inherited_body_bullet == ["•"]

        slide1 = etree.fromstring(package.read("ppt/slides/slide1.xml"))
        title = slide1.xpath('.//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="title"]]', namespaces=ns)[0]
        assert title.xpath(".//a:rPr/@sz", namespaces=ns) == ["4800"]
        slide2 = etree.fromstring(package.read("ppt/slides/slide2.xml"))
        title2 = slide2.xpath('.//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="title"]]', namespaces=ns)[0]
        assert not title2.xpath(".//a:rPr/@sz", namespaces=ns)

    reopened = Presentation(output)
    assert len(reopened.slides) == 4
    assert reopened.slides[0].shapes.title.text == "Slide One"
    assert reopened.slides[1].shapes.title.text == "Slide Two"

    imported = tmp_path / "roundtrip-typst"
    run_forward(output, imported)
    generated = (imported / "main.typ").read_text(encoding="utf-8")
    assert generated.count("PptxPlaceholder(") == 8
    assert 'text:"Slide One"' in generated
    assert 'bullet:"•"' in generated
    if typst:
        subprocess.run([typst, "compile", "main.typ", "after.pdf"], cwd=imported,
                       check=True, capture_output=True, text=True, timeout=120)


def test_init_uses_current_master_layout_macro(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from pptx_typst.cli import main as cli_main
    from pptx_typst.typst_to_pptx import convert

    project = tmp_path / "starter"
    monkeypatch.setattr(sys, "argv", ["pptx-typst", "init", str(project)])
    cli_main()

    packaged_macro = (ROOT / "src/pptx_typst/templates/pptx-typst.typ").read_text(encoding="utf-8")
    copied_macro = (project / "pptx-typst.typ").read_text(encoding="utf-8")
    starter = (project / "main.typ").read_text(encoding="utf-8")
    assert copied_macro == packaged_macro
    packaged_assets = ROOT / "src/pptx_typst/templates/assets"
    assert {path.name for path in (project / "assets").iterdir()} == {path.name for path in packaged_assets.iterdir()}
    assert "#let FieldNotes = PptxMaster(" in starter
    assert "#let CoverLayout = PptxLayout(" in starter
    assert "#let StoryLayout = PptxLayout(" in starter
    assert "#let InsightLayout = PptxLayout(" in starter

    typst = shutil.which("typst")
    if typst:
        subprocess.run(
            [typst, "compile", "main.typ", "starter.pdf"], cwd=project,
            check=True, capture_output=True, text=True, timeout=120,
        )

    output = project / "starter.pptx"
    report = convert(project / "main.typ", output)
    deck = Presentation(output)
    assert report["hierarchy_written"] is True
    assert len(deck.slides) == 3
    assert len(deck.slide_masters) == 1
    assert [slide.slide_layout.name for slide in deck.slides] == ["CoverLayout", "StoryLayout", "InsightLayout"]
    master_part = deck.slide_masters[0].part.partname
    assert all(slide.slide_layout.slide_master.part.partname == master_part for slide in deck.slides)
    svg_warnings = [item for item in report["unsupported"] if item["kind"] == "svg_renderer_unavailable"]
    master_artwork_count = sum(not shape.is_placeholder for shape in deck.slide_masters[0].shapes)
    if svg_warnings:
        assert master_artwork_count == 2
    else:
        assert master_artwork_count == 3
    slide_text = text_in(deck)
    assert any("Systems for a more resilient future" in text for text in slide_text)
    assert any("Care should flow around people" in text for text in slide_text)
    assert any("Small shifts, measurable change" in text for text in slide_text)


def text_in(deck: Presentation) -> list[str]:
    return [shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame]


def picture_count(deck: Presentation) -> int:
    return sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in deck.slides for shape in slide.shapes)


def test_pptx_converts_to_compilable_typst(tmp_path: Path) -> None:
    typst = shutil.which("typst")
    if not typst:
        pytest.skip("Typst CLI is required for the compile smoke test")
    image_path = tmp_path / "pixel.png"
    source = tmp_path / "source.pptx"
    project = tmp_path / "typst-project"
    write_pixel(image_path)
    make_source_pptx(source, image_path)

    run_forward(source, project)
    assert (project / "main.typ").is_file()
    subprocess.run(
        [typst, "compile", "main.typ", "smoke.pdf"],
        cwd=project,
        check=True,
        capture_output=True,
        text=True,
        timeout=120,
    )
    assert (project / "smoke.pdf").is_file()


def test_tiff_media_gets_a_typst_compatible_derivative(tmp_path: Path) -> None:
    typst = shutil.which("typst")
    if not typst:
        pytest.skip("Typst CLI is required for the TIFF media regression test")
    image_path = tmp_path / "source.tiff"
    Image.new("RGBA", (12, 8), (40, 120, 200, 128)).save(image_path, format="TIFF")
    source = tmp_path / "tiff.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(image_path), Pt(20), Pt(20), width=Pt(60), height=Pt(40))
    deck.save(source)

    project = tmp_path / "typst-project"
    run_forward(source, project)
    manifest = json.loads((project / "asset_manifest.json").read_text(encoding="utf-8"))
    media = next(iter(manifest["media"].values()))
    assert media["path"].endswith(".tiff")
    assert media["render_path"].endswith("_typst.png")
    assert (project / media["path"]).is_file()
    with Image.open(project / media["render_path"]) as converted:
        assert converted.format == "PNG"
        assert converted.mode == "RGBA"
        assert converted.getpixel((0, 0))[3] == 128

    subprocess.run(
        [typst, "compile", "main.typ", "tiff.pdf"], cwd=project,
        check=True, capture_output=True, text=True, timeout=120,
    )
    assert (project / "tiff.pdf").is_file()


def test_default_text_layout_does_not_insert_a_full_line_between_paragraphs(tmp_path: Path) -> None:
    source = tmp_path / "paragraphs.pptx"
    project = tmp_path / "typst-project"
    image_path = tmp_path / "pixel.png"
    write_pixel(image_path)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    frame = slide.shapes.add_textbox(Pt(40), Pt(40), Pt(500), Pt(300)).text_frame
    frame.clear()
    frame.paragraphs[0].text = "First paragraph"
    frame.add_paragraph().text = "Second paragraph"
    for paragraph in frame.paragraphs:
        paragraph.runs[0].font.size = Pt(20)
    slide.shapes.add_picture(str(image_path), Pt(60), Pt(360), width=Pt(48), height=Pt(32))
    deck.save(source)

    env = os.environ.copy()
    env["PYTHONPATH"] = str(ROOT / "src") + (os.pathsep + env["PYTHONPATH"] if env.get("PYTHONPATH") else "")
    subprocess.run(
        [sys.executable, "-m", "pptx_typst.pptx_to_typst", str(source), "--output", str(project)],
        cwd=ROOT, env=env, check=True, capture_output=True, text=True, timeout=120,
    )

    typst_source = (project / "main.typ").read_text(encoding="utf-8")
    assert "First paragraph" in typst_source and "Second paragraph" in typst_source
    assert "#v(26.00pt)" not in typst_source
    assert "#v(5.20pt)" not in typst_source
    assert "PptxTextFrame(" in typst_source
    assert "PptxParagraph(" in typst_source
    typst = shutil.which("typst")
    if typst:
        subprocess.run(
            [typst, "compile", "main.typ", "paragraphs.pdf"], cwd=project,
            check=True, capture_output=True, text=True, timeout=120,
        )
        assert (project / "paragraphs.pdf").is_file()

    legacy_project = tmp_path / "legacy-project"
    subprocess.run(
        [sys.executable, "-m", "pptx_typst.pptx_to_typst", str(source), "--output", str(legacy_project),
         "--text-layout", "anchors"],
        cwd=ROOT, env=env, check=True, capture_output=True, text=True, timeout=120,
    )
    assert "#v(26.00pt)" not in (legacy_project / "main.typ").read_text(encoding="utf-8")
    default_manifest = json.loads((project / "asset_manifest.json").read_text(encoding="utf-8"))
    legacy_manifest = json.loads((legacy_project / "asset_manifest.json").read_text(encoding="utf-8"))
    non_text = lambda manifest: [
        placement for slide in manifest["slides"] for placement in slide["placements"]
        if placement.get("type") != "text"
    ]
    assert non_text(default_manifest) == non_text(legacy_manifest)


def test_default_text_layout_compiles_multiline_bullets_and_vertical_anchor(tmp_path: Path) -> None:
    source = tmp_path / "bullets.pptx"
    project = tmp_path / "typst-project"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Pt(40), Pt(40), Pt(360), Pt(200))
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    first = shape.text_frame.paragraphs[0]
    first.text = "A multiline bullet\nwith a continued line"
    first._p.get_or_add_pPr().append(OxmlElement("a:buChar"))
    first._p.pPr[-1].set("char", "•")
    second = shape.text_frame.add_paragraph()
    second.text = "A second bullet"
    second._p.get_or_add_pPr().append(OxmlElement("a:buChar"))
    second._p.pPr[-1].set("char", "•")
    deck.save(source)

    env = os.environ.copy()
    env["PYTHONPATH"] = str(ROOT / "src") + (os.pathsep + env["PYTHONPATH"] if env.get("PYTHONPATH") else "")
    subprocess.run(
        [sys.executable, "-m", "pptx_typst.pptx_to_typst", str(source), "--output", str(project)],
        cwd=ROOT, env=env, check=True, capture_output=True, text=True, timeout=120,
    )
    generated = (project / "main.typ").read_text(encoding="utf-8")
    assert "A multiline bullet" in generated and "with a continued line" in generated
    assert "A second bullet" in generated
    assert 'bullet:"•"' in generated
    assert "PptxTextFrame(" in generated
    assert 'anchor:"center"' in generated
    typst = shutil.which("typst")
    if typst:
        subprocess.run(
            [typst, "compile", "main.typ", "bullets.pdf"], cwd=project,
            check=True, capture_output=True, text=True, timeout=120,
        )
        assert (project / "bullets.pdf").is_file()


def test_supported_typst_converts_to_pptx(tmp_path: Path) -> None:
    from pptx_typst.typst_to_pptx import convert

    image_path = tmp_path / "pixel.png"
    write_pixel(image_path)
    main = tmp_path / "main.typ"
    main.write_text(
        '''#PptxSlide(width:320pt,height:180pt,fill:rgb("#FFFFFF"),children:[
  #PptxPlace(x:20pt,y:20pt,child:PptxBox(width:250pt,height:40pt)[
    #PptxText(font:"Arial",size:18pt,fill:rgb("#111111"))[Supported reverse conversion]
  ])
  #PptxPlace(x:20pt,y:80pt,child:PptxImage("pixel.png",width:32pt,height:24pt))
  #PptxPlace(x:20pt,y:120pt,child:PptxText(font:"Arial",size:16pt,fill:rgb("#222222"))[Direct starter text call])
])
''',
        encoding="utf-8",
    )
    output = tmp_path / "reverse.pptx"

    report = convert(main, output)
    deck = Presentation(output)
    assert report["slides"] == 1
    assert deck.slide_width == Pt(320)
    assert deck.slide_height == Pt(180)
    assert any("Supported reverse conversion" in text for text in text_in(deck))
    assert any("Direct starter text call" in text for text in text_in(deck))
    assert picture_count(deck) == 1


def test_reverse_conversion_preserves_image_fit_and_rounded_panel_style(tmp_path: Path) -> None:
    from pptx_typst.typst_to_pptx import convert

    wide = tmp_path / "wide.png"
    Image.new("RGB", (200, 100), (80, 120, 160)).save(wide)
    main = tmp_path / "main.typ"
    main.write_text(
        '''#PptxSlide(width:320pt,height:180pt,fill:rgb("#FFFFFF"),children:[
  #PptxPlace(x:20pt,y:20pt,child:PptxImage("wide.png",width:40pt,height:40pt,fit:"cover"))
  #PptxPlace(x:80pt,y:20pt,child:PptxRect(width:100pt,height:60pt,fill:rgb("#FF000080"),stroke:1pt + rgb("#223344"),radius:8pt))
])
''',
        encoding="utf-8",
    )
    output = tmp_path / "fit-and-panel.pptx"

    convert(main, output)
    deck = Presentation(output)
    picture = next(shape for shape in deck.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    panel = next(shape for shape in deck.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE)
    assert picture.crop_left == pytest.approx(0.25, abs=0.01)
    assert picture.crop_right == pytest.approx(0.25, abs=0.01)
    assert panel.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
    assert panel.line.width == Pt(1)
    effect_ref = panel._element.find(qn("p:style")).find(qn("a:effectRef"))
    assert effect_ref.get("idx") == "0"
    assert panel._element.spPr.find(".//" + qn("a:outerShdw")) is None


def test_roundtrip_preserves_dimensions_text_and_shared_image(tmp_path: Path) -> None:
    from pptx_typst.typst_to_pptx import convert

    image_path = tmp_path / "pixel.png"
    source = tmp_path / "source.pptx"
    project = tmp_path / "typst-project"
    output = tmp_path / "roundtrip.pptx"
    write_pixel(image_path)
    make_source_pptx(source, image_path)

    run_forward(source, project)
    manifest = json.loads((project / "asset_manifest.json").read_text(encoding="utf-8"))
    assert manifest["unique_images"] == 1
    report = convert(project / "main.typ", output)
    result = Presentation(output)
    assert report["slides"] == 2
    assert result.slide_width == Pt(960)
    assert result.slide_height == Pt(540)
    assert "PPTX smoke slide 1" in text_in(result)
    assert "PPTX smoke slide 2" in text_in(result)
    assert picture_count(result) == 2
    image_hashes = {
        hashlib.sha256(shape.image.blob).hexdigest()
        for slide in result.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    }
    assert len(image_hashes) == 1


def test_master_layout_hierarchy_is_shared_and_roundtrips(tmp_path: Path) -> None:
    from pptx_typst.typst_to_pptx import convert

    write_pixel(tmp_path / "pixel.png")
    main = tmp_path / "main.typ"
    main.write_text(
        '''#import "pptx-typst.typ": PptxMaster, PptxLayout, PptxSlide, Place, PptxPlace, PptxRect, PptxImage
#let Brand = PptxMaster(fill:rgb("#112233"),children:[
  #PptxPlace(x:10pt,y:10pt,child:PptxRect(width:12pt,height:12pt,fill:rgb("#123456")))
  #PptxPlace(x:30pt,y:10pt,child:PptxImage("pixel.png",width:20pt,height:20pt))
])
#let Other = PptxMaster(children:[
  #Place(x:10pt,y:10pt,child:rect(width:14pt,height:14pt,fill:rgb("#654321")))
])
#let Alpha = PptxLayout(master:Brand,fill:rgb("#223344"),children:[
  #PptxPlace(x:10pt,y:40pt,child:PptxRect(width:30pt,height:8pt,fill:rgb("#CC0000")))
])
#let Beta = PptxLayout(master:Brand,children:[
  #Place(x:10pt,y:60pt,child:rect(width:30pt,height:8pt,fill:rgb("#00CC00")))
])
#let Gamma = PptxLayout(master:Other,children:[
  #Place(x:10pt,y:80pt,child:rect(width:30pt,height:8pt,fill:rgb("#0000CC")))
])
#PptxSlide(width:320pt,height:180pt,layout:Alpha,children:[#Place(x:80pt,y:20pt,child:text(size:18pt)[Slide one])])
#pagebreak()
#PptxSlide(width:320pt,height:180pt,layout:Alpha,children:[#Place(x:80pt,y:20pt,child:text(size:18pt)[Slide two])])
#pagebreak()
#PptxSlide(width:320pt,height:180pt,layout:Beta,children:[#Place(x:80pt,y:20pt,child:text(size:18pt)[Slide three])])
#pagebreak()
#PptxSlide(width:320pt,height:180pt,layout:Gamma,children:[#Place(x:80pt,y:20pt,child:text(size:18pt)[Slide four])])
''',
        encoding="utf-8",
    )
    # The starter module supports both composition and the bounded static parser.
    (tmp_path / "template.typ").write_text((ROOT / "src/pptx_typst/template.typ").read_text(encoding="utf-8"), encoding="utf-8")
    (tmp_path / "pptx-typst.typ").write_text((ROOT / "src/pptx_typst/templates/pptx-typst.typ").read_text(encoding="utf-8"), encoding="utf-8")
    typst = shutil.which("typst")
    if typst:
        subprocess.run([typst, "compile", "main.typ", "authored.pdf"], cwd=tmp_path, check=True, capture_output=True, text=True, timeout=120)
        pdftotext = shutil.which("pdftotext")
        if pdftotext:
            extracted = subprocess.run([pdftotext, "authored.pdf", "-"], cwd=tmp_path, check=True, capture_output=True, text=True, timeout=30).stdout
            assert "," not in extracted
    output = tmp_path / "hierarchy.pptx"
    report = convert(main, output)

    deck = Presentation(output)
    assert report["hierarchy_written"] is True
    assert len(deck.slides) == 4
    assert len(deck.slide_masters) == 2
    assert [slide.slide_layout.name for slide in deck.slides] == ["Alpha", "Alpha", "Beta", "Gamma"]
    assert [str(slide.slide_layout.slide_master.part.partname) for slide in deck.slides] == [
        str(deck.slide_masters[0].part.partname), str(deck.slide_masters[0].part.partname),
        str(deck.slide_masters[0].part.partname), str(deck.slide_masters[1].part.partname),
    ]
    assert [len(slide.shapes) for slide in deck.slides] == [1, 1, 1, 1]
    assert "Slide four" in text_in(deck)[-1]
    assert sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in deck.slide_masters[0].shapes) == 1
    assert sum(not shape.is_placeholder for shape in deck.slide_masters[0].shapes) == 2
    assert sum(not shape.is_placeholder for shape in deck.slide_masters[1].shapes) == 1
    assert [len(master.slide_layouts) for master in deck.slide_masters] == [2, 1]
    assert len({str(master.part.related_part(rel.rId).partname) for master in deck.slide_masters for rel in master.part.rels.values() if rel.reltype.endswith("/theme")}) == 2

    # Inspect package ownership and relationship targets directly. Slide parts
    # contain only slide-owned text; shared artwork lives in its master/layout.
    ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships", "rel": "http://schemas.openxmlformats.org/package/2006/relationships"}
    with ZipFile(output) as package:
        names = set(package.namelist())
        presentation = etree.fromstring(package.read("ppt/presentation.xml"))
        presentation_rels = etree.fromstring(package.read("ppt/_rels/presentation.xml.rels"))
        rel_map = {rel.get("Id"): rel.get("Target") for rel in presentation_rels}
        slide_ids = presentation.xpath("./p:sldIdLst/p:sldId", namespaces=ns)
        assert len(slide_ids) == 4
        slide_parts = []
        for slide_id in slide_ids:
            target = rel_map[slide_id.get("{%s}id" % ns["r"])]
            slide_part = "ppt/" + target
            slide_parts.append(slide_part)
            slide_xml = etree.fromstring(package.read(slide_part))
            assert len(slide_xml.xpath("./p:cSld/p:spTree/p:sp", namespaces=ns)) == 1
            rel_path = str(Path(slide_part).parent / "_rels" / (Path(slide_part).name + ".rels")).replace("\\", "/")
            slide_rels = etree.fromstring(package.read(rel_path))
            layout_rel = next(rel for rel in slide_rels if rel.get("Type", "").endswith("/slideLayout"))
            assert layout_rel.get("Target").startswith("../slideLayouts/")
        assert slide_parts[0] != slide_parts[1]
        master_parts = [name for name in names if name.startswith("ppt/slideMasters/slideMaster") and name.endswith(".xml")]
        assert len(master_parts) == 2
        layout_ids = []
        master1_rels = etree.fromstring(package.read("ppt/slideMasters/_rels/slideMaster1.xml.rels"))
        master1_targets = {rel.get("Target") for rel in master1_rels if rel.get("Type", "").endswith("/slideLayout")}
        assert len(master1_targets) == 2
        for master_part in master_parts:
            master_xml = etree.fromstring(package.read(master_part))
            layout_ids.extend(master_xml.xpath("./p:sldLayoutIdLst/p:sldLayoutId/@id", namespaces=ns))
            assert len(master_xml.xpath("./p:cSld/p:spTree/p:sp", namespaces=ns)) >= 1
            master_rels_path = str(Path(master_part).parent / "_rels" / (Path(master_part).name + ".rels")).replace("\\", "/")
            owner_rels = etree.fromstring(package.read(master_rels_path))
            for rel in owner_rels:
                if rel.get("Type", "").endswith("/image"):
                    target = posixpath.normpath(posixpath.join(posixpath.dirname(master_part), rel.get("Target")))
                    assert target in names
        assert len(layout_ids) == len(set(layout_ids))
        master1_xml = etree.fromstring(package.read("ppt/slideMasters/slideMaster1.xml"))
        alpha_layout = next(name for name in names if name.endswith("slideLayout12.xml"))
        alpha_xml = etree.fromstring(package.read(alpha_layout))
        assert master1_xml.xpath("string(./p:cSld/p:bg/p:bgPr/a:solidFill/a:srgbClr/@val)", namespaces={**ns, "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}) == "112233"
        assert alpha_xml.xpath("string(./p:cSld/p:bg/p:bgPr/a:solidFill/a:srgbClr/@val)", namespaces={**ns, "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}) == "223344"

    # Forward conversion must recover the same explicit declarations and use
    # every shared object only once in the generated Typst source.
    project = tmp_path / "forward-project"
    run_forward(output, project)
    manifest = json.loads((project / "asset_manifest.json").read_text(encoding="utf-8"))
    generated = (project / "main.typ").read_text(encoding="utf-8")
    assert len(manifest["masters"]) == 2
    assert len(manifest["layouts"]) == 3
    assert generated.count(" = PptxMaster(") == 2
    assert generated.count(" = PptxLayout(") == 3
    assert generated.count('"Slide one"') == 1
    assert "#PptxSlide(\n" in generated
    assert "#Place(\n    x:" in generated
    assert 'PptxMaster(fill:rgb("#112233")' in generated
    assert 'PptxLayout(master:Master01,fill:rgb("#223344")' in generated

    if typst:
        subprocess.run([typst, "compile", "main.typ", "hierarchy.pdf"], cwd=project, check=True, capture_output=True, text=True, timeout=120)
        pdftotext = shutil.which("pdftotext")
        if pdftotext:
            extracted = subprocess.run([pdftotext, "hierarchy.pdf", "-"], cwd=project, check=True, capture_output=True, text=True, timeout=30).stdout
            assert extracted.count("Slide one") == 1
            assert "," not in extracted


def test_table_fill_transparency_survives_pptx_to_typst(tmp_path: Path) -> None:
    source = tmp_path / "transparent-table.pptx"
    project = tmp_path / "typst-project"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(1, 1, Pt(20), Pt(20), Pt(160), Pt(40)).table
    cell = table.cell(0, 0)
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(51, 102, 153)
    srgb = cell._tc.get_or_add_tcPr().xpath('./a:solidFill/a:srgbClr')[0]
    alpha = OxmlElement('a:alpha')
    alpha.set('val', '40000')
    srgb.append(alpha)
    cell.text = "Transparent cell"
    deck.save(source)

    run_forward(source, project)
    typst = (project / "main.typ").read_text(encoding="utf-8")
    assert 'fill:rgb("#33669966")' in typst


def test_table_fill_transparency_survives_typst_to_pptx(tmp_path: Path) -> None:
    from pptx_typst.typst_to_pptx import convert

    main = tmp_path / "main.typ"
    main.write_text(
        '''#PptxSlide(width:320pt,height:180pt,fill:rgb("#FFFFFF"),children:[
  #PptxPlace(x:20pt,y:20pt,child:table(columns:(100pt),rows:(40pt),
    table.cell(fill:rgb("#33669966"),stroke:none)[Transparent cell]))
])
''',
        encoding="utf-8",
    )
    output = tmp_path / "table.pptx"

    convert(main, output)
    deck = Presentation(output)
    cell = next(shape.table.cell(0, 0) for shape in deck.slides[0].shapes if shape.has_table)
    alpha = cell._tc.get_or_add_tcPr().xpath('./a:solidFill/a:srgbClr/a:alpha/@val')
    assert alpha == ['40000']


def test_image_path_cannot_escape_project_to_prefix_sibling(tmp_path: Path) -> None:
    from pptx_typst.typst_to_pptx import convert

    project = tmp_path / "project"
    sibling = tmp_path / "project-evil"
    project.mkdir()
    sibling.mkdir()
    write_pixel(sibling / "outside.png")
    main = project / "main.typ"
    main.write_text(
        '''#PptxSlide(width:320pt,height:180pt,fill:rgb("#FFFFFF"),children:[
  #PptxPlace(x:10pt,y:10pt,child:PptxImage("../project-evil/outside.png",width:32pt,height:24pt))
])
''',
        encoding="utf-8",
    )

    report = convert(main, tmp_path / "blocked.pptx")
    assert report["pictures"] == 0
    assert any(item["kind"] == "missing_media" for item in report["unsupported"])


def test_forward_report_classifies_unsupported_chart(tmp_path: Path) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    source = tmp_path / "unsupported-chart.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("Values", [1, 2])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(40), Pt(40), Pt(240), Pt(160), data)
    deck.save(source)

    output = tmp_path / "converted"
    run_forward(source, output)
    report = json.loads((output / "conversion_report.json").read_text(encoding="utf-8"))
    warning = next(item for item in report["warnings"] if item["issue"] == "unsupported_object")

    assert warning["outcome"] == "unsupported"
    assert warning["source_type"] == "chart"
    assert warning["slide"] == 1
    assert warning["shape_id"] is not None
    assert warning["fallback"] == "none"
    assert "Unsupported chart" in warning["reason"]


def test_forward_uses_embedded_ole_fallback_picture(tmp_path: Path) -> None:
    from copy import deepcopy
    from lxml import etree

    source = tmp_path / "ole-fallback.pptx"
    image = tmp_path / "preview.png"
    write_pixel(image)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    picture = slide.shapes.add_picture(str(image), Pt(40), Pt(50), width=Pt(160), height=Pt(90))
    picture_element = deepcopy(picture._element)
    slide.shapes._spTree.remove(picture._element)

    frame = etree.Element(qn("p:graphicFrame"))
    nonvisual = etree.SubElement(frame, qn("p:nvGraphicFramePr"))
    properties = etree.SubElement(nonvisual, qn("p:cNvPr"))
    properties.set("id", "99")
    properties.set("name", "Embedded Chart Preview")
    etree.SubElement(nonvisual, qn("p:cNvGraphicFramePr"))
    etree.SubElement(nonvisual, qn("p:nvPr"))
    transform = etree.SubElement(frame, qn("p:xfrm"))
    offset = etree.SubElement(transform, qn("a:off"))
    offset.set("x", str(int(Pt(40))))
    offset.set("y", str(int(Pt(50))))
    extent = etree.SubElement(transform, qn("a:ext"))
    extent.set("cx", str(int(Pt(160))))
    extent.set("cy", str(int(Pt(90))))
    graphic = etree.SubElement(frame, qn("a:graphic"))
    data = etree.SubElement(graphic, qn("a:graphicData"))
    data.set("uri", "http://schemas.openxmlformats.org/presentationml/2006/ole")
    mc = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    alternate = etree.SubElement(data, f"{{{mc}}}AlternateContent")
    fallback = etree.SubElement(alternate, f"{{{mc}}}Fallback")
    ole = etree.SubElement(fallback, qn("p:oleObj"))
    ole.set("progId", "Package")
    ole.append(picture_element)
    slide.shapes._spTree.insert(2, frame)
    deck.save(source)

    output = tmp_path / "converted"
    run_forward(source, output)
    report = json.loads((output / "conversion_report.json").read_text(encoding="utf-8"))
    manifest = json.loads((output / "asset_manifest.json").read_text(encoding="utf-8"))
    warning = next(item for item in report["warnings"] if item.get("source_type") == "ole_object")
    placement = next(item for item in manifest["slides"][0]["placements"] if item.get("source_type") == "ole_object")

    assert warning["outcome"] == "approximated"
    assert warning["fallback"] == "embedded OLE fallback picture"
    assert "not editable in Typst" in warning["reason"]
    assert placement["type"] == "image"
    assert placement["render_source"].startswith("assets/media/")
    assert (output / placement["render_source"]).is_file()


def test_forward_reports_placeholder_and_later_picture_overlap(tmp_path: Path) -> None:
    source = tmp_path / "placeholder-layering.pptx"
    image = tmp_path / "cover.png"
    write_pixel(image)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    body = slide.placeholders[1]
    body.text = "Text behind the picture"
    slide.shapes.add_picture(str(image), body.left, body.top, width=body.width, height=body.height)
    deck.save(source)

    output = tmp_path / "converted"
    run_forward(source, output)
    report = json.loads((output / "conversion_report.json").read_text(encoding="utf-8"))
    warning = next(item for item in report["warnings"] if item["issue"] == "placeholder_layering_risk")

    assert warning["outcome"] == "approximated"
    assert warning["source_type"] == "placeholder:object"
    assert warning["slide"] == 1
    assert warning["shape_id"] is not None
    assert warning["overlapping_shape_id"] is not None
    assert "may appear over a picture" in warning["expected_loss"]


def test_layout_validator_reports_source_objects_without_geometry(tmp_path: Path) -> None:
    source = tmp_path / "unpositioned.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Pt(10), Pt(10), Pt(120), Pt(30))
    shape.text_frame.text = "Unpositioned object"
    shape._element.spPr.remove(shape._element.spPr.xfrm)
    shape_id = shape.shape_id
    deck.save(source)

    project = tmp_path / "project"
    project.mkdir()
    (project / "asset_manifest.json").write_text(json.dumps({
        "unique_images": 0,
        "counts": {"group": 0},
        "slides": [{
            "slide": 1,
            "groups": [],
            "placements": [{
                "source_shape_id": shape_id,
                "type": "unsupported",
                "z_index": 0,
                "group_path": [],
            }],
        }],
    }), encoding="utf-8")
    output = tmp_path / "validation.json"
    result = subprocess.run(
        [
            sys.executable,
            str(ROOT / "src" / "pptx_typst" / "validate_layout.py"),
            str(source),
            str(project),
            "--out",
            str(output),
        ],
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    validation = json.loads(output.read_text(encoding="utf-8"))
    assert validation["summary"]["source_objects_without_geometry"] == 1
    assert validation["slides"][0]["issues"] == [{
        "shape": shape_id,
        "issue": "source geometry unavailable; position and size cannot be compared",
    }]


def test_forward_cli_help_does_not_require_native_cairo(monkeypatch: pytest.MonkeyPatch) -> None:
    import runpy

    monkeypatch.setitem(sys.modules, "cairosvg", None)
    monkeypatch.setattr(sys, "argv", ["pptx-to-typst", "--help"])
    with pytest.raises(SystemExit) as result:
        runpy.run_module("pptx_typst.pptx_to_typst", run_name="__main__")

    assert result.value.code == 0


def test_reverse_svg_without_native_cairo_is_reported(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import builtins

    from pptx_typst.typst_to_pptx import convert

    project = tmp_path / "project"
    assets = project / "assets"
    assets.mkdir(parents=True)
    (assets / "art.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="20" height="10"><rect width="20" height="10"/></svg>',
        encoding="utf-8",
    )
    main = project / "main.typ"
    main.write_text(
        '''#PptxSlide(width:320pt,height:180pt,fill:rgb("#FFFFFF"),children:[
  #PptxPlace(x:20pt,y:20pt,child:PptxImage("assets/art.svg",width:40pt,height:20pt))
])
''',
        encoding="utf-8",
    )
    real_import = builtins.__import__

    def without_cairo(name, *args, **kwargs):
        if name == "cairosvg":
            raise OSError("native Cairo library is unavailable")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", without_cairo)
    report = convert(main, tmp_path / "output.pptx")

    assert any(item["kind"] == "svg_renderer_unavailable" for item in report["unsupported"])
