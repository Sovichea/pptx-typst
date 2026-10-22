from __future__ import annotations

import os
from pathlib import Path
import shutil
import subprocess
import sys
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt
import pytest

from pptx_typst.paragraph_layout import resolve_paragraph_style, typst_target_pitch, typst_spacing
from pptx_typst.typst_to_pptx import convert as convert_to_pptx, find_macro_records


ROOT = Path(__file__).resolve().parents[1]
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}


def _ppr(**attrs):
    node = OxmlElement("a:pPr")
    for name, value in attrs.items():
        node.set(name, str(value))
    return node


def _spacing(ppr, tag, kind, value):
    spacing = OxmlElement("a:" + tag)
    child = OxmlElement("a:spcPct" if kind == "percent" else "a:spcPts")
    child.set("val", str(value))
    spacing.append(child)
    ppr.append(spacing)


def test_paragraph_model_resolves_each_property_independently():
    direct = _ppr(algn="ctr")
    layout = _ppr(marL=254000, indent=-101600)
    _spacing(layout, "lnSpc", "percent", 120000)
    master = _ppr()
    _spacing(master, "spcBef", "points", 450)
    _spacing(master, "spcAft", "percent", 50000)

    style = resolve_paragraph_style(
        (("direct", direct), ("layout", layout), ("master_style", master)), level=2
    )
    assert style.alignment.value == "center" and style.alignment.explicit
    assert style.line_spacing.value.kind == "percent"
    assert style.line_spacing.value.value == 120
    assert style.line_spacing.source == "layout"
    assert style.space_before.value.kind == "points"
    assert style.space_before.value.value == 4.5
    assert style.space_after.value.kind == "percent"
    assert style.space_after.value.value == 50
    assert style.margin_left_pt.value == 20 and style.indent_pt.value == -8
    assert typst_target_pitch(style, 20) == pytest.approx(28.8)
    assert typst_spacing(style.space_before, 20) == 4.5
    # Unspecified values remain unspecified instead of receiving a source guess.
    assert resolve_paragraph_style((("direct", _ppr()),)).line_spacing is None


@pytest.mark.parametrize(("kind", "value", "expected"), [
    ("percent", 100000, 24.0),
    ("percent", 120000, 28.8),
    ("percent", 150000, 36.0),
    ("percent", 80000, 19.2),
    ("points", 1800, 18.0),
])
def test_typst_target_pitch_uses_powerpoint_spacing_semantics(kind, value, expected):
    ppr = _ppr()
    _spacing(ppr, "lnSpc", kind, value)
    style = resolve_paragraph_style((("direct", ppr),))
    assert typst_target_pitch(style, 20) == pytest.approx(expected)


def test_autofit_line_spacing_reduction_scales_the_source_pitch():
    ppr = _ppr()
    _spacing(ppr, "lnSpc", "percent", 120000)
    style = resolve_paragraph_style((("direct", ppr),))
    assert typst_target_pitch(style, 20, 0.1) == pytest.approx(25.92)
    assert typst_target_pitch(resolve_paragraph_style(()), 20, 0.1) == pytest.approx(21.6)


def _set_paragraph_format(paragraph, line_kind, line_value):
    ppr = paragraph._p.get_or_add_pPr()
    ppr.set("algn", "l")
    ppr.set("marL", str(20 * 12700))
    ppr.set("indent", str(-8 * 12700))
    _spacing(ppr, "lnSpc", line_kind, line_value)
    _spacing(ppr, "spcBef", "points", 400)
    _spacing(ppr, "spcAft", "points", 600)
    bullet = OxmlElement("a:buChar")
    bullet.set("char", "•")
    ppr.append(bullet)
    font = OxmlElement("a:buFont");font.set("typeface", "Arial");ppr.append(font)
    size = OxmlElement("a:buSzPct");size.set("val", "80000");ppr.append(size)
    color = OxmlElement("a:buClr")
    rgb = OxmlElement("a:srgbClr");rgb.set("val", "FF0000");color.append(rgb);ppr.append(color)


def _make_frame_text(text_frame):
    text_frame.clear()
    first = text_frame.paragraphs[0]
    run = first.add_run()
    run.text = "First line"
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run = first.add_run()
    run.text = " continued"
    run.font.name = "Arial"
    run.font.size = Pt(24)
    br = OxmlElement("a:br")
    first._p.insert(list(first._p).index(run._r), br)
    first.level = 1
    _set_paragraph_format(first, "percent", 120000)
    second = text_frame.add_paragraph()
    run = second.add_run()
    run.text = "Second paragraph"
    run.font.name = "Arial"
    run.font.size = Pt(14)
    second.level = 2
    _set_paragraph_format(second, "points", 1800)
    empty = text_frame.add_paragraph()  # Distinct from the soft line break above.
    empty_ppr = empty._p.get_or_add_pPr()
    empty_ppr.set("marL", "0"); empty_ppr.set("indent", "0")
    _spacing(empty_ppr, "lnSpc", "percent", 100000)
    _spacing(empty_ppr, "spcBef", "points", 0)
    _spacing(empty_ppr, "spcAft", "points", 0)
    empty_bullet = OxmlElement("a:buChar")
    empty_bullet.set("char", "-")
    empty_ppr.append(empty_bullet)
    end_style = OxmlElement("a:endParaRPr")
    end_style.set("sz", "4000")  # Caret style must not size an invisible blank line.
    empty._p.append(end_style)


def test_ordinary_text_and_placeholder_share_spacing_and_roundtrip(tmp_path: Path):
    source = tmp_path / "paragraphs.pptx"
    project = tmp_path / "typst-project"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[2])
    placeholder = next(shape for shape in slide.shapes.placeholders
                       if shape.placeholder_format.type == PP_PLACEHOLDER.BODY)
    placeholder.left = Pt(80); placeholder.top = Pt(130)
    placeholder.width = Pt(700); placeholder.height = Pt(300)
    placeholder.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    placeholder.text_frame.margin_left = Pt(6)
    placeholder.text_frame.margin_right = Pt(6)
    placeholder.text_frame.margin_top = Pt(3)
    placeholder.text_frame.margin_bottom = Pt(3)
    _make_frame_text(placeholder.text_frame)

    ordinary = slide.shapes.add_textbox(Pt(80), Pt(130), Pt(700), Pt(300))
    ordinary.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    ordinary.text_frame.margin_left = Pt(6)
    ordinary.text_frame.margin_right = Pt(6)
    ordinary.text_frame.margin_top = Pt(3)
    ordinary.text_frame.margin_bottom = Pt(3)
    _make_frame_text(ordinary.text_frame)
    for frame in (placeholder.text_frame, ordinary.text_frame):
        auto = OxmlElement("a:normAutofit")
        auto.set("fontScale", "90000")
        auto.set("lnSpcReduction", "10000")
        frame._txBody.bodyPr.append(auto)
    second_slide = deck.slides.add_slide(deck.slide_layouts[2])
    second_ordinary = second_slide.shapes.add_textbox(Pt(80), Pt(130), Pt(700), Pt(300))
    second_ordinary.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    second_ordinary.text_frame.margin_left = Pt(6)
    second_ordinary.text_frame.margin_right = Pt(6)
    second_ordinary.text_frame.margin_top = Pt(3)
    second_ordinary.text_frame.margin_bottom = Pt(3)
    _make_frame_text(second_ordinary.text_frame)
    auto = OxmlElement("a:normAutofit")
    auto.set("fontScale", "90000")
    auto.set("lnSpcReduction", "10000")
    second_ordinary.text_frame._txBody.bodyPr.append(auto)
    deck.save(source)

    env = os.environ.copy()
    env["PYTHONPATH"] = str(ROOT / "src") + os.pathsep + env.get("PYTHONPATH", "")
    subprocess.run(
        [sys.executable, "-m", "pptx_typst.pptx_to_typst", str(source), "--output", str(project)],
        cwd=ROOT, env=env, check=True, capture_output=True, text=True, timeout=120,
    )
    generated = (project / "main.typ").read_text(encoding="utf-8")
    assert "PptxTextFrame(" in generated
    ordinary_paras = find_macro_records(generated, "PptxParagraph")[:3]
    placeholder_record = next(item for item in find_macro_records(generated, "PptxPlaceholder")
                              if item.get("type") == '"body"')
    placeholder_paras = find_macro_records(placeholder_record.get("paragraphs", "()"),
                                           "PptxPlaceholderParagraph")
    assert len(ordinary_paras) == len(placeholder_paras) == 3
    for left, right in zip(ordinary_paras, placeholder_paras):
        for field in ("line-kind", "line-value", "leading", "before-kind", "before-value",
                      "after-kind", "after-value", "margin-left", "indent", "bullet"):
            assert left.get(field) == right.get(field), field
    assert ordinary_paras[0]["line-kind"] == '"percent"'
    assert ordinary_paras[0]["line-value"] == "120.0000"
    assert ordinary_paras[1]["line-kind"] == '"points"'
    assert ordinary_paras[1]["line-value"] == "18.0000"
    assert ordinary_paras[0]["before-value"] == "4.0000"
    assert ordinary_paras[0]["after-value"] == "6.0000"
    assert ordinary_paras[0]["bullet-font-source"] == '"direct"'
    assert ordinary_paras[0]["bullet-size-kind"] == '"percent"'
    assert ordinary_paras[0]["bullet-size-value"] == "80.0000"
    assert ordinary_paras[0]["bullet-fill-source"] == '"direct"'
    assert ordinary_paras[0]["level"] == placeholder_paras[0]["level"] == "1"
    assert ordinary_paras[1]["level"] == placeholder_paras[1]["level"] == "2"
    assert ordinary_paras[2].get("text") == placeholder_paras[2].get("text") == '""'
    assert ordinary_paras[2].get("bullet") is None
    assert placeholder_paras[2].get("bullet") is None
    assert ordinary_paras[2]["target-pitch"] == placeholder_paras[2]["target-pitch"] == "13.6080pt"
    assert 'font-scale:0.90000' in generated and 'line-spacing-reduction:0.10000' in generated
    assert generated.count('PptxPlaceholderRun(text:"\\n", font:"Arial"') >= 3

    typst = shutil.which("typst")
    if typst:
        pdf_path = project / "output.pdf"
        subprocess.run([typst, "compile", "main.typ", str(pdf_path)], cwd=project,
                       check=True, capture_output=True, text=True, timeout=120)
        try:
            import pymupdf
        except ImportError:  # Rendering-position check is optional in basic test installs.
            pymupdf = None
        if pymupdf is not None:
            with pymupdf.open(pdf_path) as rendered:
                assert len(rendered) == 2

                def first_line_origins(page):
                    return [tuple(span["origin"]) for block in page.get_text("dict")["blocks"]
                            if "lines" in block for line in block["lines"] for span in line["spans"]
                            if span["text"] == "First line"]

                paired_origins = first_line_origins(rendered[0])
                ordinary_origin = first_line_origins(rendered[1])
                assert len(paired_origins) == 2 and len(ordinary_origin) == 1
                assert all(abs(left - right) < 0.01
                           for left, right in zip(paired_origins[0], ordinary_origin[0]))
                assert all(abs(left - right) < 0.01
                           for left, right in zip(paired_origins[1], ordinary_origin[0]))

                def text_origins(page, value):
                    return [tuple(span["origin"]) for block in page.get_text("dict")["blocks"]
                            if "lines" in block for line in block["lines"] for span in line["spans"]
                            if span["text"] == value]

                paired_paragraph_origins = text_origins(rendered[0], "Second paragraph")
                ordinary_paragraph_origins = text_origins(rendered[1], "Second paragraph")
                assert len(paired_paragraph_origins) == 2 and len(ordinary_paragraph_origins) == 1
                assert all(abs(left - right) < 0.01
                           for left, right in zip(paired_paragraph_origins[0], ordinary_paragraph_origins[0]))
                assert all(abs(left - right) < 0.01
                           for left, right in zip(paired_paragraph_origins[1], ordinary_paragraph_origins[0]))

    output = tmp_path / "paragraphs-roundtrip.pptx"
    convert_to_pptx(project / "main.typ", output)
    with ZipFile(output) as package:
        slide_xml = etree.fromstring(package.read("ppt/slides/slide1.xml"))
        ordinary_shapes = slide_xml.xpath(".//p:sp[not(p:nvSpPr/p:nvPr/p:ph)]", namespaces=NS)
        placeholder_shapes = slide_xml.xpath(".//p:sp[p:nvSpPr/p:nvPr/p:ph[@type='body']]", namespaces=NS)
        assert len(ordinary_shapes) == 1 and len(placeholder_shapes) == 1
        ordinary_ln = ordinary_shapes[0].xpath("string(.//a:p[1]/a:pPr/a:lnSpc/a:spcPct/@val)", namespaces=NS)
        ordinary_exact = ordinary_shapes[0].xpath("string(.//a:p[2]/a:pPr/a:lnSpc/a:spcPts/@val)", namespaces=NS)
        assert ordinary_ln == "120000" and ordinary_exact == "1800"
        assert ordinary_shapes[0].xpath("string(.//a:p[1]/a:pPr/a:spcBef/a:spcPts/@val)", namespaces=NS) == "400"
        assert ordinary_shapes[0].xpath("string(.//a:p[1]/a:pPr/a:spcAft/a:spcPts/@val)", namespaces=NS) == "600"
        assert ordinary_shapes[0].xpath("string(.//a:p[1]/a:pPr/a:buFont/@typeface)", namespaces=NS) == "Arial"
        assert ordinary_shapes[0].xpath("string(.//a:p[1]/a:pPr/a:buSzPct/@val)", namespaces=NS) == "80000"
        assert ordinary_shapes[0].xpath("string(.//a:p[1]/a:pPr/a:buClr/a:srgbClr/@val)", namespaces=NS) == "FF0000"
        assert ordinary_shapes[0].xpath("string(.//a:p[1]/a:pPr/@lvl)", namespaces=NS) == "1"
        assert ordinary_shapes[0].xpath("string(.//a:p[2]/a:pPr/@lvl)", namespaces=NS) == "2"
        assert ordinary_shapes[0].xpath("count(.//a:p[1]/a:br)", namespaces=NS) == 1
        assert len(ordinary_shapes[0].xpath(".//a:p", namespaces=NS)) == 3
        assert placeholder_shapes[0].xpath("string(.//a:p[1]/a:pPr/a:lnSpc/a:spcPct/@val)", namespaces=NS) == "120000"
        ordinary_fit = ordinary_shapes[0].xpath("string(.//a:bodyPr/a:normAutofit/@fontScale)", namespaces=NS)
        assert ordinary_fit == "90000"
        placeholder_fit = placeholder_shapes[0].xpath("string(.//a:bodyPr/a:normAutofit/@fontScale)", namespaces=NS)
        assert placeholder_fit == "90000"


def test_placeholder_spacing_inheritance_stays_on_layout(tmp_path: Path):
    source = tmp_path / "inherited.pptx"
    project = tmp_path / "typst-project"
    deck = Presentation()
    layout = deck.slide_layouts[2]
    layout_body = next(shape for shape in layout.placeholders
                       if shape.placeholder_format.type == PP_PLACEHOLDER.BODY)
    inherited_ppr = layout_body.text_frame.paragraphs[0]._p.get_or_add_pPr()
    _spacing(inherited_ppr, "lnSpc", "percent", 150000)
    slides = []
    for content, override in (("Inherited spacing", None), ("Slide override", 80000)):
        slide = deck.slides.add_slide(layout)
        body = next(shape for shape in slide.shapes.placeholders
                    if shape.placeholder_format.type == PP_PLACEHOLDER.BODY)
        body.text = content
        slides.append(body)
        if override is not None:
            _spacing(body.text_frame.paragraphs[0]._p.get_or_add_pPr(), "lnSpc", "percent", override)
    deck.save(source)

    env = os.environ.copy()
    env["PYTHONPATH"] = str(ROOT / "src") + os.pathsep + env.get("PYTHONPATH", "")
    subprocess.run(
        [sys.executable, "-m", "pptx_typst.pptx_to_typst", str(source), "--output", str(project)],
        cwd=ROOT, env=env, check=True, capture_output=True, text=True, timeout=120,
    )
    generated = (project / "main.typ").read_text(encoding="utf-8")
    paragraph_records = find_macro_records(generated, "PptxPlaceholderParagraph")
    assert any(item.get("line-source") == '"layout"' and item.get("line-value") == "150.0000"
               for item in paragraph_records)
    assert any(item.get("line-source") == '"direct"' and item.get("line-value") == "80.0000"
               for item in paragraph_records)

    output = tmp_path / "inherited-roundtrip.pptx"
    convert_to_pptx(project / "main.typ", output)
    with ZipFile(output) as package:
        layout_files = [name for name in package.namelist()
                        if name.startswith("ppt/slideLayouts/slideLayout") and name.endswith(".xml")]
        layout_xml = next(etree.fromstring(package.read(name)) for name in layout_files
                          if etree.fromstring(package.read(name)).xpath(
                              './/p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]', namespaces=NS))
        assert layout_xml.xpath(
            'string(.//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]//a:lvl1pPr/a:lnSpc/a:spcPct/@val)',
            namespaces=NS,
        ) == "150000"
        first_slide = etree.fromstring(package.read("ppt/slides/slide1.xml"))
        second_slide = etree.fromstring(package.read("ppt/slides/slide2.xml"))
        first_body = first_slide.xpath('.//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]', namespaces=NS)[0]
        second_body = second_slide.xpath('.//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]', namespaces=NS)[0]
        assert not first_body.xpath(".//a:p[1]/a:pPr/a:lnSpc", namespaces=NS)
        assert second_body.xpath("string(.//a:p[1]/a:pPr/a:lnSpc/a:spcPct/@val)", namespaces=NS) == "80000"


def test_master_placeholder_spacing_is_inherited_and_roundtrips_at_master(tmp_path: Path):
    source = tmp_path / "master-spacing.pptx"
    project = tmp_path / "typst-project"
    deck = Presentation()
    layout = deck.slide_layouts[2]
    master_body = next(shape for shape in layout.slide_master.placeholders
                       if shape.placeholder_format.type == PP_PLACEHOLDER.BODY)
    _spacing(master_body.text_frame.paragraphs[0]._p.get_or_add_pPr(), "lnSpc", "percent", 135000)
    slide = deck.slides.add_slide(layout)
    body = next(shape for shape in slide.shapes.placeholders
                if shape.placeholder_format.type == PP_PLACEHOLDER.BODY)
    paragraph = body.text_frame.paragraphs[0]
    paragraph.text = "Master spacing"
    paragraph._p.get_or_add_pPr().set("algn", "ctr")
    deck.save(source)

    env = os.environ.copy()
    env["PYTHONPATH"] = str(ROOT / "src") + os.pathsep + env.get("PYTHONPATH", "")
    subprocess.run(
        [sys.executable, "-m", "pptx_typst.pptx_to_typst", str(source), "--output", str(project)],
        cwd=ROOT, env=env, check=True, capture_output=True, text=True, timeout=120,
    )
    generated = (project / "main.typ").read_text(encoding="utf-8")
    record = next(item for item in find_macro_records(generated, "PptxPlaceholderParagraph")
                  if item.get("line-source") == '"master"')
    assert record["line-source"] == '"master"'
    assert record["line-value"] == "135.0000"
    assert record["align-source"] == '"direct"'

    output = tmp_path / "master-spacing-roundtrip.pptx"
    convert_to_pptx(project / "main.typ", output)
    with ZipFile(output) as package:
        master_files = [name for name in package.namelist()
                        if name.startswith("ppt/slideMasters/slideMaster") and name.endswith(".xml")]
        master = next(etree.fromstring(package.read(name)) for name in master_files
                      if etree.fromstring(package.read(name)).xpath(
                          './/p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]', namespaces=NS))
        assert master.xpath(
            'string(.//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]//a:lvl1pPr/a:lnSpc/a:spcPct/@val)',
            namespaces=NS,
        ) == "135000"
        slide_xml = etree.fromstring(package.read("ppt/slides/slide1.xml"))
        slide_body = slide_xml.xpath('.//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]', namespaces=NS)[0]
        assert not slide_body.xpath(".//a:p[1]/a:pPr/a:lnSpc", namespaces=NS)
        assert slide_body.xpath("string(.//a:p[1]/a:pPr/@algn)", namespaces=NS) == "ctr"
