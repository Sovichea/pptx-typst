#!/usr/bin/env python3
"""Convert PowerPoint slides to editable Typst source and shared media assets."""
from __future__ import annotations
from pathlib import Path
from copy import deepcopy
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from zipfile import ZipFile
import numpy as np
import json,math,hashlib,os,shutil,re
from html import escape as xml_escape
from lxml import etree
from pptx.oxml.ns import qn
from pptx.dml.color import RGBColor
from collections import Counter,defaultdict

import argparse
from .pptx_geometry import shape_paths, preset_path
from .pptx_mask import mask_image
from .media_compat import convert_metafile
from .paragraph_layout import resolve_paragraph_style, typst_target_pitch, typst_spacing

cli=argparse.ArgumentParser(description="Convert PPTX to structured Typst with shared master/layout declarations and original media assets")
cli.add_argument("input_pptx",type=Path,help="source PowerPoint document")
cli.add_argument("--output",type=Path,default=Path("./typst-output"),help="Typst project output directory")
cli.add_argument("--template",type=Path,default=Path(__file__).with_name("template.typ"),help="Typst template defining Root, Place, and related layout helpers")
cli.add_argument("--text-layout",choices=["v3","anchors","spacing","anchors-spacing"],default="anchors-spacing",help="legacy compatibility selector for vertical anchoring; paragraph spacing uses the shared paragraph model")
cli.add_argument("--text-features", choices=["legacy","inherit","structure","all"], default="all", help="Text inheritance/bullets/tabs features (default: all)")
cli.add_argument("--group-mode",choices=["safe","nested","flatten"],default="safe",help="Emit editable nested GroupContent containers or v5 flattened placements")
cli.add_argument("--clip-mode",choices=["component","legacy"],default="legacy",help="Rectangular picture crop representation")
cli.add_argument("--mask-mode",choices=["legacy","validated","experimental"],default="legacy",help="Conservative legacy default; opt in to preset or experimental curved masks after visual validation")
cli.add_argument("--mask-allowlist",type=Path,help="Optional JSON list of placements allowed to use nonrectangular masking, keyed by slide and shape ID")
cli.add_argument("--layout-mode",choices=["structured","flat"],default="structured",help="Keep reusable PowerPoint master/layout artwork as declarations, or repeat it in each slide for legacy output")
args=cli.parse_args()
MASK_ALLOWLIST=None
if args.mask_allowlist:
    data=json.loads(args.mask_allowlist.read_text())
    MASK_ALLOWLIST={(int(v["slide"]),int(v["shape_id"])) for v in data["placements"]}
SRC=args.input_pptx.resolve()
OUT=args.output.resolve()
ASSETS=OUT/'assets/media';ASSETS.mkdir(parents=True,exist_ok=True)
EMU=12700
SLIDE_W=Presentation(SRC).slide_width/EMU
SLIDE_H=Presentation(SRC).slide_height/EMU
# PowerPoint scheme colours are resolved through the master theme and clrMap.
THEME_CACHE={}
CURRENT_THEME={}
CURRENT_MAP={}
CURRENT_FONTS={"major":"Aptos Display","minor":"Aptos"}
CURRENT_SLIDE=None
DRAWINGML_NS={
    'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
}

def xpath_ns(element,path):
    """Run an XPath against either python-pptx XML or a plain lxml element."""
    try:return element.xpath(path)
    except etree.XPathEvalError:return element.xpath(path,namespaces=DRAWINGML_NS)

def theme_context(slide):
    master=slide.slide_layout.slide_master
    theme_part=next((rel.target_part for rel in master.part.rels.values() if rel.reltype.endswith('/theme') and not rel.is_external),None)
    key=str(theme_part.partname) if theme_part else 'default'
    if key not in THEME_CACHE:
        palette={}
        if theme_part:
            root=etree.fromstring(theme_part.blob)
            schemes=root.xpath('//*[local-name()="clrScheme"]')
            if schemes:
                for el in schemes[0]:
                    tag=etree.QName(el).localname
                    ch=list(el)
                    if ch:
                        c=ch[0];rgb=c.get('val') if etree.QName(c).localname=='srgbClr' else c.get('lastClr')
                        if rgb:palette[tag]='#'+rgb.upper()
        THEME_CACHE[key]=palette
    remap={}
    master_maps=master._element.xpath('./p:clrMap')
    if master_maps:remap.update(master_maps[0].attrib)
    override=slide.slide_layout._element.xpath('./p:clrMapOvr/a:overrideClrMapping')
    if override:remap.update(override[0].attrib)
    override=slide._element.xpath('./p:clrMapOvr/a:overrideClrMapping')
    if override:remap.update(override[0].attrib)
    return THEME_CACHE[key],remap

def theme_font_context(slide):
    master=slide.slide_layout.slide_master
    theme_part=next((rel.target_part for rel in master.part.rels.values() if rel.reltype.endswith('/theme') and not rel.is_external),None)
    out={'major':'Aptos Display','minor':'Aptos'}
    if theme_part is not None:
        root=etree.fromstring(theme_part.blob)
        for key in ('major','minor'):
            el=root.xpath('//*[local-name()="'+key+'Font"]/*[local-name()="latin"]')
            if el and el[0].get('typeface'):out[key]=el[0].get('typeface')
    return out

def oxml_font(node):
    if node is None:return None
    els=node.xpath('./*[local-name()="latin"]')
    return els[0].get('typeface') if els else None

def oxml_size(node):
    if node is None:return None
    value=node.get('sz')
    if not value:return None
    try:return int(value)/100
    except Exception:return None

def expand_theme_font(name):
    if not name:return None
    if name.startswith('+mj'):return CURRENT_FONTS['major']
    if name.startswith('+mn'):return CURRENT_FONTS['minor']
    return name

def layout_placeholder(sh):
    """Find the layout peer by placeholder ID, not collection position."""
    if not sh.is_placeholder or CURRENT_SLIDE is None:return None
    idx=sh.placeholder_format.idx
    return next((ph for ph in CURRENT_SLIDE.slide_layout.placeholders
                 if ph.placeholder_format.idx==idx),None)

def text_style_defaults(sh,pa):
    # Run formatting cascades through pPr/defRPr, then lstStyle, then
    # placeholder defaults and finally major/minor theme fonts.
    lvl=min(9,max(1,pa.level+1))
    result=[]
    def add(node):
        if node is not None and node not in result:result.append(node)
    pp=pa._p.xpath('./a:pPr/a:defRPr')
    if pp:add(pp[0])
    ls=sh._element.xpath(f'./p:txBody/a:lstStyle/a:lvl{lvl}pPr/a:defRPr')
    if ls:add(ls[0])
    if sh.is_placeholder and CURRENT_SLIDE is not None:
        try:
            ph=layout_placeholder(sh)
            if ph:
                x=ph._element.xpath(f'./p:txBody/a:lstStyle/a:lvl{lvl}pPr/a:defRPr')
                if x:add(x[0])
        except Exception:pass
    # Master text styles vary for title, body, and ordinary text boxes.
    if CURRENT_SLIDE is not None:
        master=CURRENT_SLIDE.slide_layout.slide_master
        is_title=(sh.is_placeholder and 'TITLE' in str(sh.placeholder_format.type).upper())
        contexts=['titleStyle'] if is_title else (['bodyStyle','otherStyle'] if sh.is_placeholder else ['otherStyle'])
        for context in contexts:
            x=master._element.xpath(f'./p:txStyles/p:{context}/a:lvl{lvl}pPr/a:defRPr')
            if x:add(x[0])
    return result

def effective_run_style(sh,pa,run):
    rp=run._r.xpath('./a:rPr');rp=rp[0] if rp else None
    defaults=text_style_defaults(sh,pa)
    font=run.font.name or oxml_font(rp) or next((oxml_font(el) for el in defaults if oxml_font(el)),None)
    if not font:
        style=sh._element.xpath('./p:style/a:fontRef')
        idx=style[0].get('idx') if style else 'minor'
        font=('+mj-lt' if idx=='major' else '+mn-lt')
    font=font_mapping(expand_theme_font(font))
    sz=(run.font.size.pt if run.font.size is not None else None)
    if sz is None:sz=oxml_size(rp) or next((oxml_size(el) for el in defaults if oxml_size(el)),None) or 20
    c=color_of(run) or next((color_in(el,'./a:solidFill') for el in defaults if color_in(el,'./a:solidFill')),None) or shape_text_color(sh)
    bold=run.font.bold
    if bold is None:
        bold=next((el.get('b') in ('1','true') for el in ([rp]+defaults) if el is not None and el.get('b') is not None),False)
    italic=run.font.italic
    if italic is None:
        italic=next((el.get('i') in ('1','true') for el in ([rp]+defaults) if el is not None and el.get('i') is not None),False)
    return font,sz,c,('bold' if bold else 'regular'),bool(italic)

def paragraph_style_source_records(sh,pa):
    """Return DrawingML paragraph nodes in inheritance priority order."""
    lvl=min(9,max(1,pa.level+1))
    result=[]
    def add(source,nodes):
        for node in nodes or []:
            if all(node is not current for _,current in result):
                result.append((source,node))
    add('direct',pa._p.xpath('./a:pPr'))
    add('shape',sh._element.xpath(f'./p:txBody/a:lstStyle/a:lvl{lvl}pPr'))
    add('shape',sh._element.xpath('./p:txBody/a:lstStyle/a:defPPr'))
    if sh.is_placeholder and CURRENT_SLIDE is not None:
        idx=sh.placeholder_format.idx
        layout=CURRENT_SLIDE.slide_layout
        master=layout.slide_master
        for origin,owner in (('layout',layout),('master',master)):
            peer=next((item for item in owner.placeholders if item.placeholder_format.idx==idx),None)
            if peer is not None and peer is not sh:
                # Placeholder-local paragraph properties precede its level list
                # style. This matters when only one paragraph property is set.
                peer_paragraphs=peer._element.xpath('./p:txBody/a:p/a:pPr')
                add(origin,peer_paragraphs[:1])
                add(origin,peer._element.xpath(f'./p:txBody/a:lstStyle/a:lvl{lvl}pPr'))
                add(origin,peer._element.xpath('./p:txBody/a:lstStyle/a:defPPr'))
        is_title='TITLE' in str(sh.placeholder_format.type).upper()
        contexts=['titleStyle'] if is_title else ['bodyStyle','otherStyle']
        for context in contexts:
            add('master_style',master._element.xpath(f'./p:txStyles/p:{context}/a:lvl{lvl}pPr'))
            add('master_style',master._element.xpath(f'./p:txStyles/p:{context}/a:defPPr'))
    elif CURRENT_SLIDE is not None:
        master=CURRENT_SLIDE.slide_layout.slide_master
        add('master_style',master._element.xpath(f'./p:txStyles/p:otherStyle/a:lvl{lvl}pPr'))
        add('master_style',master._element.xpath('./p:txStyles/p:otherStyle/a:defPPr'))
    return result

def paragraph_style_sources(sh,pa):
    """Compatibility view of paragraph properties ordered by precedence."""
    return [node for _,node in paragraph_style_source_records(sh,pa)]

def paragraph_indent(sh,pa):
    """Resolve DrawingML margin and hanging indent, expressed in points."""
    sources=paragraph_style_sources(sh,pa)
    def value(name):
        for ppr in sources:
            raw=ppr.get(name)
            if raw is not None:
                try:return int(raw)/EMU
                except ValueError:pass
        return 0.0
    return value('marL'),value('indent')

def normalized_bullet_glyph(glyph, ppr):
    """Map legacy font-encoded bullet codepoints to portable Unicode glyphs."""
    fonts=ppr.xpath('./*[local-name()="buFont"]/@typeface') if ppr is not None else []
    if glyph and ord(glyph[0])==0x00A7 and fonts and fonts[0].lower()=="wingdings":
        return chr(0x25AA)
    return glyph

def bullet_spec(sh,pa,fontsize):
    source_records=paragraph_style_source_records(sh,pa)
    sources=[node for _,node in source_records]
    node=None;owner=None
    owner_source=None
    for source,ppr in source_records:
        found=ppr.xpath('./*[local-name()="buNone" or local-name()="buChar" or local-name()="buAutoNum" or local-name()="buBlip"]')
        if found:
            node=found[0];owner=ppr;owner_source=source;break
    if node is None:return None
    tag=etree.QName(node).localname
    if tag=='buNone':return None
    if tag=='buBlip':spec={'kind':'unsupported','glyph':chr(0x2022)}
    elif tag=='buChar':
        glyph=normalized_bullet_glyph(node.get('char',chr(0x2022)),owner)
        spec={'kind':'char','glyph':glyph}
    elif tag=='buAutoNum':spec={'kind':'number','type':node.get('type','arabicPeriod'),'start':int(node.get('startAt','1'))}
    else:return None
    spec['fontsize']=fontsize
    spec['source']=owner_source
    for source,ppr in source_records:
        fonts=ppr.xpath('./*[local-name()="buFont"]')
        if fonts and fonts[0].get('typeface'):
            spec['font']=expand_theme_font(fonts[0].get('typeface'))
            spec['font_source']=source;break
    if owner is not None:
        pct=owner.xpath('./*[local-name()="buSzPct"]')
        pts=owner.xpath('./*[local-name()="buSzPts"]')
        if pts and pts[0].get('sz'):
            spec['fontsize']=int(pts[0].get('sz'))/100;spec['absolute_size']=True
            spec['size_kind']='points';spec['size_value']=spec['fontsize'];spec['size_source']=owner_source
        elif pct and pct[0].get('val'):
            spec['size_kind']='percent';spec['size_value']=int(pct[0].get('val'))/1000
            spec['size_source']=owner_source
            spec['fontsize']=fontsize*spec['size_value']/100
        colors=owner.xpath('./*[local-name()="buClr"]/*[local-name()="srgbClr" or local-name()="schemeClr" or local-name()="sysClr" or local-name()="prstClr"] | ./*[local-name()="buClr"]/*[local-name()="solidFill"]/*[local-name()="srgbClr" or local-name()="schemeClr" or local-name()="sysClr" or local-name()="prstClr"]')
        spec['color']=resolve_color_node(colors[0]) if colors else None
        if colors:spec['color_source']=owner_source
    return spec

def number_glyph(n,kind):
    def letters(i):
        z=''
        while i:
            i,x=divmod(i-1,26);z=chr(65+x)+z
        return z
    def roman(i):
        result=''
        for value,glyph in ((1000,'M'),(900,'CM'),(500,'D'),(400,'CD'),(100,'C'),(90,'XC'),(50,'L'),(40,'XL'),(10,'X'),(9,'IX'),(5,'V'),(4,'IV'),(1,'I')):
            while i>=value:i-=value;result+=glyph
        return result
    if kind.startswith('alpha'):
        t=letters(n);t=t.lower() if 'Lc' in kind else t
    elif kind.startswith('roman'):
        t=roman(n);t=t.lower() if 'Lc' in kind else t
    else:t=str(n)
    if kind.endswith('ParenBoth'):return '('+t+')'
    if kind.endswith('ParenR'):return t+')'
    if kind.endswith('Period'):return t+'.'
    return t

def pptx_tabs(pa):
    # Positions in EMU, relative to the left inner text box edge.
    x=pa._p.xpath('./a:pPr/a:tabLst/a:tab')
    return [(max(0,int(n.get('pos','0')))/EMU,n.get('algn','l')) for n in x]

def body_tab_size(sh,pa):
    ppr=pa._p.xpath('./a:pPr')
    if ppr and ppr[0].get('defTabSz'):
        return int(ppr[0].get('defTabSz'))/EMU
    return 36.0

def resolve_color_node(node, default=None, color_map=None):
    if node is None:return default
    k=etree.QName(node).localname
    if k=='srgbClr':color='#'+(node.get('val') or '000000')
    elif k=='sysClr':color='#'+(node.get('lastClr') or '000000')
    elif k=='schemeClr':
        key=node.get('val','')
        key=(CURRENT_MAP if color_map is None else color_map).get(key,key)
        color=CURRENT_THEME.get(key,default or '#000000')
    elif k=='prstClr':
        table={'black':'#000000','white':'#FFFFFF','red':'#FF0000','blue':'#0000FF','gray':'#808080','yellow':'#FFFF00'}
        color=table.get(node.get('val',''),default or '#000000')
    else:return default
    try:
        rgb=[int(color[i:i+2],16) for i in (1,3,5)]
        for change in node:
            t=etree.QName(change).localname;v=int(change.get('val','100000'))/100000
            if t=='lumMod':rgb=[int(max(0,min(255,c*v))) for c in rgb]
            elif t=='lumOff':rgb=[int(max(0,min(255,c+255*v))) for c in rgb]
            elif t=='tint':rgb=[int(max(0,min(255,c*(1-v)+255*v))) for c in rgb]
            elif t=='shade':rgb=[int(max(0,min(255,c*v))) for c in rgb]
        color='#'+''.join(f'{c:02X}' for c in rgb)
    except Exception:pass
    return color

def color_in(parent, prefix, default=None, color_map=None):
    el=xpath_ns(parent,prefix+'/a:srgbClr | '+prefix+'/a:schemeClr | '+prefix+'/a:sysClr | '+prefix+'/a:prstClr')
    return resolve_color_node(el[0],default,color_map) if el else default

def color_opacity(node):
    """Resolve DrawingML alpha transforms on one color node."""
    opacity=1.0
    if node is None:return opacity
    for transform in node:
        kind=etree.QName(transform).localname
        value=float(transform.get('val','100000'))/100000
        if kind=='alpha':opacity=value
        elif kind=='alphaMod':opacity*=value
        elif kind=='alphaOff':opacity+=value
    return max(0.0,min(1.0,opacity))

def color_opacity_in(parent, prefix):
    nodes=xpath_ns(parent,prefix+'/a:srgbClr | '+prefix+'/a:schemeClr | '+prefix+'/a:sysClr | '+prefix+'/a:prstClr')
    return color_opacity(nodes[0]) if nodes else 1.0

def typst_color_expr(node, default='#FFFFFF'):
    """Represent a DrawingML color and its alpha in Typst's 8-digit hex form."""
    color=resolve_color_node(node,default)
    if not color:return 'none'
    opacity=color_opacity(node)
    if opacity<.99999:
        color+=f'{round(opacity*255):02X}'
    return f'rgb("{color}")'

def gradient_expr(gf,color_map=None):
    stops=[]
    for gs in xpath_ns(gf,'./a:gsLst/a:gs'):
        node=next((child for child in gs if etree.QName(child).localname in ('srgbClr','schemeClr','sysClr','prstClr')),None)
        if node is None:continue
        rgb=resolve_color_node(node,'#000000',color_map)
        alpha=next((child for child in node if etree.QName(child).localname=='alpha'),None)
        opacity=float(alpha.get('val','100000'))/100000 if alpha is not None else 1
        if opacity<.999:rgb=rgb+f'{round(opacity*255):02X}'
        stops.append((int(gs.get('pos','0')),rgb))
    stops.sort()
    if len(stops)<2:return 'white'
    # Typst requires endpoints to be explicitly present. OOXML permits 1..99%.
    if stops[0][0]>0:stops.insert(0,(0,stops[0][1]))
    if stops[-1][0]<100000:stops.append((100000,stops[-1][1]))
    values=[f'(rgb("{c}"),{pos/1000:.3f}%)' for pos,c in stops]
    ln=xpath_ns(gf,'./a:lin')
    if not ln:return 'gradient.radial('+','.join(values)+')'
    a=float(ln[0].get('ang','0'))/60000
    return 'gradient.linear('+','.join(values)+f',angle:{a:.3f}deg)'

def owner_background_color_map(owner):
    """Resolve theme color mappings at the level that owns a background."""
    part=str(owner.part.partname)
    if '/slideMasters/' in part:
        master=owner; overrides=[]
    elif '/slideLayouts/' in part:
        master=owner.slide_master; overrides=[owner]
    else:
        master=owner.slide_layout.slide_master
        overrides=[owner.slide_layout,owner]
    maps=master._element.xpath('./p:clrMap')
    result=dict(maps[0].attrib) if maps else {}
    for item in overrides:
        override=item._element.xpath('./p:clrMapOvr/a:overrideClrMapping')
        if override:result.update(override[0].attrib)
    return result


def owner_theme_part(owner):
    part=str(owner.part.partname)
    if '/slideMasters/' in part:master=owner
    elif '/slideLayouts/' in part:master=owner.slide_master
    else:master=owner.slide_layout.slide_master
    return next((rel.target_part for rel in master.part.rels.values()
                 if rel.reltype.endswith('/theme') and not rel.is_external),None)


def theme_background_style(owner,bg_ref,color_map):
    """Resolve a p:bgRef style through the owning master theme."""
    try:style_index=int(bg_ref.get('idx','0'))-1001
    except (TypeError,ValueError):return None
    if style_index<0:return None
    theme_part=owner_theme_part(owner)
    if theme_part is None:return None
    root=etree.fromstring(theme_part.blob)
    styles=xpath_ns(root,'./a:themeElements/a:fmtScheme/a:bgFillStyleLst/*')
    if style_index>=len(styles):return None
    style=deepcopy(styles[style_index])
    ref_colors=owner._element.xpath(
        './p:cSld/p:bg/p:bgRef/a:srgbClr | ./p:cSld/p:bg/p:bgRef/a:schemeClr | '
        './p:cSld/p:bg/p:bgRef/a:sysClr | ./p:cSld/p:bg/p:bgRef/a:prstClr')
    if ref_colors:
        ref_color=ref_colors[0]
        for placeholder in xpath_ns(style,'.//a:schemeClr[@val="phClr"]'):
            replacement=deepcopy(ref_color)
            for transform in placeholder:
                replacement.append(deepcopy(transform))
            placeholder.getparent().replace(placeholder,replacement)
    style_name=etree.QName(style).localname
    solid=xpath_ns(style,'./a:srgbClr | ./a:schemeClr | ./a:sysClr | ./a:prstClr') if style_name=='solidFill' else []
    if solid:
        color=resolve_color_node(solid[0],color_map=color_map)
        if color:
            opacity=color_opacity(solid[0])
            if opacity<.99999:color+=f'{round(opacity*255):02X}'
            return ('color',color,None)
    if style_name=='gradFill':return ('gradient',style,None)
    blips=xpath_ns(style,'./a:blip') if style_name=='blipFill' else []
    if blips:
        rid=blips[0].get(qn('r:embed'))
        if rid and rid in theme_part.rels:
            return ('image',theme_part.related_part(rid),style)
    return None


def owner_background(owner,color_map=None):
    """Return only the background declared by this PPTX part, if any."""
    color_map=owner_background_color_map(owner) if color_map is None else color_map
    # PowerPoint commonly stores a master/layout background as a theme style
    # reference instead of a p:bgPr fill. Resolve the referenced scheme color
    # through the current slide's theme and color map just like shape colors.
    bg_ref=owner._element.xpath('./p:cSld/p:bg/p:bgRef')
    if bg_ref:
        if bg_ref[0].get('idx')=='0':return None
        styled=theme_background_style(owner,bg_ref[0],color_map)
        if styled:return styled
        color_nodes=owner._element.xpath(
            './p:cSld/p:bg/p:bgRef/a:srgbClr | ./p:cSld/p:bg/p:bgRef/a:schemeClr | '
            './p:cSld/p:bg/p:bgRef/a:sysClr | ./p:cSld/p:bg/p:bgRef/a:prstClr')
        if color_nodes:
            color=resolve_color_node(color_nodes[0],color_map=color_map)
            if color:
                opacity=color_opacity(color_nodes[0])
                if opacity<.99999:color+=f'{round(opacity*255):02X}'
                return ('color',color,None)
    bg=owner._element.xpath('./p:cSld/p:bg/p:bgPr')
    if not bg:return None
    el=bg[0]
    if el.xpath('./a:noFill'):return None
    col=color_in(el,'./a:solidFill',color_map=color_map)
    if col:return ('color',col,None)
    grad=el.xpath('./a:gradFill')
    if grad:return ('gradient',grad[0],None)
    blip=el.xpath('./a:blipFill/a:blip')
    if blip:
        rid=blip[0].get(qn('r:embed'))
        if rid and rid in owner.part.rels:return ('image',owner.part.related_part(rid),el)
    return None


def source_background(slide):
    # Resolve slide, layout, then master without changing the owning layer.
    for owner in (slide,slide.slide_layout,slide.slide_layout.slide_master):
        # Inherited background references are colored in the slide's effective
        # clrMap context, which may differ from the master's own theme mapping.
        background=owner_background(owner,color_map=CURRENT_MAP)
        if background:return background
    return ('color','#FFFFFF',None)


# Preserve user-provided Place-based template, which remains the presentation API.
if args.template.resolve() != (OUT/'template.typ').resolve():
    shutil.copy2(args.template,OUT/'template.typ')

prs=Presentation(SRC)
TABLE_STYLES_ROOT=None
for rel in prs.part.rels.values():
    if rel.reltype.endswith('/tableStyles') and not rel.is_external:
        try:TABLE_STYLES_ROOT=etree.fromstring(rel.target_part.blob)
        except Exception:pass
def stylize(s):
 return '\\"'.join([]) if False else '"'+s.replace('\\','\\\\').replace('"','\\"').replace('\n','\\n').replace('\t',' ' ).replace('\v','\\n')+'"'

def typst_comment(value):
    """Collapse controls and line separators so source metadata stays a comment."""
    return re.sub(r'\s+',' ',re.sub(r'[\x00-\x1f\x7f-\x9f\u2028\u2029]+',' ',str(value))).strip()

def typst_place(x,y,child,rotation=0,indent='    '):
    """Write a readable multiline Place call while keeping its child editable."""
    lines=[f'{indent}#Place(',f'{indent}  x:{x:.3f}pt,',f'{indent}  y:{y:.3f}pt,',
           f'{indent}  orientation:{rotation:.3f}deg,',f'{indent}  child:']
    lines.extend(f'{indent}    {line}' for line in str(child).splitlines())
    lines[-1]+=','
    lines.append(f'{indent})')
    return '\n'.join(lines)

PLACEHOLDER_KINDS={
    'TITLE':'title', 'CENTER_TITLE':'centered-title', 'BODY':'body',
}

def placeholder_kind(shape):
    if not shape.is_placeholder:return None
    raw=str(shape.placeholder_format.type).split(' ')[0]
    return PLACEHOLDER_KINDS.get(raw)

def placeholder_id(kind,idx):
    return f'{kind}-{int(idx)}'

def placeholder_direct_geometry(shape):
    xfrm=shape._element.xpath('./p:spPr/a:xfrm')
    if not xfrm:return None
    off=xfrm[0].xpath('./a:off');ext=xfrm[0].xpath('./a:ext')
    if not off or not ext:return None
    return {
        'x':float(off[0].get('x','0'))/EMU,
        'y':float(off[0].get('y','0'))/EMU,
        'width':float(ext[0].get('cx','0'))/EMU,
        'height':float(ext[0].get('cy','0'))/EMU,
        'rotation':float(xfrm[0].get('rot','0'))/60000,
    }

def placeholder_style_fields(shape, include_inherited=False):
    """Read local placeholder styles and optionally resolve inherited font family."""
    fields={}
    tx=shape._element.xpath('./p:txBody')
    if not tx:return fields
    ppr=tx[0].xpath('./a:p/a:pPr')
    levels=tx[0].xpath('./a:lstStyle/a:lvl1pPr')
    align_sources=ppr+levels
    if include_inherited and getattr(shape,'is_placeholder',False) and CURRENT_SLIDE is not None:
        idx=shape.placeholder_format.idx
        for owner in (CURRENT_SLIDE.slide_layout,CURRENT_SLIDE.slide_layout.slide_master):
            peer=next((item for item in owner.placeholders
                       if item.placeholder_format.idx==idx),None)
            if peer is not None and peer is not shape:
                align_sources.extend(peer._element.xpath('./p:txBody/a:p/a:pPr | ./p:txBody/a:lstStyle/a:lvl1pPr'))
        if not align_sources:
            kind=placeholder_kind(shape)
            context='titleStyle' if kind in ('title','centered-title') else 'bodyStyle'
            master=CURRENT_SLIDE.slide_layout.slide_master
            align_sources.extend(master._element.xpath(f'./p:txStyles/p:{context}/a:lvl1pPr'))
    alignment=next((node.get('algn') for node in align_sources if node.get('algn')),None)
    if alignment:
        fields['align']={'l':'left','ctr':'center','r':'right','just':'justify'}.get(alignment)
    body=tx[0].xpath('./a:bodyPr')
    if body:
        attrs={'lIns':'margin-left','rIns':'margin-right','tIns':'margin-top','bIns':'margin-bottom'}
        for attr,key in attrs.items():
            if body[0].get(attr) is not None:fields[key]=float(body[0].get(attr))/EMU
        if body[0].get('anchor'):
            fields['anchor']={'t':'top','ctr':'center','b':'bottom'}.get(body[0].get('anchor'))
    background=shape_background_expr(shape)
    if background is not None:
        fields['background']=background
    def font_defaults(owner):
        tx_body=owner._element.xpath('./p:txBody')
        if not tx_body:return []
        return (tx_body[0].xpath('./a:lstStyle/a:lvl1pPr/a:defRPr')+
                tx_body[0].xpath('./a:p/a:pPr/a:defRPr')+
                tx_body[0].xpath('./a:p/a:endParaRPr'))

    local_style_sources=font_defaults(shape)
    font_sources=list(local_style_sources)
    if include_inherited and getattr(shape,'is_placeholder',False) and CURRENT_SLIDE is not None:
        idx=shape.placeholder_format.idx
        for owner in (CURRENT_SLIDE.slide_layout, CURRENT_SLIDE.slide_layout.slide_master):
            peer=next((item for item in owner.placeholders
                       if item.placeholder_format.idx==idx),None)
            if peer is not None and peer is not shape:
                font_sources.extend(font_defaults(peer))
    typeface=next((node.xpath('./a:latin/@typeface')[0] for node in font_sources
                   if node.xpath('./a:latin/@typeface')),None)
    if typeface:
        fields['font']=expand_theme_font(typeface)
    elif include_inherited and getattr(shape,'is_placeholder',False):
        master=CURRENT_SLIDE.slide_layout.slide_master if CURRENT_SLIDE is not None else None
        kind=placeholder_kind(shape)
        context='titleStyle' if kind in ('title','centered-title') else 'bodyStyle'
        style=(master._element.xpath(f'./p:txStyles/p:{context}/a:lvl1pPr/a:defRPr')
               if master is not None else [])
        style_font=style[0].xpath('./a:latin/@typeface') if style else []
        if style_font:
            fields['font']=expand_theme_font(style_font[0])
        elif master is not None:
            fields['font']=CURRENT_FONTS['major' if context=='titleStyle' else 'minor']
    rpr=next((node for node in local_style_sources if node is not None),None)
    if rpr is not None:
        if rpr.get('sz'):fields['size']=float(rpr.get('sz'))/100
        if rpr.get('b') is not None:fields['weight']='bold' if rpr.get('b') in ('1','true') else 'regular'
        colors=rpr.xpath('./a:solidFill/a:srgbClr | ./a:solidFill/a:schemeClr | ./a:solidFill/a:sysClr | ./a:solidFill/a:prstClr')
        if colors:
            fields['fill']=typst_color_expr(colors[0])
    bullets=(tx[0].xpath('./a:lstStyle/a:lvl1pPr/a:buChar')+
             tx[0].xpath('./a:p/a:pPr/a:buChar'))
    if bullets and bullets[0].get('char'):
        fields['bullet']=normalized_bullet_glyph(bullets[0].get('char'),bullets[0].getparent())
    return fields

def placeholder_slot_expression(shape):
    kind=placeholder_kind(shape)
    if not kind:return None
    idx=shape.placeholder_format.idx
    fields={'id':placeholder_id(kind,idx),'type':kind,'idx':idx}
    fields.update(placeholder_direct_geometry(shape) or {})
    fields.update(placeholder_style_fields(shape,include_inherited=True))
    args=[]
    for key,value in fields.items():
        if key in ('id','type','font','weight','align','anchor','bullet'):
            rendered=stylize(str(value))
        elif key in ('fill','background'):rendered=value
        elif key=='idx':rendered=str(value)
        elif key=='rotation':rendered=f'{float(value):.3f}deg'
        else:rendered=f'{float(value):.3f}pt'
        args.append(f'{key}:{rendered}')
    return 'PptxPlaceholderSlot('+', '.join(args)+')',fields

def typst_font_fallback(f):
    """Prefer the requested family, then Typst's bundled fallback family."""
    families=[]
    for family in (font_mapping(f), 'Libertinus Serif'):
        if family and family not in families:families.append(family)
    return '('+', '.join(stylize(family) for family in families)+')'

def font_mapping(f):
    if not f:return 'Arial'
    # Preserve the original family name; typst_font_fallback adds local choices.
    return f.strip() or 'Arial'

def weight_mapping(run):
 f=(run.font.name or '').lower()
 if run.font.bold or 'bold' in f or 'semibold' in f:return 'bold'
 return 'regular'

def color_of(run):
    c=run.font.color
    try:
        if c.rgb is not None:return '#'+str(c.rgb)
    except Exception:pass
    # run may inherit its colour from a:schemeClr
    try:
        if run._r.rPr is not None:
            col=color_in(run._r.rPr,'./a:solidFill')
            if col:return col
    except Exception:pass
    return None

def p_align(p):
 if p.alignment==PP_ALIGN.CENTER:return 'center'
 if p.alignment==PP_ALIGN.RIGHT:return 'right'
 # LEFT/JUSTIFY default left; justify can be set for paragraphs
 if p.alignment==PP_ALIGN.JUSTIFY:return 'left'
 return 'left'

def inherited_run_style(sh,pa):
    # a:pPr/a:defRPr is a PowerPoint paragraph default, overridden by a:rPr.
    arr=pa._p.xpath('./a:pPr/a:defRPr')
    if not arr:
        level=pa.level+1
        arr=sh._element.xpath(f'./p:txBody/a:lstStyle/a:lvl{level}pPr/a:defRPr')
    return arr[0] if arr else None

def shape_text_color(sh):
    try:
        val=color_in(sh._element,'./p:style/a:fontRef')
        if val:return val
        dr=sh._element.xpath('./p:txBody/a:lstStyle/a:lvl1pPr/a:defRPr')
        if dr:
            val=color_in(dr[0],'./a:solidFill')
            if val:return val
    except Exception:pass
    return TEXT_FALLBACK

def text_autofit(sh):
    body=sh._element.xpath('./p:txBody/a:bodyPr')
    if not body:return ('none',1.0,0.0)
    bb=body[0]
    nf=bb.xpath('./a:normAutofit')
    if nf:
        v=nf[0]
        scale=max(.1,min(1.0,int(v.get('fontScale','100000'))/100000))
        reduced=max(0.0,min(1.0,int(v.get('lnSpcReduction','0'))/100000))
        return ('shrink',scale,reduced)
    if bb.xpath('./a:spAutoFit'):return ('shape_grows',1.0,0.0)
    if bb.xpath('./a:noAutofit'):return ('no_autofit',1.0,0.0)
    return ('none',1.0,0.0)

def _paragraph_run_expressions(sh,pa,shrink=1.0):
    """Return shared Typst run records in DrawingML source order."""
    expressions=[]
    run_iter=iter(pa.runs)
    max_size=0.0
    first_style=None
    last_style=None
    for child in pa._p:
        tag=etree.QName(child).localname
        if tag=='br':
            # A soft break has no text run of its own. Carry the preceding
            # run's metric/style so it does not fall back to the containing
            # text frame's size (which differs for ordinary frames and
            # placeholders even when their paragraph runs are identical).
            if last_style is None:
                expressions.append('PptxPlaceholderRun(text:"\\n")')
            else:
                font,size,fill,weight,italic=last_style
                expressions.append('PptxPlaceholderRun('+', '.join((
                    'text:"\\n"',f'font:{stylize(font)}',f'size:{size:.3f}pt',
                    f'fill:rgb("{fill}")',f'weight:{stylize(weight)}',
                    f'italic:{str(bool(italic)).lower()}'))+')')
            continue
        if tag not in ('r','fld'):
            continue
        run=next(run_iter,None)
        if run is None or not run.text:
            continue
        font,size,fill,weight,italic=effective_run_style(sh,pa,run)
        last_style=(font,size,fill,weight,italic)
        max_size=max(max_size,size*shrink)
        if first_style is None:
            first_style=(font,size,fill,weight)
        fields=[f'text:{stylize(run.text.replace(chr(11), chr(10)))}',
                f'font:{stylize(font)}',f'size:{size:.3f}pt',
                f'fill:rgb("{fill}")',f'weight:{stylize(weight)}',
                f'italic:{str(bool(italic)).lower()}']
        expressions.append('PptxPlaceholderRun('+', '.join(fields)+')')
    return expressions,max_size,first_style


def paragraph_render_styles(sh, paragraphs):
    """Return source font size and family used for each paragraph line box.

    An empty DrawingML paragraph's ``endParaRPr`` describes the insertion
    point in PowerPoint, not a visible run. Use the nearest non-empty
    paragraph's font size for its line box instead of treating that caret
    style as visible content.
    """
    styles = []
    for paragraph in paragraphs:
        _, size, first_style = _paragraph_run_expressions(sh, paragraph)
        styles.append((size, first_style[0] if first_style else None))
    nonempty = [index for index, (paragraph, (size, _)) in enumerate(zip(paragraphs, styles))
                if size > 0 and paragraph.text.strip()]
    for index, (paragraph, (size, _)) in enumerate(zip(paragraphs, styles)):
        if size > 0 or paragraph.text.strip():
            continue
        nearest = min(nonempty, key=lambda other: abs(other - index), default=None)
        styles[index] = styles[nearest] if nearest is not None else (20.0, None)
    return styles


def paragraph_record_expression(sh,pa,*,shrink=1.0,spacing_reduction=0.0,bullet_counters=None,record_name='PptxParagraph',empty_paragraph_size=None,empty_paragraph_font=None):
    """Serialize one paragraph with source spacing semantics and resolved origin."""
    bullet_counters=bullet_counters if bullet_counters is not None else {}
    runs,max_size,first_style=_paragraph_run_expressions(sh,pa,shrink)
    empty_paragraph = not pa.text.strip() and not runs
    if max_size<=0:
        max_size=(empty_paragraph_size or 20.0)*shrink
    style=resolve_paragraph_style(paragraph_style_source_records(sh,pa),pa.level)
    spec=bullet_spec(sh,pa,max_size) if args.text_features in ('structure','all') and not empty_paragraph else None
    if spec:
        if spec['kind']=='number':
            key=(pa.level,spec['type'])
            number=bullet_counters.get(key,spec['start']-1)+1
            bullet_counters[key]=number
            marker=number_glyph(number,spec['type'])
        elif spec['kind']=='char':
            marker=spec['glyph']
        else:
            marker='•'
    else:
        marker=None
    record=style.effective_record()
    target_pitch=typst_target_pitch(style,max_size,spacing_reduction)
    before=typst_spacing(style.space_before,max_size)
    after=typst_spacing(style.space_after,max_size)
    fields=[f'level:{pa.level}']
    if not runs:
        fields.insert(0,f'text:{stylize(pa.text.replace(chr(11), chr(10)))}')
    if empty_paragraph:
        fields.append(f'empty-size:{(empty_paragraph_size or 20.0):.3f}pt')
        if empty_paragraph_font:
            fields.append(f'empty-font:{typst_font_fallback(empty_paragraph_font)}')
    if marker is not None:
        fields.append(f'bullet:{stylize(marker)}')
        bullet_font=(spec.get('font') if spec else None) or (first_style[0] if first_style else CURRENT_FONTS.get('minor','Arial'))
        bullet_fill=(spec.get('color') if spec else None) or (first_style[2] if first_style else shape_text_color(sh))
        bullet_weight=first_style[3] if first_style else 'regular'
        bullet_size=spec.get('fontsize',max_size) if spec else max_size
        if spec and spec.get('absolute_size'):bullet_size*=shrink
        fields.extend((f'bullet-font:{typst_font_fallback(bullet_font)}',
                       f'bullet-size:{bullet_size:.3f}pt',
                       f'bullet-fill:rgb("{bullet_fill}")',
                       f'bullet-weight:{stylize(bullet_weight)}'))
        if spec:
            for key,value in (('bullet-font-source',spec.get('font_source')),
                              ('bullet-size-kind',spec.get('size_kind')),
                              ('bullet-size-value',spec.get('size_value')),
                              ('bullet-size-source',spec.get('size_source')),
                              ('bullet-fill-source',spec.get('color_source'))):
                if value is not None:
                    fields.append(f'{key}:{stylize(str(value))}' if key.endswith('kind') or key.endswith('source')
                                  else f'{key}:{float(value):.4f}')
    bullet_style=record.get('bullet') or {}
    if bullet_style:
        fields.append(f'bullet-kind:{stylize(bullet_style.get("kind","char"))}')
        if bullet_style.get('type'):
            fields.append(f'bullet-type:{stylize(bullet_style["type"])}')
        if bullet_style.get('start') is not None:
            fields.append(f'bullet-start:{int(bullet_style["start"])}')
        if record.get('bullet_source'):
            fields.append(f'bullet-source:{stylize(record["bullet_source"])}')
    align=record.get('align')
    if align:fields.append(f'align:{stylize(align)}')
    for key in ('align_source','margin_left_source','indent_source'):
        value=record.get(key)
        if value:fields.append(f'{key.replace("_","-")}:{stylize(value)}')
    for key,value in (('line-kind',record['line_kind']),('line-value',record['line_value']),
                      ('line-source',record['line_source']),('target-pitch',target_pitch),
                      ('before-kind',record['before_kind']),('before-value',record['before_value']),
                      ('before-source',record['before_source']),('before',before),
                      ('after-kind',record['after_kind']),('after-value',record['after_value']),
                      ('after-source',record['after_source']),('after',after),
                      ('margin-left',record['margin_left']),('indent',record['indent'])):
        if value is None:continue
        if key.endswith('kind') or key.endswith('source'):
            fields.append(f'{key}:{stylize(str(value))}')
        elif key in ('line-value','before-value','after-value'):
            fields.append(f'{key}:{float(value):.4f}')
        else:
            fields.append(f'{key}:{float(value):.4f}pt')
    if runs:
        fields.append('runs:('+', '.join(runs)+( ',' if len(runs)==1 else '')+')')
    return record_name+'('+', '.join(fields)+')',first_style,max_size,style


def content_for_text(sh, ew, eh, sx, sy):
    """Build ordinary text with the same paragraph model used by placeholders."""
    autofit_mode,shrink,spacing_reduction=text_autofit(sh)
    paragraphs=[]
    first_style=None
    bullet_counters={}
    source_paragraphs = list(sh.text_frame.paragraphs)
    paragraph_styles = paragraph_render_styles(sh, source_paragraphs)
    for pa, (empty_size, empty_font) in zip(source_paragraphs, paragraph_styles):
        expression,run_style,_,_=paragraph_record_expression(
            sh,pa,shrink=shrink,spacing_reduction=spacing_reduction,bullet_counters=bullet_counters,
            empty_paragraph_size=empty_size,empty_paragraph_font=empty_font)
        paragraphs.append(expression)
        if first_style is None and run_style is not None:first_style=run_style
    if first_style is None:first_style=(CURRENT_FONTS.get('minor','Arial'),20*shrink,shape_text_color(sh),'regular')
    font,size,fill,weight=first_style
    body=sh._element.xpath('./p:txBody/a:bodyPr')
    attrs=body[0] if body else None
    anchor={'ctr':'center','b':'bottom'}.get(attrs.get('anchor','t') if attrs is not None else 't','top')
    if args.text_layout not in ('anchors','anchors-spacing'):
        anchor='top'
    margins=(('margin-left',round(float(sh.text_frame.margin_left or 0)/EMU,3)),
             ('margin-right',round(float(sh.text_frame.margin_right or 0)/EMU,3)),
             ('margin-top',round(float(sh.text_frame.margin_top or 0)/EMU,3)),
             ('margin-bottom',round(float(sh.text_frame.margin_bottom or 0)/EMU,3)))
    fields=[f'width:{ew:.3f}pt',f'height:{eh:.3f}pt',
            'paragraphs:('+', '.join(paragraphs)+( ',' if len(paragraphs)==1 else '')+')',
            f'font:{typst_font_fallback(font)}',f'size:{size:.3f}pt',
            f'fill:rgb("{fill}")',f'weight:{stylize(weight)}',f'anchor:{stylize(anchor)}',
            f'autofit-mode:{stylize(autofit_mode)}',f'font-scale:{shrink:.5f}',
            f'line-spacing-reduction:{spacing_reduction:.5f}']
    fields.extend(f'{key}:{value:.3f}pt' for key,value in margins)
    return 'PptxTextFrame('+',\n  '.join(fields)+')'


def table_style_color(scheme,transform=None):
    transform=transform or ''
    node=etree.fromstring(f'<a:schemeClr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="{scheme}">{transform}</a:schemeClr>'.encode())
    return resolve_color_node(node,'#FFFFFF')

def table_cell_typst(sh, row_idx, col_idx, cell, table_style, banded):
    """Build an editable Typst table cell with source text and direct formatting."""
    tcPr=cell._tc.tcPr or cell._tc.get_or_add_tcPr()
    row_count=len(sh.table.rows);col_count=len(sh.table.columns)
    fill_node=tcPr.xpath('./a:solidFill/*')
    fill=typst_color_expr(fill_node[0]) if fill_node else 'none'
    # A table style can be materialized in the package, or be the built-in
    # default style whose definition is intentionally omitted by PowerPoint.
    style_cell=None
    if table_style is not None:
        names=['wholeTbl']
        if row_idx==0 and sh.table.first_row:names.append('firstRow')
        elif banded and row_idx%2==1:names.append('band1H')
        if col_idx==0 and sh.table.first_col:names.append('firstCol')
        if col_idx==col_count-1 and sh.table.last_col:names.append('lastCol')
        if row_idx==row_count-1 and sh.table.last_row:names.append('lastRow')
        for name in names:
            zone=table_style.xpath('./*[local-name()="'+name+'"]/*[local-name()="tcStyle"]')
            if zone:style_cell=zone[0]
            if zone and not fill_node:
                sf=zone[0].xpath('./*[local-name()="fill"]/*[local-name()="solidFill"]/*')
                if sf:fill=typst_color_expr(sf[0])
                elif zone[0].xpath('./*[local-name()="fill"]/*[local-name()="noFill"]'):fill='none'
    elif not fill_node and table_style is None:
        # PowerPoint omits definitions of its built-in table styles. Approximate
        # the default Medium 2 / Accent 1 style used when its id is in effect.
        tbl=sh._element.xpath('.//a:tbl/a:tblPr')[0]
        style_id=tbl.xpath('./a:tableStyleId')
        style_id=[style_id[0].text] if style_id else []
        default_id=TABLE_STYLES_ROOT.get('def') if TABLE_STYLES_ROOT is not None else None
        if style_id and default_id and style_id[0]==default_id:
            if row_idx==0 and sh.table.first_row:fill='rgb("'+table_style_color('accent1')+'")'
            elif banded and row_idx%2==0:fill='rgb("'+table_style_color('accent1','<a:tint val="40000"/>')+'")'
            else:fill='rgb("'+table_style_color('accent1','<a:tint val="20000"/>')+'")'
    margin_x=(cell.margin_left+cell.margin_right)/(2*EMU)
    margin_y=(cell.margin_top+cell.margin_bottom)/(2*EMU)
    valign={1:'top',3:'center + horizon',4:'bottom'}.get(int(cell.vertical_anchor or 3),'center + horizon')
    align=valign
    paragraphs=[]
    for pa in cell.text_frame.paragraphs:
        ha='center' if pa.alignment==PP_ALIGN.CENTER else 'right' if pa.alignment==PP_ALIGN.RIGHT else 'left'
        align=ha+(' + horizon' if valign=='center + horizon' else '')
        bits=[]
        for run in pa.runs:
            if not run.text:continue
            font=font_mapping(run.font.name)
            size=run.font.size.pt if run.font.size else 14.0
            color=color_of(run) or TEXT_FALLBACK
            bold='bold' if run.font.bold else 'regular'
            italic=run.font.italic
            for j,line in enumerate(run.text.replace('\v','\n').split('\n')):
                if line:
                    bits.append(f'#text(font: {typst_font_fallback(font)}, size: {size:.2f}pt, weight: "{bold}", fill: rgb("{color}"), '+('style: "italic", ' if italic else '')+f'{stylize(line)})')
                if j<len(run.text.replace('\v','\n').split('\n'))-1:bits.append('#linebreak()')
        paragraphs.append(''.join(bits))
    body='#parbreak()'.join(paragraphs)
    stroke='none'
    # Respect direct cell edge lines. Missing edges use a light, thin gridline.
    edges=[]
    edge_names={'left':'lnL','right':'lnR','top':'lnT','bottom':'lnB'}
    for side,tag in edge_names.items():
        edge=tcPr.xpath('./a:'+tag)
        if not edge and style_cell is not None:
            edge=style_cell.xpath('./*[local-name()="tcBdr"]/*[local-name()="'+side+'"]')
        if edge:
            nofill=edge[0].xpath('./*[local-name()="noFill"] | ./*[local-name()="solidFill"]/*[local-name()="noFill"]')
            if nofill:edges.append(side+': none');continue
            col=edge[0].xpath('./*[local-name()="solidFill"]/*')
            color=resolve_color_node(col[0],'#FFFFFF') if col else '#FFFFFF'
            line=edge[0].xpath('./*[local-name()="ln"]')
            weight=(float(line[0].get('w','6350'))/EMU) if line else 0.5
            edges.append(side+f': {weight:.3f}pt + rgb("{color}")')
        else:edges.append(side+': 0.5pt + rgb("#FFFFFF")')
    stroke='('+', '.join(edges)+')'
    return f'table.cell(fill:{fill},stroke:{stroke},inset:(x:{margin_x:.3f}pt,y:{margin_y:.3f}pt),align:{align})[{body}]'

def normalize_table_row_heights(frame_height, row_heights):
    """Fit source row proportions to the PPTX table frame when OOXML is stale."""
    heights=list(row_heights)
    total=sum(heights)
    if frame_height>0 and total>0 and abs(frame_height-total)>1.0:
        return [height*frame_height/total for height in heights]
    return heights


def typst_table(sh):
    tbl=sh._element.xpath('.//a:tbl')[0]
    tblpr=tbl.xpath('./a:tblPr')[0]
    style_id=tblpr.xpath('./a:tableStyleId')
    style_id=[style_id[0].text] if style_id else []
    table_style=None
    if style_id and TABLE_STYLES_ROOT is not None:
        found=TABLE_STYLES_ROOT.xpath('./*[local-name()="tblStyle"][@styleId="'+style_id[0]+'"]')
        if found:table_style=found[0]
    widths=[c.width/EMU for c in sh.table.columns]
    heights=[r.height/EMU for r in sh.table.rows]
    frame_height=float(sh.height)/EMU
    heights=normalize_table_row_heights(frame_height,heights)
    column_expr='('+', '.join(f'{w:.3f}pt' for w in widths)+')'
    row_expr='('+', '.join(f'{h:.3f}pt' for h in heights)+')'
    banded=tblpr.get('bandRow') in ('1','true') or sh.table.horz_banding
    cells=[]
    for ri in range(len(sh.table.rows)):
        for ci in range(len(sh.table.columns)):
            cells.append(table_cell_typst(sh,ri,ci,sh.table.cell(ri,ci),table_style,banded))
    return (f'table(columns:{column_expr},rows:{row_expr},gutter:0pt,stroke:none,'+
            ',\n  '.join(cells)+')')

def table_has_spans(sh):
    """Merged cells need colspan/rowspan mapping and are not flattened silently."""
    cells=sh._element.xpath('.//a:tbl/a:tr/a:tc')
    return any(tc.xpath('./a:tcPr[@gridSpan or @rowSpan or @hMerge or @vMerge]') for tc in cells)


TEXT_FALLBACK='#111111'
registered={}; media_contents={}; media_uses=defaultdict(list); manifest=[]; counts=Counter(); vector_assets={}; compatible_media={}; media_warnings=[]
# Export assets on demand, reference by embedded original name and SHA256.
def register_media(part):
    name=Path(str(part.partname)).name
    data=part.blob
    sha=hashlib.sha256(data).hexdigest()
    if sha in registered:return registered[sha]
    ext=Path(name).suffix.lower()
    out_name=f'{Path(name).stem}_{sha[:10]}{ext}'
    dest=ASSETS/out_name;dest.write_bytes(data)
    rel='assets/media/'+out_name
    registered[sha]=rel
    media_contents[sha]={'path':rel,'source':str(part.partname),'bytes':len(data),'sha256':sha,'placements':[]}
    return rel

def record_compatibility_use(part,original_path,sha,render_path,context):
    ext=Path(str(part.partname)).suffix.lower()
    if ext not in ('.emf','.wmf'):
        return
    derived=next((item for item in media_contents[sha].get('derived_assets',[])
                  if item['path']==render_path),{})
    output_format=derived.get('format','compatible image')
    warning={
        'issue':'legacy_vector_media_compatibility_fallback',
        'source_media':original_path,'source_part':str(part.partname),
        'source_sha256':sha,
        'fallback':f'Used a {output_format.upper()} render derivative in Typst; the original {ext[1:].upper()} is retained in assets.',
        'expected_loss':('Raster output loses vector scalability; the original vector media remains available in the project assets.'
                         if output_format=='png' else
                         'The imported artwork remains one placed SVG image rather than native Typst shapes.'),
    }
    warning.update(context or {})
    media_warnings.append(warning)
    counts['media_raster_fallback_placements' if output_format=='png'
           else 'media_vector_fallback_placements']+=1

def typst_compatible_media(part, original_path, context=None):
    """Return a Typst-readable media path and optional actionable warning."""
    ext=Path(str(part.partname)).suffix.lower()
    sha=hashlib.sha256(part.blob).hexdigest()
    if ext not in ('.tif','.tiff','.emf','.wmf'):
        return original_path, None
    conversion_key=(sha,ext,300,4096)
    if conversion_key in compatible_media:
        rel,conversion_error=compatible_media[conversion_key]
    elif ext in ('.tif','.tiff'):
        try:
            from PIL import Image
            import io
            name=f'{Path(original_path).stem}_typst.png'
            (ASSETS/name).unlink(missing_ok=True)
            with Image.open(io.BytesIO(part.blob)) as source:
                source.convert('RGBA').save(ASSETS/name,optimize=True)
            rel='assets/media/'+name
            compatible_media[conversion_key]=(rel,None)
            media_contents[sha]['render_path']=rel
            media_contents[sha].setdefault('derived_assets',[]).append({
                'path':rel,'format':'png','converter':'Pillow','lossy':False,
                'reason':'Typst does not decode TIFF directly.'})
            counts['media_compatibility_derivatives']+=1
            return rel, None
        except Exception as exc:
            rel=None
            conversion_error=f'TIFF to PNG conversion failed: {exc}'
            compatible_media[conversion_key]=(rel,conversion_error)
    else:
        destination_stem=ASSETS/f'{Path(original_path).stem}_typst'
        converted,conversion_error=convert_metafile(
            ASSETS/Path(original_path).name,destination_stem,dpi=300,max_edge=4096)
        if converted:
            rel='assets/media/'+converted.path.name
            media_contents[sha]['render_path']=rel
            record={'path':rel,'format':converted.output_format,'converter':converted.converter,
                    'lossy':converted.output_format=='png','notes':converted.notes,
                    'source_sha256':sha}
            if converted.dpi is not None:
                record.update(dpi=converted.dpi,max_edge_px=converted.max_edge)
            media_contents[sha].setdefault('derived_assets',[]).append(record)
            compatible_media[conversion_key]=(rel,None)
            counts['media_compatibility_derivatives']+=1
            record_compatibility_use(part,original_path,sha,rel,context)
            return rel,None
        compatible_media[conversion_key]=(None,conversion_error)
    if rel:
        record_compatibility_use(part,original_path,sha,rel,context)
        return rel,None
    warning={
        'issue':'unsupported_media_format','source_media':original_path,
        'source_part':str(part.partname),'source_sha256':sha,
        'fallback':'omitted from Typst rendering; original media retained in assets and source PPTX',
        'expected_loss':'This object will not be visible in the Typst output.',
        'reason':conversion_error,
    }
    warning.update(context or {})
    media_warnings.append(warning)
    counts['media_conversion_failures']+=1
    return None,warning

def trans(x,y):return np.array([[1,0,x],[0,1,y],[0,0,1]],float)
def scale(x,y):return np.array([[x,0,0],[0,y,0],[0,0,1]],float)
def rot(a):
    t=math.radians(a);c=math.cos(t);s=math.sin(t)
    return np.array([[c,-s,0],[s,c,0],[0,0,1]],float)

def shape_matrix(sh):
    e=sh._element
    if etree.QName(e).localname=='graphicFrame':
        xfrm=e.xpath('./p:xfrm')
    elif sh.shape_type==MSO_SHAPE_TYPE.PICTURE:
        xfrm=e.xpath('./p:spPr/a:xfrm')
    else:xfrm=e.xpath('./p:spPr/a:xfrm')
    if xfrm:
        xf=xfrm[0]; flip_h=xf.get('flipH')=='1' or xf.get('flipH')=='true';flip_v=xf.get('flipV')=='1' or xf.get('flipV')=='true'; a=int(xf.get('rot','0'))/60000
    else: flip_h=flip_v=False;a=0
    w=sh.width/EMU;h=sh.height/EMU;x=sh.left/EMU;y=sh.top/EMU
    return trans(x+w/2,y+h/2) @ rot(a) @ scale(-1 if flip_h else 1,-1 if flip_v else 1) @ scale(w,h) @ trans(-0.5,-0.5)

def group_child_matrix(sh):
    xf=sh._element.grpSpPr.xfrm
    x,y=xf.off.x/EMU,xf.off.y/EMU
    w,h=xf.ext.cx/EMU,xf.ext.cy/EMU
    chx,chy=xf.chOff.x/EMU,xf.chOff.y/EMU
    chw,chh=xf.chExt.cx/EMU,xf.chExt.cy/EMU
    a=int(xf.get('rot','0'))/60000
    flip_h=xf.get('flipH') in ('1','true');flip_v=xf.get('flipV') in ('1','true')
    return trans(x+w/2,y+h/2) @ rot(a) @ scale(-1 if flip_h else 1,-1 if flip_v else 1) @ trans(-w/2,-h/2) @ scale(w/chw if chw else 1,h/chh if chh else 1) @ trans(-chx,-chy)

def image_fill_params(sh):
    src=sh._element.xpath('./p:spPr/a:blipFill/a:srcRect') or sh._element.xpath('./p:blipFill/a:srcRect')
    stretch=sh._element.xpath('./p:spPr/a:blipFill/a:stretch/a:fillRect') or sh._element.xpath('./p:blipFill/a:stretch/a:fillRect')
    blip=sh._element.xpath('./p:spPr/a:blipFill/a:blip') or sh._element.xpath('./p:blipFill/a:blip')
    def vals(el):return {k:float(el.get(k,'0'))/100000 for k in ['l','t','r','b']}
    sr=vals(src[0]) if src else None
    fr=vals(stretch[0]) if stretch else None
    am=blip[0].xpath('./a:alphaModFix') if blip else []
    alpha=float(am[0].get('amt','100000'))/100000 if am else 1.0
    if sr and not any(sr.values()):sr=None
    if fr and not any(fr.values()):fr=None
    return sr,fr,alpha

ALPHA_CACHE={}
def alpha_variant(part, raw_path, alpha):
    if alpha>=.999:return raw_path
    from PIL import Image
    import io
    sha=hashlib.sha256(part.blob).hexdigest()
    key=(sha,round(alpha,5))
    if key in ALPHA_CACHE:return ALPHA_CACHE[key]
    try:
        im=Image.open(io.BytesIO(part.blob)).convert('RGBA')
        a=im.getchannel('A').point(lambda x:round(x*alpha))
        im.putalpha(a)
        name=f'{Path(raw_path).stem}_alpha{round(alpha*10000)}.png'
        im.save(ASSETS/name,optimize=True)
        rel='assets/media/'+name
        ALPHA_CACHE[key]=rel
        return rel
    except Exception:return raw_path

def image_blip(sh):
    bl=sh._element.xpath('./p:spPr/a:blipFill/a:blip') or sh._element.xpath('./p:blipFill/a:blip')
    if not bl:return None
    from pptx.oxml.ns import qn
    # PowerPoint can store SVG artwork in an Office extension relationship and
    # a raster fallback in a:blip. Prefer the SVG for crisp master/layout art.
    svg_blips=bl[0].xpath('.//*[local-name()="svgBlip"]')
    rids=([svg_blips[0].get(qn('r:embed'))] if svg_blips else [])+[bl[0].get(qn('r:embed'))]
    for rid in rids:
        if not rid:continue
        try:return sh.part.related_part(rid)
        except KeyError:continue
    return None

def effective_shape_geometry(sh):
    """Resolve geometry inherited by a slide placeholder from its layout peer."""
    own=sh._element.xpath('./p:spPr/a:prstGeom | ./p:spPr/a:custGeom')
    if own or not sh.is_placeholder or not hasattr(sh.part,'slide_layout'):
        return sh._element
    ph=sh.placeholder_format
    for candidate in sh.part.slide_layout.placeholders:
        other=candidate.placeholder_format
        if other.idx==ph.idx and other.type==ph.type:
            inherited=candidate._element.xpath('./p:spPr/a:prstGeom | ./p:spPr/a:custGeom')
            if inherited:return candidate._element
    return sh._element

def geometry_kind(sh, element=None):
    # p:pic can carry nonrectangular prstGeom/custGeom as well as p:sp.
    # Never infer 'rect' solely because its OOXML element is a picture.
    if element is None:element=sh._element
    if element.xpath('./p:spPr/a:prstGeom'):
        return element.xpath('./p:spPr/a:prstGeom')[0].get('prst')
    path=element.xpath('./p:spPr/a:custGeom/a:pathLst/a:path')
    if not path:return 'rect'
    pts=[]
    for tag in path[0]:
        t=tag.tag.split('}')[-1]
        if t not in ('moveTo','lnTo','close'):return 'complex'
        if t=='close':continue
        coords=tag.xpath('./a:pt')
        if len(coords)!=1:return 'complex'
        pts.append((float(coords[0].get('x')),float(coords[0].get('y'))))
    # nearly every Canva shape is a rectangle written with 5 points, sometimes corner order reversed
    if len(pts) in (4,5):
        if pts[0]==pts[-1]:pts=pts[:-1]
        xs=set(x for x,y in pts);ys=set(y for x,y in pts)
        if len(xs)==2 and len(ys)==2:return 'rect'
    return 'polygon'

def placement_info(sh, parent):
    M=parent @ shape_matrix(sh)
    ctr=M @ np.array([0.5,0.5,1.])
    vx=M[:2,:2] @ np.array([1.,0.]);vy=M[:2,:2] @ np.array([0.,1.])
    w=float(np.linalg.norm(vx));h=float(np.linalg.norm(vy))
    a=math.degrees(math.atan2(vx[1],vx[0]))
    det=float(np.linalg.det(M[:2,:2]));flip_y=det<0
    skew=float(np.dot(vx,vy)/(w*h)) if w*h>0 else 0.
    return dict(x=float(ctr[0]-w/2),y=float(ctr[1]-h/2),w=w,h=h,rotation=a,flip_y=flip_y,skew=skew)


TRANSPARENCY_CACHE={}
def original_asset_has_alpha(path):
    """Keep established clipping for alpha-backed picture overlays.
    The source raster is never changed or regenerated here.
    """
    if path not in TRANSPARENCY_CACHE:
        try:
            from PIL import Image
            with Image.open(OUT/path) as im:
                a=im.getchannel('A') if 'A' in im.getbands() else None
                TRANSPARENCY_CACHE[path]=bool(a and a.getextrema()[0] < 255)
        except Exception:
            TRANSPARENCY_CACHE[path]=False
    return TRANSPARENCY_CACHE[path]


def shape_fill_element(shape):
    """Resolve a shape's direct fill or its theme fill-style reference."""
    direct = shape._element.xpath(
        './p:spPr/a:solidFill | ./p:spPr/a:gradFill | ./p:spPr/a:blipFill | ./p:spPr/a:noFill'
    )
    if direct:
        return direct[0]
    fill_ref = shape._element.xpath('./p:style/a:fillRef')
    if not fill_ref:
        return None
    try:
        index = int(fill_ref[0].get('idx', '0')) - 1
    except (TypeError, ValueError):
        return None
    if index < 0 or CURRENT_SLIDE is None:
        return None
    theme_part = owner_theme_part(CURRENT_SLIDE)
    if theme_part is None:
        return None
    theme = etree.fromstring(theme_part.blob)
    styles = theme.xpath('./a:themeElements/a:fmtScheme/a:fillStyleLst/*', namespaces={
        'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'})
    if index >= len(styles):
        return None
    resolved = deepcopy(styles[index])
    ref_colors = fill_ref[0].xpath('./a:srgbClr | ./a:schemeClr | ./a:sysClr | ./a:prstClr', namespaces={
        'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'})
    if ref_colors:
        for placeholder in resolved.xpath('.//a:schemeClr[@val="phClr"]', namespaces={
                'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}):
            replacement = deepcopy(ref_colors[0])
            for transform in placeholder:
                replacement.append(deepcopy(transform))
            placeholder.getparent().replace(placeholder, replacement)
    return resolved


def shape_background_expr(shape):
    """Convert a resolved simple placeholder fill to a Typst paint expression."""
    fill = shape_fill_element(shape)
    if fill is None or etree.QName(fill).localname == 'noFill':
        return None
    if etree.QName(fill).localname == 'solidFill':
        color = fill.xpath('./*[local-name()="srgbClr" or local-name()="schemeClr" or local-name()="sysClr" or local-name()="prstClr"]')
        return typst_color_expr(color[0]) if color else None
    if etree.QName(fill).localname == 'gradFill':
        return gradient_expr(fill, CURRENT_MAP)
    return None

def typst_img(path, placement, source_crop=None, fillrect=None):
    w,h=placement['w'],placement['h']
    no_op = (source_crop is None and fillrect is None)
    if no_op:
        core=f'image("{path}",width:{w:.3f}pt,height:{h:.3f}pt,fit:"stretch")'
    else:
        # PPTX a:stretch/a:fillRect uses negative offsets to extend a picture
        # beyond the shape boundary; this matters for large zoomed photos.
        fr=fillrect or {k:0 for k in ['l','t','r','b']}
        cr=source_crop or {k:0 for k in ['l','t','r','b']}
        l,t,r,b=(fr[k] for k in ['l','t','r','b'])
        cl,ct,cright,cb=(cr[k] for k in ['l','t','r','b'])
        DW=(1-l-r)*w;DH=(1-t-b)*h
        if DW>0 and DH>0 and cl+cright<.98 and ct+cb<.98:
            IW=DW/(1-cl-cright);IH=DH/(1-ct-cb)
            ix=l*w-cl*IW;iy=t*h-ct*IH
            if args.clip_mode == 'component' and not original_asset_has_alpha(path):
                core=(f'ClipContent(width:{w:.3f}pt,height:{h:.3f}pt,children:['
                      f'#place(top+left,dx:{ix:.3f}pt,dy:{iy:.3f}pt)['
                      f'#image("{path}",width:{IW:.3f}pt,height:{IH:.3f}pt,fit:"stretch")]])')
            else:
                core=(f'box(width:{w:.3f}pt,height:{h:.3f}pt,clip:true)['
                      f'#place(top+left,dx:{ix:.3f}pt,dy:{iy:.3f}pt)['
                      f'#image("{path}",width:{IW:.3f}pt,height:{IH:.3f}pt,fit:"stretch")]]')
        else:core=f'image("{path}",width:{w:.3f}pt,height:{h:.3f}pt,fit:"stretch")'
    if placement['flip_y']:
        core=f'scale(x:100%,y:-100%,origin:center+horizon,reflow:false,{core})'
    return core

def vector_svg(sh):
    """Serialize source OOXML freeform geometry/strokes as shared vector artwork.
    Image-filled shapes never enter this function: they retain original media bytes.
    """
    e=sh._element
    ln=e.xpath('./p:spPr/a:ln')
    fill_element=shape_fill_element(sh)
    solid=(fill_element.xpath('./*[local-name()="srgbClr" or local-name()="schemeClr" or local-name()="sysClr" or local-name()="prstClr"]')
           if fill_element is not None and etree.QName(fill_element).localname=='solidFill' else [])
    if not (ln or solid or (fill_element is not None and etree.QName(fill_element).localname=='gradFill')):return None
    w0=max(.1,sh.width/EMU);h0=max(.1,sh.height/EMU)
    width=float(w0);height=float(h0)
    rgb='none'
    if solid:rgb=resolve_color_node(solid[0],'none')
    fill_opacity=color_opacity(solid[0]) if solid else 1.0
    native_grad=[fill_element] if fill_element is not None and etree.QName(fill_element).localname=='gradFill' else []
    defs=''
    if native_grad:
        gf=native_grad[0]; stops=[]
        for gs in xpath_ns(gf,'./a:gsLst/a:gs'):
            color=(color_in(gs,'./a:srgbClr') or color_in(gs,'./a:schemeClr') or
                   color_in(gs,'./a:sysClr') or '#000000')
            alpha_node=xpath_ns(gs,'./a:srgbClr/a:alpha | ./a:schemeClr/a:alpha | ./a:sysClr/a:alpha')
            opacity=float(alpha_node[0].get('val','100000'))/100000 if alpha_node else 1
            offset=float(gs.get('pos','0'))/1000
            stops.append(f'<stop offset="{offset:.2f}%" stop-color="{color}" stop-opacity="{opacity:.5f}"/>')
        linear=xpath_ns(gf,'./a:lin')
        if linear:
            # In OOXML, angle is clockwise relative to a left-to-right vector.
            a=math.radians(float(linear[0].get('ang','0'))/60000)
            x1=50-50*math.cos(a);y1=50-50*math.sin(a)
            x2=50+50*math.cos(a);y2=50+50*math.sin(a)
            defs=f'<defs><linearGradient id="pptgrad" x1="{x1:.4f}%" y1="{y1:.4f}%" x2="{x2:.4f}%" y2="{y2:.4f}%">{"".join(stops)}</linearGradient></defs>'
        else:
            defs=f'<defs><radialGradient id="pptgrad">{"".join(stops)}</radialGradient></defs>'
        if stops:rgb='url(#pptgrad)'
    stroke='none';stroke_w=0.;linecap='butt';dash='';stroke_opacity=1.0
    if ln:
        lin=ln[0]
        stroke=color_in(lin,'./a:solidFill','none')
        stroke_opacity=color_opacity_in(lin,'./a:solidFill')
        if lin.xpath('./a:noFill'):stroke='none'
        stroke_w=float(lin.get('w','0'))/EMU
        linecap='round' if lin.get('cap')=='rnd' else 'square' if lin.get('cap')=='sq' else 'butt'
        pd=lin.xpath('./a:prstDash')
        if pd and pd[0].get('val') in ('sysDot','dot'):
            dash=f' stroke-dasharray="0 {stroke_w*3:.3f}"'
    paths=e.xpath('./p:spPr/a:custGeom/a:pathLst/a:path')
    geom=e.xpath('./p:spPr/a:prstGeom')
    if paths and any(op.tag.split('}')[-1]=='arcTo' for path in paths for op in path):
        _,normpaths=shape_paths(e)
        art=''.join(f'<path d="{path}" transform="scale({w0:.6f},{h0:.6f})" fill="{rgb}" stroke="{stroke}" stroke-width="{stroke_w:.4f}" vector-effect="non-scaling-stroke"/>' for path in normpaths)
    elif paths:
        path=paths[0];pw=max(1.,float(path.get('w',max(1,sh.width))));ph=max(1.,float(path.get('h',max(1,sh.height))))
        cmds=[]
        for op in path:
            typ=op.tag.split('}')[-1]
            if typ in ('moveTo','lnTo'):
                pt=op.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}pt')
                if pt:
                    x=float(pt[0].get('x'))/pw*w0;y=float(pt[0].get('y'))/ph*h0
                    cmds.append(f'{"M" if typ=="moveTo" else "L"}{x:.3f},{y:.3f}')
            elif typ=='cubicBezTo':
                pp=op.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}pt')
                if len(pp)==3:
                    coords=['%.3f,%.3f'%(float(pt.get('x'))/pw*w0,float(pt.get('y'))/ph*h0) for pt in pp]
                    cmds.append('C'+' '.join(coords))
            elif typ=='quadBezTo':
                pp=op.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}pt')
                if len(pp)==2:
                    coords=['%.3f,%.3f'%(float(pt.get('x'))/pw*w0,float(pt.get('y'))/ph*h0) for pt in pp]
                    cmds.append('Q'+' '.join(coords))
            elif typ=='close':cmds.append('Z')
        d=' '.join(cmds)
        if not d:return None
        art=f'<path d="{d}" fill="{rgb}" stroke="{stroke}" stroke-width="{stroke_w:.4f}" stroke-linecap="{linecap}"{dash}/>'
    elif geom:
        g=geom[0].get('prst')
        if g=='ellipse':art=f'<ellipse cx="{w0/2}" cy="{h0/2}" rx="{w0/2}" ry="{h0/2}" fill="{rgb}" stroke="{stroke}" stroke-width="{stroke_w}"/>'
        elif g in ('rect','roundRect'):
            r= min(w0,h0)*.1 if g=='roundRect' else 0
            art=f'<rect x="0" y="0" width="{w0}" height="{h0}" rx="{r}" fill="{rgb}" stroke="{stroke}" stroke-width="{stroke_w}"/>'
        elif g=='line':art=f'<path d="M0,0 L{w0:.3f},{h0:.3f}" stroke="{stroke}" stroke-width="{stroke_w}"/>'
        else:
            _,normpaths=shape_paths(e)
            if not normpaths:return None
            art=''.join(f'<path d="{path}" transform="scale({w0:.6f},{h0:.6f})" fill="{rgb}" stroke="{stroke}" stroke-width="{stroke_w:.4f}" vector-effect="non-scaling-stroke"/>' for path in normpaths)
    elif abs(sh.height/EMU)<.01 and sh.width>0:
        art=f'<path d="M0,0 L{w0:.3f},0" stroke="{stroke}" stroke-width="{stroke_w}"/>'
    else:return None
    # Arrow markers only for source shapes that actually have arrowheads.
    if ln and stroke not in ('none',None):
        heads=[]
        for tag,attach,id in [('headEnd','start','arrstart'),('tailEnd','end','arrend')]:
            ends=ln[0].xpath('./a:'+tag)
            if not ends:continue
            typ=ends[0].get('type','none')
            if typ not in ('triangle','stealth','diamond','oval','arrowhead'):continue
            if typ=='oval':mark=f'<ellipse cx="5" cy="3" rx="3" ry="3" fill="{stroke}"/>'
            elif typ=='diamond':mark=f'<path d="M0,3 L4,0 L8,3 L4,6 Z" fill="{stroke}"/>'
            else:mark=f'<path d="M0,0 L9,3 L0,6 Z" fill="{stroke}"/>'
            marker_opacity=f' opacity="{stroke_opacity:.5f}"' if stroke_opacity<.99999 else ''
            heads.append(f'<marker id="{id}" markerWidth="9" markerHeight="6" refX="{0 if attach=="start" else 9}" refY="3" orient="auto-start-reverse" markerUnits="strokeWidth"{marker_opacity}>{mark}</marker>')
            art=art.replace('<path ', f'<path marker-{attach}="url(#{id})" ', 1)
        if heads:defs=defs.replace('</defs>',''.join(heads)+'</defs>') if defs else '<defs>'+''.join(heads)+'</defs>'
    # permit stroke beyond boundaries, as in PPTX ring decorations
    pad=max(0,stroke_w/2)
    opacity_attrs=[]
    if fill_opacity<.99999:opacity_attrs.append(f'fill-opacity="{fill_opacity:.5f}"')
    if stroke_opacity<.99999:opacity_attrs.append(f'stroke-opacity="{stroke_opacity:.5f}"')
    if opacity_attrs:art=f'<g {" ".join(opacity_attrs)}>{art}</g>'
    svg=f'<svg xmlns="http://www.w3.org/2000/svg" width="{w0+2*pad:.3f}" height="{h0+2*pad:.3f}" viewBox="{-pad:.3f} {-pad:.3f} {w0+2*pad:.3f} {h0+2*pad:.3f}">{defs}{art}</svg>'
    sha=hashlib.sha256(svg.encode()).hexdigest()
    if sha not in vector_assets:
        out_name=f'shape_{sha[:12]}.svg';(ASSETS/out_name).write_text(svg)
        vector_assets[sha]='assets/media/'+out_name
    return vector_assets[sha],pad,w0,h0

def smartart_drawing(sh):
    """Return the materialized DrawingML drawing for a SmartArt graphic frame."""
    graphic=sh._element.xpath('./a:graphic/a:graphicData')
    if not graphic or not any(etree.QName(child).localname=='relIds' for child in graphic[0]):
        return None
    if not any(etree.QName(child).namespace.endswith('/diagram') for child in graphic[0]):
        return None
    for rel in sh.part.rels.values():
        if rel.reltype.endswith('/diagramDrawing') and not rel.is_external:
            return etree.fromstring(rel.target_part.blob)
    return None

def smartart_svg(sh, drawing):
    """Serialize SmartArt's positioned drawing cache to a reusable vector SVG.

    PowerPoint stores the calculated layout (including its z-order) in drawing*.xml.
    Reusing that drawing cache is substantially more faithful than rebuilding its
    data model and layout algorithm. Text remains vector text in the SVG.
    """
    width=max(1,float(sh.width));height=max(1,float(sh.height))
    ns={'dsp':'http://schemas.microsoft.com/office/drawing/2008/diagram',
        'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    root=drawing.xpath('.//dsp:spTree',namespaces=ns)
    if not root:return None
    out=[]
    def shape_color(parent,default='none'):
        fill=parent.xpath('./*[local-name()="solidFill"]/*')
        return resolve_color_node(fill[0],default) if fill else default
    def color_opacity_path(parent):
        fill=parent.xpath('./*[local-name()="solidFill"]/*')
        return color_opacity(fill[0]) if fill else 1.0
    def run_text(run):
        return ''.join(run.xpath('.//*[local-name()="t"]/text()'))
    for shape_index,sp in enumerate(root[0]):
        if etree.QName(sp).localname!='sp':continue
        xfs=sp.xpath('./dsp:spPr/a:xfrm',namespaces=ns)
        if not xfs:continue
        xf=xfs[0];offs=xf.xpath('./a:off',namespaces=ns);exts=xf.xpath('./a:ext',namespaces=ns)
        if not offs or not exts:continue
        x=float(offs[0].get('x','0'));y=float(offs[0].get('y','0'))
        w=float(exts[0].get('cx','0'));h=float(exts[0].get('cy','0'))
        if w<0 or h<0 or (w==0 and h==0):continue
        rot=int(xf.get('rot','0'))/60000
        flipx=-1 if xf.get('flipH') in ('1','true') else 1
        flipy=-1 if xf.get('flipV') in ('1','true') else 1
        transform=f'transform="translate({x+w/2:.3f} {y+h/2:.3f}) rotate({rot:.4f}) scale({flipx} {flipy}) translate({-w/2:.3f} {-h/2:.3f})"'
        geom=sp.xpath('./dsp:spPr/a:prstGeom',namespaces=ns)
        preset=geom[0].get('prst','rect') if geom else 'rect'
        solid=sp.xpath('./dsp:spPr/a:solidFill/*',namespaces=ns)
        fill=resolve_color_node(solid[0],'none') if solid else 'none'
        fill_op=color_opacity(solid[0]) if solid else 1
        ln=sp.xpath('./dsp:spPr/a:ln',namespaces=ns)
        stroke='none';sw=0.;stroke_op=1.
        if ln:
            stroke=shape_color(ln[0])
            stroke_op=color_opacity_path(ln[0])
            if ln[0].xpath('./a:noFill',namespaces=ns):stroke='none'
            sw=float(ln[0].get('w','0'))
        dash=''
        dash_nodes=ln[0].xpath('./a:prstDash/@val',namespaces=ns) if ln else []
        if dash_nodes and dash_nodes[0] not in ('solid','sys','solidFill'):
            ratios={'dash':(4,2),'lgDash':(8,3),'dashDot':(4,2,1,2),
                    'lgDashDot':(8,3,1,3),'lgDashDotDot':(8,3,1,3,1,3),
                    'sysDash':(3,2),'sysDot':(1,2),'sysDashDot':(3,2,1,2),
                    'sysDashDotDot':(3,2,1,2,1,2),'dot':(1,2)}
            ratio=ratios.get(dash_nodes[0])
            if ratio:dash=f' stroke-dasharray="{" ".join(f"{sw*v:.3f}" for v in ratio)}"'
        if preset=='ellipse':
            artwork=f'<ellipse cx="{w/2:.3f}" cy="{h/2:.3f}" rx="{w/2:.3f}" ry="{h/2:.3f}" fill="{fill}" stroke="{stroke}" stroke-width="{sw:.3f}"/>'
        elif preset in ('line','straightConnector1'):
            marker_defs=[];marker_attrs=[]
            for tag,side,ident in (('headEnd','start',f'smartart-start-{shape_index}'),('tailEnd','end',f'smartart-end-{shape_index}')):
                end=ln[0].xpath('./a:'+tag,namespaces=ns) if ln else []
                typ=end[0].get('type','none') if end else 'none'
                if typ not in ('triangle','stealth','diamond','oval','arrowhead'):continue
                if typ=='diamond':mark='<path d="M0,3 L4,0 L8,3 L4,6 Z"/>'
                elif typ=='oval':mark='<ellipse cx="4" cy="3" rx="3" ry="3"/>'
                else:mark='<path d="M0,0 L9,3 L0,6 Z"/>'
                marker_defs.append(f'<defs><marker id="{ident}" markerWidth="9" markerHeight="6" refX="{0 if side=="start" else 9}" refY="3" orient="auto-start-reverse" markerUnits="strokeWidth" fill="{stroke}">{mark}</marker></defs>')
                marker_attrs.append(f'marker-{side}="url(#{ident})"')
            arrow=''.join(marker_defs)
            marks=' '.join(marker_attrs)
            artwork=f'{arrow}<path d="M0,0 L{w:.3f},{h:.3f}" fill="none" stroke="{stroke}" stroke-width="{max(sw,1):.3f}"{dash} {marks}/>'
        elif preset in ('rect','roundRect'):
            r=min(w,h)*.08 if preset=='roundRect' else 0
            artwork=f'<rect x="0" y="0" width="{w:.3f}" height="{h:.3f}" rx="{r:.3f}" fill="{fill}" stroke="{stroke}" stroke-width="{sw:.3f}"/>'
        else:
            path=preset_path(preset)
            if path:artwork=f'<path d="{path}" transform="scale({w:.6f} {h:.6f})" fill="{fill}" stroke="{stroke}" stroke-width="{sw:.3f}"/>'
            else:artwork=f'<rect x="0" y="0" width="{w:.3f}" height="{h:.3f}" fill="{fill}" stroke="{stroke}" stroke-width="{sw:.3f}"/>'
        op=[]
        if fill_op<.99999:op.append(f'fill-opacity="{fill_op:.5f}"')
        if stroke_op<.99999:op.append(f'stroke-opacity="{stroke_op:.5f}"')
        out.append(f'<g {transform} {" ".join(op)}>{artwork}')
        bodies=sp.xpath('./dsp:txBody',namespaces=ns)
        if bodies:
            body=bodies[0];bodypr=body.xpath('./a:bodyPr',namespaces=ns)
            bp=bodypr[0] if bodypr else None
            left=float(bp.get('lIns','91440')) if bp is not None else 91440
            right=float(bp.get('rIns','91440')) if bp is not None else 91440
            top=float(bp.get('tIns','45720')) if bp is not None else 45720
            anchor=bp.get('anchor','ctr') if bp is not None else 'ctr'
            paras=body.xpath('./a:p',namespaces=ns)
            font_default=1800
            total=0
            parsed=[]
            for para in paras:
                runs=para.xpath('./a:r | ./a:fld | ./a:endParaRPr',namespaces=ns)
                pieces=[];psize=None
                for run in runs:
                    if etree.QName(run).localname=='endParaRPr':continue
                    txt=run_text(run)
                    if not txt:continue
                    rpr=run.xpath('./a:rPr',namespaces=ns)
                    rp=rpr[0] if rpr else None
                    sz=int(rp.get('sz'))*127 if rp is not None and rp.get('sz') else font_default
                    psize=psize or sz
                    typeface=rp.xpath('./a:latin/@typeface',namespaces=ns) if rp is not None else []
                    fam=typeface[0] if typeface and not typeface[0].startswith('+') else 'Arial'
                    color_nodes=rp.xpath('./a:solidFill/*',namespaces=ns) if rp is not None else []
                    tc=resolve_color_node(color_nodes[0],'#111111') if color_nodes else '#111111'
                    pieces.append(f'<tspan font-family="{xml_escape(fam)}" font-size="{sz:.1f}" fill="{tc}">{xml_escape(txt)}</tspan>')
                if pieces:
                    sz=psize or font_default;total+=sz*1.2;parsed.append((para,pieces,sz))
            ytext=top if anchor=='t' else max(top,(h-total)/2) if anchor=='ctr' else max(top,h-total)
            for para,pieces,sz in parsed:
                ppr=para.xpath('./a:pPr',namespaces=ns)
                align=ppr[0].get('algn','l') if ppr else 'l'
                anch={'ctr':'middle','r':'end','just':'start','l':'start'}.get(align,'start')
                tx=left if anch=='start' else w-right if anch=='end' else w/2
                ytext+=sz
                out.append(f'<text x="{tx:.3f}" y="{ytext:.3f}" text-anchor="{anch}">{"".join(pieces)}</text>')
                ytext+=sz*.2
        out.append('</g>')
    if not out:return None
    svg=f'<svg xmlns="http://www.w3.org/2000/svg" width="{width/EMU:.3f}pt" height="{height/EMU:.3f}pt" viewBox="0 0 {width:.3f} {height:.3f}">{"".join(out)}</svg>'
    sha=hashlib.sha256(svg.encode()).hexdigest()
    if sha not in vector_assets:
        out_name=f'smartart_{sha[:12]}.svg';(ASSETS/out_name).write_text(svg,encoding='utf8')
        vector_assets[sha]='assets/media/'+out_name
    return vector_assets[sha]

def slide_background(slide):
    bg=slide._element.xpath('./p:cSld/p:bg/p:bgPr')
    if not bg:return 'white'
    solid=bg[0].xpath('./a:solidFill/a:srgbClr')
    if solid:return 'rgb("#'+solid[0].get('val')+'")'
    grad=bg[0].xpath('./a:gradFill')
    if grad:
        stops=[]
        for gs in grad[0].xpath('./a:gsLst/a:gs'):
            col=gs.xpath('./a:srgbClr')
            if col:
                color='#'+col[0].get('val')
                opacity=col[0].xpath('./a:alpha')
                if opacity:
                    from math import floor
                    op=int(opacity[0].get('val'))/100000
                    if op<1:color=f'#{color[1:]}{round(op*255):02X}'
                stops.append(f'rgb("{color}")')
        if stops:
            if grad[0].xpath('./a:path'):
                return 'gradient.radial('+','.join(stops)+')'
            lin=grad[0].xpath('./a:lin')
            angle=float(lin[0].get('ang','0'))/60000 if lin else 0
            return 'gradient.linear('+','.join(stops)+f',angle:{angle:.2f}deg)'
    return 'white'

def group_world_bounds(shapes, parent):
    """Axis-aligned bounds of existing world-space placements, used only to choose
    a stable coordinate origin. Original orientation stays on each placed object.
    This does not simulate the source transform twice or alter drawing order.
    """
    spans=[]
    for sh in shapes:
        if sh.shape_type==MSO_SHAPE_TYPE.GROUP:
            nested=group_world_bounds(sh.shapes,parent @ group_child_matrix(sh))
            if nested:spans.append(nested)
        else:
            pl=placement_info(sh,parent)
            spans.append((pl['x'],pl['y'],pl['x']+pl['w'],pl['y']+pl['h']))
    if not spans:return None
    return (min(v[0] for v in spans),min(v[1] for v in spans),
            max(v[2] for v in spans),max(v[3] for v in spans))

lines=['// Source-faithful PPTX conversion with original shared media and editable text.',
       '// Source slide format, theme-resolved colors and image backgrounds.',
       '#import "template.typ": Root, PptxMaster, PptxLayout, PptxSlide, PptxPlaceholderSlot, PptxPlaceholder, PptxParagraph, PptxTextFrame, PptxPlaceholderParagraph, PptxPlaceholderRun, PptxLine, Place, GroupContent, ClipContent',
       f'#let ppt-format=(width:{SLIDE_W:.3f}pt,height:{SLIDE_H:.3f}pt)',
       f'#set page(width:{SLIDE_W:.3f}pt,height:{SLIDE_H:.3f}pt,margin:0pt,fill:white)',
       '#set text(font:("Arial","Libertinus Serif"),size:20pt,fill:black)',
       '']

preamble=lines[:]
slide_sources=[]
master_declarations={}
layout_declarations={}
master_manifest=[]
layout_manifest=[]
master_placeholder_manifest=[]
layout_placeholder_manifest=[]
placeholder_warnings=[]
master_ids={}
layout_ids={}

for slide_no,slide in enumerate(prs.slides,1):
    CURRENT_THEME,CURRENT_MAP=theme_context(slide)
    CURRENT_FONTS=theme_font_context(slide)
    CURRENT_SLIDE=slide
    bgtype,bgval,bgxml=source_background(slide)
    slide_bg=owner_background(slide)
    if slide_bg and slide_bg[0]=='color':slide_fill=f'rgb("{slide_bg[1]}")'
    elif slide_bg and slide_bg[0]=='gradient':slide_fill=gradient_expr(slide_bg[1])
    elif slide_bg is None:slide_fill='none'
    else:slide_fill='none'
    if bgtype=='color':bg_expr=f'rgb("{bgval}")'
    elif bgtype=='gradient':
        bg_expr=gradient_expr(bgval)
    else:bg_expr='white'
    slide_map_override=bool(slide._element.xpath('./p:clrMapOvr/a:overrideClrMapping'))
    emitted_fill=slide_fill if args.layout_mode=='structured' else bg_expr
    # clrMapOvr can intentionally recolor an inherited bgRef for one slide.
    # Keep the shared master fill canonical and encode that color-map exception
    # as a slide fill override.
    if args.layout_mode=='structured' and slide_bg is None and slide_map_override and bgtype in ('color','gradient'):
        emitted_fill=bg_expr
    lines=[]
    master=slide.slide_layout.slide_master
    layout=slide.slide_layout
    master_part=str(master.part.partname)
    layout_part=str(layout.part.partname)
    master_ids.setdefault(master_part,f'Master{len(master_ids)+1:02d}')
    layout_ids.setdefault(layout_part,f'Layout{len(layout_ids)+1:02d}')
    master_id=master_ids[master_part]
    layout_id=layout_ids[layout_part]
    master_allowed=slide._element.get('showMasterSp','1') not in ('0','false')
    layout_allowed=layout._element.get('showMasterSp','1') not in ('0','false')
    slide_placeholder_records=[]
    slide_placeholder_exprs=[]
    if args.layout_mode=='structured':
        declared_slots=set()
        for owner in (master,layout):
            for candidate in owner.shapes:
                candidate_kind=placeholder_kind(candidate)
                if candidate_kind:
                    declared_slots.add(placeholder_id(candidate_kind,candidate.placeholder_format.idx))
        for shape in slide.shapes:
            kind=placeholder_kind(shape)
            if shape.is_placeholder and not kind:
                placeholder_warnings.append({'owner':'slide','slide':slide_no,'shape_id':shape.shape_id,
                    'placeholder_type':str(shape.placeholder_format.type),'action':'ordinary-object fallback'})
            if not kind or not shape.has_text_frame or not shape.text:
                continue
            idx=shape.placeholder_format.idx
            slot_id=placeholder_id(kind,idx)
            direct=placeholder_direct_geometry(shape)
            effective=placement_info(shape,np.identity(3))
            # An unresolved placeholder without its own complete geometry
            # remains in the ordinary fallback path rather than disappearing.
            if slot_id not in declared_slots and direct is None:
                placeholder_warnings.append({'owner':'slide','slide':slide_no,'shape_id':shape.shape_id,
                    'placeholder_type':kind,'idx':idx,'action':'unresolved slot; ordinary-object fallback'})
                continue
            fields={'id':slot_id,'type':kind,'idx':idx}
            if direct:
                fields.update(direct)
            fields.update(placeholder_style_fields(shape))
            autofit_mode,font_scale,line_reduction=text_autofit(shape)
            fields.update({'autofit-mode':autofit_mode,'font-scale':font_scale,
                           'line-spacing-reduction':line_reduction})
            args_out=[]
            for key,value in fields.items():
                if key in ('id','type','font','weight','align','anchor','bullet','autofit-mode'):value=stylize(str(value))
                elif key in ('fill','background'):pass
                elif key=='idx':value=str(value)
                elif key=='rotation':value=f'{float(value):.3f}deg'
                elif key in ('font-scale','line-spacing-reduction'):value=f'{float(value):.5f}'
                else:value=f'{float(value):.3f}pt'
                args_out.append(f'{key}:{value}')
            paragraph_exprs=[]
            bullet_counters={}
            source_paragraphs = list(shape.text_frame.paragraphs)
            paragraph_styles = paragraph_render_styles(shape, source_paragraphs)
            for paragraph, (empty_size, empty_font) in zip(source_paragraphs, paragraph_styles):
                expression,_,_,_=paragraph_record_expression(
                    shape,paragraph,shrink=font_scale,spacing_reduction=line_reduction,
                    bullet_counters=bullet_counters,
                    record_name='PptxPlaceholderParagraph',empty_paragraph_size=empty_size,
                    empty_paragraph_font=empty_font)
                paragraph_exprs.append(expression)
            paragraph_tuple = '(' + ', '.join(paragraph_exprs) + (',' if len(paragraph_exprs) == 1 else '') + ')'
            args_out.extend([f'text:{stylize(shape.text)}',f'paragraphs:{paragraph_tuple}'])
            slide_placeholder_exprs.append('PptxPlaceholder('+', '.join(args_out)+')')
            slide_placeholder_records.append({
                'id':slot_id,'type':kind,'idx':idx,'text':shape.text,
                'source_shape_id':shape.shape_id,'source_name':shape.name,
                'layout_reference':slot_id if slot_id in declared_slots else None,
                'geometry_override':direct,
                'unresolved_layout_slot':slot_id not in declared_slots,
            })
            if slot_id not in declared_slots:
                placeholder_warnings.append({'owner':'slide','slide':slide_no,'shape_id':shape.shape_id,
                    'placeholder_type':kind,'idx':idx,'action':'native slide placeholder has no reusable owner slot'})
    slide_call='PptxSlide' if args.layout_mode=='structured' else 'Root'
    lines.extend([f'// Slide {slide_no:02d}',f'#{slide_call}(',
                  f'  width:{SLIDE_W:.3f}pt,',f'  height:{SLIDE_H:.3f}pt,',
                  f'  fill:{emitted_fill},','  format:ppt-format,'])
    if args.layout_mode=='structured':lines.append(f'  layout:{layout_id},')
    if args.layout_mode=='structured' and slide_placeholder_exprs:
        lines.append('  placeholders:(')
        lines.extend(f'    {expr},' for expr in slide_placeholder_exprs)
        lines.append('  ),')
    lines.append('  children:[')
    if bgtype=='color':
        try:
            vals=[int(bgval[i:i+2],16) for i in (1,3,5)]
            TEXT_FALLBACK='#FFFFFF' if sum(vals)<315 else '#111111'
        except Exception:TEXT_FALLBACK='#111111'
    else:TEXT_FALLBACK='#111111'
    image_bg=slide_bg if args.layout_mode=='structured' else (('image',bgval,bgxml) if bgtype=='image' else None)
    if image_bg and image_bg[0]=='image':
        bgasset=register_media(image_bg[1])
        bgsha=hashlib.sha256(image_bg[1].blob).hexdigest()
        media_contents[bgsha]['placements'].append({'owner':'slide' if slide_bg else 'inherited-background','owner_id':slide_no,'slide':slide_no,'background':True})
        bgasset,bg_warning=typst_compatible_media(image_bg[1],bgasset,{
            'slide':slide_no,'owner':'slide' if slide_bg else 'inherited-background',
            'issue_location':'slide background'})
        if bgasset:
            lines.append(typst_place(0,0,f'image("{bgasset}",width:{SLIDE_W:.3f}pt,height:{SLIDE_H:.3f}pt,fit:"stretch")'))
            counts['background_images']+=1
    elif bgtype=='color':counts['background_colors']+=1
    elif bgtype=='gradient':counts['background_gradients']+=1
    sm={'slide':slide_no,'layout':layout_id,'master':master_id,'placements':[],'groups':[],
        'placeholders':slide_placeholder_records}
    def recur(shapes,parent=np.identity(3),group_path=(),origin=(0.,0.),indent='  '):
        for zindex,sh in enumerate(shapes):
            if sh.shape_type==MSO_SHAPE_TYPE.GROUP:
                counts['group']+=1
                # Group structure and authored child coordinate system must be
                # retained even if Typst currently flattens physical placement.
                xf=sh._element.grpSpPr.xfrm
                gid='g'+str(slide_no)+'_'+str(sh.shape_id)
                xfmat=group_child_matrix(sh)
                xfdata={k:round(float(v)/EMU,4) for k,v in {
                    'off_x':xf.off.x,'off_y':xf.off.y,
                    'ext_w':xf.ext.cx,'ext_h':xf.ext.cy,
                    'child_off_x':xf.chOff.x,'child_off_y':xf.chOff.y,
                    'child_ext_w':xf.chExt.cx,'child_ext_h':xf.chExt.cy,
                }.items()}
                xfdata.update(rotation_degrees=float(xf.get('rot','0'))/60000,
                              flip_h=xf.get('flipH') in ('1','true'),
                              flip_v=xf.get('flipV') in ('1','true'))
                sm['groups'].append({'id':gid,'parent':group_path[-1] if group_path else None,
                                     'name':sh.name,'source_shape_id':sh.shape_id,
                                     'transform':xfdata,'matrix_pt':xfmat.round(8).tolist()})
                group_parent=parent@xfmat
                bounds=group_world_bounds(sh.shapes,group_parent)
                if bounds and (args.group_mode=='nested' or (args.group_mode=='safe' and bounds[0]>=0 and bounds[1]>=0 and bounds[2]<=SLIDE_W and bounds[3]<=SLIDE_H)):
                    bx,by,br,bb=bounds
                    # Each child retains its v5-computed absolute geometry; only
                    # the coordinate origin changes. This is intentionally NOT an
                    # additional group rotation, which would double-transform it.
                    # Original rotation, flips and child-space are in the manifest.
                    lines.append(f'{indent}// Group: {typst_comment(sh.name)} | id: {gid}')
                    lines.extend([f'{indent}#Place(',f'{indent}  x:{bx-origin[0]:.3f}pt,',f'{indent}  y:{by-origin[1]:.3f}pt,',
                                  f'{indent}  child:GroupContent(width:{br-bx:.3f}pt,height:{bb-by:.3f}pt,children:[' ])
                    counts['group_containers']+=1
                    sm['groups'][-1]['render_bounds_pt']=[round(v,4) for v in bounds]
                    recur(sh.shapes,group_parent,group_path+(gid,),(bx,by),indent+'    ')
                    lines.append(f'{indent}  ]))')
                else:
                    recur(sh.shapes,group_parent,group_path+(gid,),origin,indent)
                continue
            if None in (sh.left,sh.top,sh.width,sh.height):
                # Some slide-number/date/footer placeholders have no transform
                # on the slide, layout, or master. Keep converting the deck and
                # report the unplaced placeholder instead of aborting all slides.
                counts['unsupported_unpositioned_shape']+=1
                counts['unsupported_placements']+=1
                obj={'z_index':zindex,'group_path':list(group_path),
                     'source_shape_id':sh.shape_id,'name':sh.name,
                     'type':'unsupported','unsupported_kind':'unpositioned_shape',
                     'text':sh.text[:200] if sh.has_text_frame else '',
                     'warning':'Shape has no resolvable position or size in slide, layout, or master.'}
                sm['placements'].append(obj)
                continue
            pl=placement_info(sh,parent)
            obj={'z_index':zindex,'group_path':list(group_path),'source_shape_id':sh.shape_id,'name':sh.name,'x':round(pl['x'],2),'y':round(pl['y'],2),'w':round(pl['w'],2),'h':round(pl['h'],2),'rot':round(pl['rotation'],2),'flip_y':pl['flip_y'],'group_skew':round(pl['skew'],5)}
            # Preserve established fixes from first converter for text boxes known to differ in rendered metrics
            smartart=smartart_drawing(sh)
            if smartart is not None:
                asset=smartart_svg(sh,smartart)
                if asset:
                    lines.append(f'{indent}// SmartArt layout drawing: {typst_comment(sh.name)}')
                    lines.append(typst_place(pl['x']-origin[0],pl['y']-origin[1],
                                             f'image("{asset}",width:{pl["w"]:.3f}pt,height:{pl["h"]:.3f}pt,fit:"stretch")',pl['rotation'],indent))
                    obj.update(type='smartart',source=asset,rendering='materialized vector drawing')
                    counts['smartart_placements']+=1
                    sm['placements'].append(obj)
                    continue
            if getattr(sh,'has_table',False) and not table_has_spans(sh):
                table=typst_table(sh)
                lines.append(f'{indent}// Editable PowerPoint table: {typst_comment(sh.name)}')
                lines.append(typst_place(pl['x']-origin[0],pl['y']-origin[1],table,pl['rotation'],indent))
                obj.update(type='table',rows=len(sh.table.rows),columns=len(sh.table.columns))
                counts['table_placements']+=1
                sm['placements'].append(obj)
                continue
            istext=sh.has_text_frame and bool(sh.text.strip())
            imgpart=image_blip(sh)
            if abs(pl['skew'])>0.001:
                counts['nonorthogonal_group_objects']+=1
                obj['transform_warning']='Nonorthogonal group transform: Place orientation approximates shear.'
            if imgpart:
                counts['media_placements']+=1
                rel=register_media(imgpart)
                sha=next(k for k,v in registered.items() if v==rel)
                crop,fillrect,alpha=image_fill_params(sh)
                authored_crop=({**crop} if crop else None)
                authored_fillrect=({**fillrect} if fillrect else None)
                rendered_rel,media_warning=typst_compatible_media(imgpart,rel,{
                    'slide':slide_no,'shape_id':sh.shape_id,'shape':sh.name,
                    'owner':sm.get('owner','slide'),'owner_id':sm.get('id',slide_no)})
                if rendered_rel is None:
                    media_contents[sha]['placements'].append({
                        'owner':sm.get('owner','slide'),'owner_id':sm.get('id',slide_no),
                        'slide':slide_no if sm.get('owner','slide')=='slide' else None,
                        'shape':sh.name,'crop':authored_crop,'fillrect':authored_fillrect,
                        'alpha':alpha,'geometry':geometry_kind(sh,effective_shape_geometry(sh)),
                        'render_warning':media_warning['reason']})
                    obj.update(type='image',source=rel,render_source=None,crop=authored_crop,
                               fillrect=authored_fillrect,alpha=alpha,media_warning=media_warning['reason'],
                               expected_loss=media_warning['expected_loss'])
                    counts['unsupported_media_placements']+=1
                    sm['placements'].append(obj)
                    continue
                rendered_rel=alpha_variant(imgpart,rendered_rel,alpha)
                geometry_element=effective_shape_geometry(sh)
                shape_kind=geometry_kind(sh,geometry_element)
                # Only genuinely masked photos receive a derived transparent image.
                # The exact original file remains registered and shared; all ordinary
                # (including rect-coded custom) photos reference it directly.
                if shape_kind != 'rect':
                    preset,mask_paths=shape_paths(geometry_element)
                    # Fail closed on adjustable Office preset guides and compound
                    # rotated freeform picture fills: v2's original-image placement
                    # is visually safer until guide/transform parity is verified.
                    # Existing v6 polygon support remains unchanged. Source-space
                    # masking happens *before* the parent Place orientation, so a
                    # rotated ellipse is safe when the local geometry is known.
                    legacy_polygon=(preset=='custom' and shape_kind=='polygon'
                                    and abs(pl['rotation'])<.01 and not pl['flip_y']
                                    and abs(pl['skew'])<.001)
                    inherited_layout_mask=(geometry_element is not sh._element and
                                           preset=='custom' and shape_kind=='complex' and
                                           abs(pl['rotation'])<.01 and not pl['flip_y'] and
                                           abs(pl['skew'])<.001)
                    validated_curve=(preset in ('ellipse','roundRect') and
                                     abs(pl['skew'])<.001 and not pl['flip_y']
                                     and abs(pl['rotation'])<.01)
                    experimental_curve=(shape_kind=='complex' and preset=='custom' and
                                        abs(pl['skew'])<.001 and not pl['flip_y'])
                    rotated_preset=(preset in ('ellipse','roundRect') and
                                    abs(pl['skew'])<.001 and not pl['flip_y'])
                    mask_is_safe=(legacy_polygon or inherited_layout_mask or
                                  (args.mask_mode in ('validated','experimental') and validated_curve) or
                                  (args.mask_mode=='experimental' and (experimental_curve or rotated_preset)))
                    # Allow users to opt in only specific PPTX shapes after a
                    # side-by-side comparison. Useful when an Office mask uses
                    # undocumented image-position semantics.
                    if MASK_ALLOWLIST is not None and (slide_no,sh.shape_id) not in MASK_ALLOWLIST:
                        mask_is_safe=False
                        counts['mask_excluded_by_allowlist']+=1
                    if mask_paths and mask_is_safe:
                        sha_mask=hashlib.sha256(imgpart.blob).hexdigest()
                        masked=mask_image(imgpart.blob,pl['w'],pl['h'],mask_paths,crop,fillrect,alpha,ASSETS,sha_mask)
                        if masked:
                            rendered_rel=masked;crop=fillrect=None
                            counts['clipped_images']+=1;counts['masked_'+shape_kind]+=1
                        else:counts['clip_mask_failures']+=1
                    else:
                        counts['unsupported_clip_masks']+=1
                        obj['clip_warning']='Mask uses adjustable presets, curves, or rotation not yet validated. Original picture kept without a generated mask.'
                core=typst_img(rendered_rel,pl,crop,fillrect)
                lines.append(f'{indent}// Source media: {typst_comment(rel)} | shape: {typst_comment(sh.name)} | {typst_comment(shape_kind)}')
                lines.append(typst_place(pl['x']-origin[0],pl['y']-origin[1],core,pl['rotation'],indent))
                media_contents[sha]['placements'].append({'owner':sm.get('owner','slide'),'owner_id':sm.get('id',slide_no),'slide':slide_no if sm.get('owner','slide')=='slide' else None,'shape':sh.name,'crop':authored_crop,'fillrect':authored_fillrect,'alpha':alpha,'geometry':shape_kind})
                obj.update(type='image',source=rel,render_source=rendered_rel,geometry=shape_kind,crop=authored_crop,fillrect=authored_fillrect,render_crop=crop,render_fillrect=fillrect,alpha=alpha,mask_applied=(rendered_rel!=rel and rendered_rel.startswith('assets/media/clip_')))
                counts['geometry_'+shape_kind]+=1
            elif sh.shape_type==MSO_SHAPE_TYPE.LINE or sh._element.xpath('./p:spPr/a:prstGeom[@prst="straightConnector1"]'):
                ln=sh._element.xpath('./p:spPr/a:ln')
                if ln:
                    color=color_in(ln[0],'./a:solidFill','#000000')
                    width=max(.1,float(ln[0].get('w','12700'))/EMU)
                    start_y=pl['h'] if pl['flip_y'] else 0.0
                    end_y=0.0 if pl['flip_y'] else pl['h']
                    core=(f'PptxLine(start:(0pt,{start_y:.3f}pt), '
                          f'end:({pl["w"]:.3f}pt,{end_y:.3f}pt), '
                          f'stroke:{width:.3f}pt + rgb("{color}"))')
                    lines.append(f'{indent}// Native editable PowerPoint connector: {typst_comment(sh.name)}')
                    lines.append(typst_place(pl['x']-origin[0],pl['y']-origin[1],core,pl['rotation'],indent))
                    obj.update(type='line',stroke=color,stroke_width_pt=width)
                    ends=ln[0].xpath('./a:headEnd[@type!="none"] | ./a:tailEnd[@type!="none"]')
                    if ends:
                        obj['warning']='Arrowhead style is not represented by the Typst line subset.'
                        counts['unsupported_line_markers']+=1
                    counts['line_shapes']+=1
                else:
                    obj.update(type='unsupported',warning='Connector has no resolvable line style.')
                    counts['unsupported_placements']+=1
            elif istext:
                # A PowerPoint text-bearing shape also paints its own fill and
                # stroke before the text. v3 dropped that shape completely,
                # which made many cards and labels disappear.
                backing=vector_svg(sh)
                if backing:
                    rel,pad,w0,h0=backing
                    pscale=(pl['w']/w0+pl['h']/h0)/2 if w0*h0 else 1
                    extra=pad*pscale
                    lines.append(f'{indent}// Fill and outline behind text shape: {typst_comment(sh.name)}')
                    lines.append(typst_place(pl['x']-extra-origin[0],pl['y']-extra-origin[1],
                                             f'image("{rel}",width:{pl["w"]+2*extra:.3f}pt,height:{pl["h"]+2*extra:.3f}pt,fit:"stretch")',pl['rotation'],indent))
                    obj['background_shape']=rel
                    counts['text_backing_shapes']+=1
                counts['text_placements']+=1
                core=content_for_text(sh,pl['w'],pl['h'],1,1)
                # Source text can contain vertical tabs and Unicode line breaks.
                # Keep it on one comment line so it cannot become Typst markup.
                comment=typst_comment(f'{sh.name} | {sh.text[:80]}')
                lines.append(f'{indent}// Editable text: {comment}')
                # Keep the text frame at the source shape origin. Text insets,
                # paragraph anchors, and point sizes are handled inside its box.
                lines.append(typst_place(pl['x']-origin[0],pl['y']-origin[1],core,pl['rotation'],indent))
                obj.update(type='text',text=sh.text[:200],anchor=(sh._element.xpath('./p:txBody/a:bodyPr')[0].get('anchor','t') if sh._element.xpath('./p:txBody/a:bodyPr') else 't'),autofit_mode=text_autofit(sh)[0],font_scale=text_autofit(sh)[1],line_spacing_reduction=text_autofit(sh)[2])
                counts['autofit_'+text_autofit(sh)[0]]+=1
            else:
                counts['other_placements']+=1
                obj.update(type='other')
                vec=vector_svg(sh)
                if vec:
                    rel,pad,w0,h0=vec
                    # Stroke padding also scales with grouped shape transformation.
                    pscale=(pl['w']/w0+pl['h']/h0)/2 if w0*h0 else 1
                    extra=pad*pscale
                    lines.append(f'{indent}// Original editable PowerPoint shape: {typst_comment(sh.name)}')
                    lines.append(typst_place(pl['x']-extra-origin[0],pl['y']-extra-origin[1],
                                             f'image("{rel}",width:{pl["w"]+2*extra:.3f}pt,height:{pl["h"]+2*extra:.3f}pt,fit:"stretch")',pl['rotation'],indent))
                    obj['type']='vector_shape';obj['source']=rel;counts['vector_shapes']+=1
                else:
                    obj['type']='unsupported'
                    xml=sh._element
                    kind='chart' if xml.xpath('.//*[local-name()="chart"]') else ('table' if xml.xpath('.//a:tbl') else 'shape')
                    obj['unsupported_kind']=kind
                    obj['warning']='Unsupported '+kind+': original preserved in source PPTX, not silently claimed as converted'
                    counts['unsupported_'+kind]+=1
                    counts['unsupported_placements']+=1
            sm['placements'].append(obj)
    # Masters are the backmost paint layers. PPTX placeholders are materialized
    # by the slide itself, so only import non-placeholder master decorations.
    def decorative_shapes(owner):
        return [sh for sh in owner.shapes if not sh.is_placeholder]
    def background_fill(owner):
        background=owner_background(owner)
        if not background:return 'none','none'
        if background[0]=='color':return f'rgb("{background[1]}")','none'
        if background[0]=='gradient':return gradient_expr(background[1],owner_background_color_map(owner)),'none'
        if background[0]=='image':
            asset=register_media(background[1])
            bgsha=hashlib.sha256(background[1].blob).hexdigest()
            media_contents[bgsha]['placements'].append({'owner':sm['owner'],'owner_id':sm['id'],'background':True})
            asset,warning=typst_compatible_media(background[1],asset,{
                'owner':sm['owner'],'owner_id':sm['id'],'issue_location':f'{sm["owner"]} background'})
            if asset:
                counts['background_images']+=1
                return 'none',f'image("{asset}",width:{SLIDE_W:.3f}pt,height:{SLIDE_H:.3f}pt,fit:"stretch")'
            return 'none','none'
        return 'none','none'
    master_allowed=slide._element.get('showMasterSp','1') not in ('0','false')
    layout_allowed=slide.slide_layout._element.get('showMasterSp','1') not in ('0','false')
    if args.layout_mode=='structured':
        if master_part not in master_declarations:
            saved_lines,saved_sm=lines,sm
            lines=[]
            sm={'owner':'master','id':master_id,'part':master_part,'placements':[],'groups':[]}
            master_fill,master_background=background_fill(master)
            master_placeholder_exprs=[]
            for candidate in master.shapes:
                if candidate.is_placeholder and not placeholder_kind(candidate):
                    placeholder_warnings.append({'owner':'master','owner_id':master_id,'shape_id':candidate.shape_id,
                        'placeholder_type':str(candidate.placeholder_format.type),'action':'unsupported; omitted from structured output'})
                resolved=placeholder_slot_expression(candidate)
                if resolved:
                    expr,record=resolved
                    if record['id'] not in {item['id'] for item in master_placeholder_manifest if item.get('owner_id')==master_id}:
                        master_placeholder_exprs.append(expr)
                        master_placeholder_manifest.append({'owner':'master','owner_id':master_id,**record})
            recur(decorative_shapes(master))
            lines.append('])')
            master_slots = ', '.join(master_placeholder_exprs)
            if len(master_placeholder_exprs) == 1:
                master_slots += ','
            master_declarations[master_part]=(f'#let {master_id} = PptxMaster(fill:{master_fill},background:{master_background},placeholders:('+master_slots+'),children:[',lines[:])
            master_manifest.append(sm)
            lines,sm=saved_lines,saved_sm
            counts['inherited_master_layers']+=1
        if layout_part not in layout_declarations:
            saved_lines,saved_sm=lines,sm
            lines=[]
            sm={'owner':'layout','id':layout_id,'master':master_id,'part':layout_part,'placements':[],'groups':[]}
            layout_fill,layout_background=background_fill(layout)
            layout_placeholder_exprs=[]
            for candidate in layout.shapes:
                if candidate.is_placeholder and not placeholder_kind(candidate):
                    placeholder_warnings.append({'owner':'layout','owner_id':layout_id,'shape_id':candidate.shape_id,
                        'placeholder_type':str(candidate.placeholder_format.type),'action':'unsupported; omitted from structured output'})
                resolved=placeholder_slot_expression(candidate)
                if resolved:
                    expr,record=resolved
                    if record['id'] not in {item['id'] for item in layout_placeholder_manifest if item.get('owner_id')==layout_id}:
                        layout_placeholder_exprs.append(expr)
                        layout_placeholder_manifest.append({'owner':'layout','owner_id':layout_id,'master':master_id,**record})
            recur(decorative_shapes(layout))
            lines.append('])')
            layout_slots = ', '.join(layout_placeholder_exprs)
            if len(layout_placeholder_exprs) == 1:
                layout_slots += ','
            layout_declarations[layout_part]=(f'#let {layout_id} = PptxLayout(master:{master_id},fill:{layout_fill},background:{layout_background},placeholders:('+layout_slots+f'),show-master:{str(layout_allowed).lower()},children:[',lines[:])
            layout_manifest.append(sm)
            lines,sm=saved_lines,saved_sm
    else:
        if master_allowed and layout_allowed:
            recur(decorative_shapes(master))
            counts['inherited_master_layers']+=1
        if layout_allowed:
            recur(decorative_shapes(layout))
            counts['inherited_layout_layers']+=1
    # PowerPoint materializes text placeholders above later slide artwork in
    # its rendered slide, even when the placeholder node precedes a picture in
    # the slide's spTree. Keep ordinary shape order intact, then paint text
    # placeholders last so Typst preserves that visible layering.
    native_placeholder_ids={record['source_shape_id'] for record in slide_placeholder_records}
    slide_shapes=[shape for shape in slide.shapes if shape.shape_id not in native_placeholder_ids]
    slide_shapes.sort(key=lambda shape: bool(
        shape.is_placeholder and shape.has_text_frame and shape.text.strip()
    ))
    recur(slide_shapes,indent='    ')
    lines.append('])')
    if slide_no<len(prs.slides):lines.append('#pagebreak()')
    slide_sources.extend(lines)
    manifest.append(sm)

master_lines=[line for decl in master_declarations.values() for line in [decl[0],*decl[1]]]
layout_lines=[line for decl in layout_declarations.values() for line in [decl[0],*decl[1]]]
(OUT/'main.typ').write_text('\n'.join(preamble+master_lines+layout_lines+slide_sources),encoding='utf8')
for owner in [*master_manifest,*layout_manifest,*manifest]:
    for placement in owner.get('placements',[]):
        for field,issue in (('warning','unsupported_or_approximate_object'),
                            ('clip_warning','unsupported_image_mask'),
                            ('transform_warning','approximate_group_transform')):
            if placement.get(field):
                media_warnings.append({
                    'issue':issue,'slide':owner.get('slide'),
                    'owner':owner.get('owner','slide'),
                    'owner_id':owner.get('id',owner.get('slide')),
                    'shape_id':placement.get('source_shape_id'),
                    'shape':placement.get('name'),
                    'fallback':placement.get('rendering','original artwork preserved or omitted as noted'),
                    'expected_loss':placement.get('expected_loss',placement.get(field)),
                })
asset_manifest={'source_file':SRC.name,'source_slide_count':len(prs.slides),
    'slide_dimensions_pt':[SLIDE_W,SLIDE_H],'unique_images':len(media_contents),
    'unique_vectors':len(vector_assets),'alpha_variants':len(ALPHA_CACHE),
    'group_mode':args.group_mode,'clip_mode':args.clip_mode,'mask_mode':args.mask_mode,
    'layout_mode':args.layout_mode,'counts':dict(counts),'masters':master_manifest,
    'layouts':layout_manifest,'master_placeholders':master_placeholder_manifest,
    'layout_placeholders':layout_placeholder_manifest,'placeholder_warnings':placeholder_warnings,
    'media_warnings':media_warnings,'media':media_contents,'slides':manifest}
(OUT/'asset_manifest.json').write_text(json.dumps(asset_manifest,ensure_ascii=False,indent=2),encoding='utf8')
(OUT/'conversion_report.json').write_text(json.dumps({
    'source_file':SRC.name,'slides':len(prs.slides),
    'status':'completed_with_warnings' if media_warnings else 'completed',
    'counts':dict(counts),'warnings':media_warnings,
},ensure_ascii=False,indent=2),encoding='utf8')
print('Saved',OUT,'assets',len(media_contents),'stats',dict(counts),'sizesMB',sum(v['bytes'] for v in media_contents.values())/1024**2)
