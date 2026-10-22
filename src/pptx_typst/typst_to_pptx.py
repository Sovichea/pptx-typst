#!/usr/bin/env python3
"""Convert the supported Place-based Typst subset into editable PowerPoint.

Static PptxMaster, PptxLayout, and PptxSlide declarations are written as real
PresentationML owner parts. This is not a general Typst interpreter. Unsupported
constructs are reported in a JSON audit.
"""
import argparse, ast, hashlib, io, json, math, os, re
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

PX_PER_PT = 1.0
NUM = r'[-+]?\d+(?:\.\d+)?'

# Calls provided by the starter module are intentionally thin aliases for the
# supported Root/Place and native Typst primitives. Normalize them before
# parsing so users can author through the module without expanding arbitrary
# Typst functions or evaluating Typst code.
STARTER_ALIASES = {
    'PptxSlide': 'Root',
    'PptxPlace': 'Place',
    'PptxGroup': 'GroupContent',
    'PptxClip': 'box',
    'PptxBox': 'box',
    'PptxText': 'text',
    'PptxImage': 'image',
    'PptxRect': 'rect',
    'PptxEllipse': 'ellipse',
    'PptxLine': 'line',
    'PptxTable': 'table',
}


def normalize_starter_aliases(content):
    for alias, target in STARTER_ALIASES.items():
        content = re.sub(
            r'(?<![\w-])(#?)' + alias + r'(?=\s*\()',
            lambda match: match.group(1) + target,
            content,
        )
    return content


def balanced(s, pos, left='(', right=')'):
    assert s[pos] == left, (s[pos:pos+30],left)
    d=1; q=None; esc=False; i=pos+1
    while i<len(s):
        c=s[i]
        if q:
            if esc: esc=False
            elif c=='\\': esc=True
            elif c==q: q=None
        elif c in ('"',"'"): q=c
        elif c==left: d+=1
        elif c==right:
            d-=1
            if d==0: return s[pos+1:i], i+1
        i+=1
    raise ValueError(f'unbalanced {left} at {pos}')


def argsplit(raw):
    arr=[]; start=0; qs=None; esc=False; stack=[]; opens={'(':')','[':']','{':'}'}
    for i,c in enumerate(raw):
        if qs:
            if esc: esc=False
            elif c=='\\': esc=True
            elif c==qs: qs=None
        elif c in ('"',"'"):qs=c
        elif c in opens:stack.append(opens[c])
        elif stack and c==stack[-1]:stack.pop()
        elif c==',' and not stack:
            arr.append(raw[start:i].strip());start=i+1
    tail=raw[start:].strip()
    if tail:arr.append(tail)
    return arr


def kwargs(raw):
    r={};pos=[]
    for a in argsplit(raw):
        m=re.match(r'^([\w-]+)\s*:\s*(.*)$',a,re.S)
        if m:r[m.group(1)]=m.group(2).strip()
        else:pos.append(a)
    if pos:r['_pos']=pos
    return r


def length(s,default=0.0):
    if s is None:return default
    m=re.search(NUM,s)
    return float(m.group()) if m else default


def find_calls(s,which):
    pat=re.compile(r'#'+re.escape(which)+r'\(')
    i=0
    while True:
        m=pat.search(s,i)
        if not m:break
        raw,end=balanced(s,m.end()-1)
        yield raw,m.start(),end
        i=end


def simple_color(s,default=(0,0,0)):
    if not s:return default
    m=re.search(r'#([\da-fA-F]{6})',s)
    if m:
        h=m.group(1);return tuple(int(h[i:i+2],16) for i in (0,2,4))
    if 'white' in s:return (255,255,255)
    if 'black' in s:return (0,0,0)
    return default


def simple_color_alpha(s):
    """Return Typst's RGB color and optional alpha from #RRGGBBAA syntax."""
    match=re.search(r'#([\da-fA-F]{8})',s or '')
    return (tuple(int(match.group(1)[i:i+2],16) for i in (0,2,4)),
            int(match.group(1)[6:8],16)/255) if match else (simple_color(s),1.0)


def txt_decode(v):
    if not v:return ''
    v=v.strip()
    if v.startswith('"'):
        try:return json.loads(v)
        except Exception:
            try:return ast.literal_eval(v)
            except Exception:return v.strip('"')
    if v.startswith('[') and v.endswith(']'):return v[1:-1]
    return v


def img_call(child):
    match = re.search(r'#?image\(',child)
    if not match:return None
    raw,end=balanced(child,match.end()-1)
    a=kwargs(raw)
    if not a.get('_pos'):return None
    path=txt_decode(a['_pos'][0]);return dict(path=path,iw=length(a.get('width')),ih=length(a.get('height')),fit=txt_decode(a.get('fit','"stretch"')))


def flip(child):
    m=re.search(r'(?<![\w])scale\(\s*x\s*:\s*(-?\d+)%\s*,\s*y\s*:\s*(-?\d+)%',child)
    return (float(m.group(1))<0,float(m.group(2))<0) if m else (False,False)


def viewport(child):
    # PowerPoint image fill: explicit clip box and inner image shifted into viewport.
    m=re.search(r'box\(\s*width\s*:\s*('+NUM+r')pt\s*,\s*height\s*:\s*('+NUM+r')pt\s*,\s*clip\s*:\s*true',child)
    if not m:return None
    w,h=map(float,m.groups())
    n=re.search(r'#place\(\s*top\s*\+\s*left\s*,\s*dx\s*:\s*('+NUM+r')pt\s*,\s*dy\s*:\s*('+NUM+r')pt',child)
    dx,dy=map(float,n.groups()) if n else (0,0)
    return w,h,dx,dy


def text_runs(child):
    out=[]
    # Starter-module calls such as ``child: PptxText(...)`` are ordinary
    # function expressions and therefore have no leading Typst markup ``#``.
    # The generated Root/Place form also commonly contains ``#text(...)``.
    pattern=re.compile(r'(?<![\w-])#?text\(')
    cursor=0
    while True:
        match=pattern.search(child,cursor)
        if not match:break
        raw,end=balanced(child,match.end()-1)
        cursor=end
        a=kwargs(raw); pos=a.get('_pos',[])
        if pos:
            content=txt_decode(pos[-1])
        else:
            block=re.match(r'\s*\[',child[end:])
            if not block:continue
            body_start=end+block.end()-1
            content,_=balanced(child,body_start,'[',']')
            content=content.strip()
            content=re.sub(r'\\([@#$\[\]\\])',r'\1',content)
        if not content:continue
        out.append(dict(text=content,font=txt_decode(a.get('font','"Aptos"')),size=length(a.get('size'),20),color=simple_color(a.get('fill')),bold='bold' in a.get('weight','').lower()))
    return out


def show_unsupported(audit,kind,expr,slide):
    audit['unsupported'].append({'slide':slide,'kind':kind,'example':expr[:170].replace('\n',' ')})


def add_image(slide, obj, base, x,y,rot,alpha_cache,audit,slide_num):
    child=obj; im=img_call(child)
    if not im:return False
    root=base.resolve()
    f=(root/im['path']).resolve()
    if not f.is_relative_to(root) or not f.is_file():
        show_unsupported(audit,'missing_media',str(f),slide_num);return True
    W=im['iw'];H=im['ih']; crop=viewport(child)
    if W<=0 or H<=0:return True
    xf,yf=flip(child)
    if crop:
        w,h,dx,dy=crop
    else:
        w,h,dx,dy=W,H,0.0,0.0
    # Case negative crop or source extends under/over requested viewport: use one object-only
    # raster fallback, retaining original artwork and alpha. This does not flatten slides.
    left=-dx/W; top=-dy/H; right=1-(w-dx)/W; bottom=1-(h-dy)/H
    safe=all(v>=-0.0001 for v in (left,top,right,bottom)) and left+right<0.999 and top+bottom<0.999
    image_path=f
    if f.suffix.lower()=='.svg':
        import cairosvg
        token=hashlib.sha256(f.read_bytes()+f'{W:.2f}x{H:.2f}'.encode()).hexdigest()[:14]
        image_path=alpha_cache/f'{token}.png'
        if not image_path.exists():
            cairosvg.svg2png(url=str(f),write_to=str(image_path),output_width=max(8,round(W*2)),output_height=max(8,round(H*2)))
        audit['svg_assets_as_png']+=1
    elif f.suffix.lower() not in ('.png','.jpeg','.jpg','.bmp','.gif','.tif','.tiff'):
        show_unsupported(audit,'media_format',im['path'],slide_num);return True
    if not safe:
        # produce a raster *instance* only when crop requires compositing beyond bounds.
        name=hashlib.sha256((str(f)+str(crop)+str((W,H))).encode()).hexdigest()[:16]+'.png'
        dst=alpha_cache/name
        if not dst.exists():
            source=Image.open(image_path).convert('RGBA')
            scale=2
            src=source.resize((max(1,round(W*scale)),max(1,round(H*scale))),Image.Resampling.LANCZOS)
            canvas=Image.new('RGBA',(max(1,round(w*scale)),max(1,round(h*scale))),(0,0,0,0))
            canvas.paste(src,(round(dx*scale),round(dy*scale)),src)
            canvas.save(dst)
        image_path=dst
        audit['cropped_instance_fallbacks']+=1
        left=right=top=bottom=0
    # Preserve Typst image fitting. Cover uses a native PowerPoint crop so the
    # source bitmap remains replaceable; contain keeps the asset's aspect ratio.
    image_w,image_h=image_path and Image.open(image_path).size or (0,0)
    fit=im.get('fit','stretch')
    fit_crop=None
    if not crop and image_w and image_h and fit in ('cover','contain'):
        source_ratio=image_w/image_h; target_ratio=w/h
        if fit=='cover':
            if source_ratio>target_ratio:
                side=(1-target_ratio/source_ratio)/2;fit_crop=(side,0.0,side,0.0)
            elif source_ratio<target_ratio:
                side=(1-source_ratio/target_ratio)/2;fit_crop=(0.0,side,0.0,side)
        elif source_ratio>target_ratio:
            fit_h=w/source_ratio; y += (h-fit_h)/2; h=fit_h
        else:
            fit_w=h*source_ratio; x += (w-fit_w)/2; w=fit_w
    shape=slide.shapes.add_picture(str(image_path),Pt(x),Pt(y),width=Pt(w),height=Pt(h))
    if safe and crop:
        shape.crop_left=max(0,left);shape.crop_top=max(0,top)
        shape.crop_right=max(0,right);shape.crop_bottom=max(0,bottom)
    elif fit_crop:
        shape.crop_left,shape.crop_top,shape.crop_right,shape.crop_bottom=fit_crop
    if abs(rot)>0.001:shape.rotation=rot%360
    if xf or yf:
        transform=shape._element.spPr.xfrm
        if xf:transform.set('flipH','1')
        if yf:transform.set('flipV','1')
    audit['pictures']+=1
    audit['source_media'].add(str(f.relative_to(base)))
    audit['items'].append({'slide':slide_num,'type':'image','asset':str(f.relative_to(base)),'x':round(x,2),'y':round(y,2),'w':round(w,2),'h':round(h,2),'rotation':round(rot,2),'crop':[round(max(0,i),3) for i in (left,top,right,bottom)]})
    return True


def add_text(slide,child,x,y,rot,audit,slide_num):
    runs=text_runs(child)
    if not runs:return False
    b=re.search(r'box\(\s*width\s*:\s*('+NUM+r')pt\s*,\s*height\s*:\s*('+NUM+r')pt',child)
    if b:w,h=map(float,b.groups())
    else:w=max(150,len(''.join(r['text'] for r in runs))*8);h=max(32,runs[0]['size']*2)
    if w<=0 or h<=0:return True
    shape=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h))
    tf=shape.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=Pt(0);tf.margin_right=Pt(0);tf.margin_top=Pt(0);tf.margin_bottom=Pt(0)
    am=re.search(r'#align\(\s*(center|left|right)',child)
    if am:
        tf.paragraphs[0].alignment={'center':PP_ALIGN.CENTER,'left':PP_ALIGN.LEFT,'right':PP_ALIGN.RIGHT}[am.group(1)]
    txt=''.join(r['text'] for r in runs)
    # Keep rich runs editable. Typst paragraph-height differences are not benchmarked.
    p=tf.paragraphs[0]
    for rr in runs:
        r=p.add_run();r.text=rr['text'];r.font.name=rr['font'];r.font.size=Pt(rr['size']);r.font.bold=rr['bold']
        r.font.color.rgb=RGBColor(*rr['color'])
    if abs(rot)>0.001:shape.rotation=rot%360
    audit['text_boxes']+=1;audit['text_content'].append(txt)
    audit['items'].append({'slide':slide_num,'type':'text','x':round(x,2),'y':round(y,2),'w':round(w,2),'h':round(h,2),'rotation':round(rot,2),'chars':len(txt)})
    return True


def add_native(slide,child,x,y,rot,audit,slide_num):
    for ty in ('rect','circle','ellipse','line'):
        m=re.match(r'#?'+ty+r'\(',child)
        if not m:continue
        a=kwargs(balanced(child,m.end()-1)[0])
        w=length(a.get('width'),100);h=length(a.get('height'),w)
        radius=length(a.get('radius'))
        if ty in ('circle','ellipse'):
            if ty=='circle':w=h=length(a.get('radius'),50)*2 if 'radius' in a else length(a.get('diameter'),100)
            shp=MSO_SHAPE.OVAL
        elif ty=='rect':shp=MSO_SHAPE.ROUNDED_RECTANGLE if radius>0 else MSO_SHAPE.RECTANGLE
        else:shp=None
        if shp:
            obj=slide.shapes.add_shape(shp,Pt(x),Pt(y),Pt(w),Pt(h))
            fill=a.get('fill','none')
            if fill=='none':obj.fill.background()
            else:
                color,opacity=simple_color_alpha(fill)
                obj.fill.solid();obj.fill.fore_color.rgb=RGBColor(*color)
                if opacity<.99999:
                    solid_fill=obj._element.spPr.xpath('./a:solidFill')
                    srgb=solid_fill[0].xpath('./a:srgbClr') if solid_fill else []
                    if srgb:
                        alpha=OxmlElement('a:alpha');alpha.set('val',str(round(opacity*100000)));srgb[0].append(alpha)
            stroke=a.get('stroke','none')
            if stroke=='none':obj.line.fill.background()
            else:
                obj.line.color.rgb=RGBColor(*simple_color(stroke,(0,0,0)))
                obj.line.width=Pt(max(.1,length(stroke,.8)))
        else:
            def point_pair(value, fallback):
                nums=re.findall(NUM,value or '')
                return (float(nums[0]),float(nums[1])) if len(nums)>=2 else fallback
            start=point_pair(a.get('start'),(0.0,0.0))
            end=point_pair(a.get('end'),(w,h))
            obj=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                Pt(x+start[0]),Pt(y+start[1]),Pt(x+end[0]),Pt(y+end[1]))
            stroke=a.get('stroke','none')
            if stroke=='none':obj.line.fill.background()
            else:
                obj.line.color.rgb=RGBColor(*simple_color(stroke,(0,0,0)))
                obj.line.width=Pt(max(.1,length(stroke,.8)))
        if rot:obj.rotation=rot%360
        audit['native_shapes']+=1
        audit['items'].append({'slide':slide_num,'type':ty,'x':x,'y':y,'w':w,'h':h})
        return True
    return False


def table_cells(raw):
    """Read the generated Typst table subset and its per-cell formatting."""
    m=re.match(r'#?table\(',raw)
    if not m:return None
    args=kwargs(balanced(raw,m.end()-1)[0])
    cells=[]
    for cell in args.get('_pos',[]):
        style={};body=cell
        cm=re.match(r'#?table\.cell\(',cell)
        if cm:
            style=kwargs(balanced(cell,cm.end()-1)[0])
            _,end=balanced(cell,cm.end()-1)
            rest=cell[end:].strip()
            if not rest.startswith('['):return None
            body=balanced(rest,0,'[',']')[0]
        elif cell.startswith('[') and cell.endswith(']'):
            body=cell[1:-1]
        else:return None
        cells.append((style,body))
    def values(name):
        return [length(v) for v in re.findall(NUM+r'\s*(?:pt|mm|cm|in)?',args.get(name,''))]
    cols=values('columns');rows=values('rows')
    if not cells or not cols or not rows or len(cells)!=len(cols)*len(rows):return None
    return cols,rows,cells


def add_table(slide,child,x,y,rot,audit,slide_num):
    parsed=table_cells(child)
    if not parsed:return False
    cols,rows,cells=parsed
    if any(v<=0 for v in cols+rows):return False
    total_w=sum(cols);total_h=sum(rows)
    shape=slide.shapes.add_table(len(rows),len(cols),Pt(x),Pt(y),Pt(total_w),Pt(total_h))
    table=shape.table
    for i,w in enumerate(cols):table.columns[i].width=Pt(w)
    for i,h in enumerate(rows):table.rows[i].height=Pt(h)
    for idx,(style,body) in enumerate(cells):
        cell=table.cell(idx//len(cols),idx%len(cols))
        fill=style.get('fill','none')
        if fill=='none':cell.fill.background()
        else:
            color,opacity=simple_color_alpha(fill)
            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor(*color)
            if opacity<.99999:
                solid_fill=cell._tc.get_or_add_tcPr().xpath('./a:solidFill')
                srgb=solid_fill[0].xpath('./a:srgbClr') if solid_fill else []
                if srgb:
                    alpha=OxmlElement('a:alpha');alpha.set('val',str(round(opacity*100000)));srgb[0].append(alpha)
        inset=style.get('inset','')
        vals=[length(v) for v in re.findall(NUM+r'\s*pt',inset)]
        if vals:
            mx=vals[0];my=vals[1] if len(vals)>1 else mx
            cell.margin_left=cell.margin_right=Pt(mx);cell.margin_top=cell.margin_bottom=Pt(my)
        align=style.get('align','left')
        cell.text_frame.paragraphs[0].alignment=(PP_ALIGN.RIGHT if 'right' in align else PP_ALIGN.CENTER if 'center' in align else PP_ALIGN.LEFT)
        cell.vertical_anchor=MSO_ANCHOR.MIDDLE if ('horizon' in align or 'center' in align) else MSO_ANCHOR.BOTTOM if 'bottom' in align else MSO_ANCHOR.TOP
        tf=cell.text_frame;tf.clear();tf.word_wrap=True;p=tf.paragraphs[0]
        for item in text_runs(body):
            r=p.add_run();r.text=item['text'];r.font.name=item['font'];r.font.size=Pt(item['size']);r.font.bold=item['bold'];r.font.color.rgb=RGBColor(*item['color'])
        stroke=style.get('stroke','')
        for side,tag in (('left','lnL'),('right','lnR'),('top','lnT'),('bottom','lnB')):
            found=re.search(side+r'\s*:\s*(none|'+NUM+r'pt\s*\+\s*rgb\("#[\da-fA-F]{6}"\))',stroke)
            if not found or found.group(1)=='none':continue
            width=length(found.group(1),.5);col=re.search(r'#[\da-fA-F]{6}',found.group(1))
            edge=OxmlElement('a:'+tag);edge.set('w',str(round(width*12700)))
            sf=OxmlElement('a:solidFill');sc=OxmlElement('a:srgbClr');sc.set('val',col.group(0)[1:].upper());sf.append(sc);edge.append(sf)
            cell._tc.get_or_add_tcPr().append(edge)
    if abs(rot)>0.001:shape._element.xfrm.set('rot',str(round(rot*60000)))
    audit['tables']+=1
    audit['items'].append({'slide':slide_num,'type':'table','x':round(x,2),'y':round(y,2),'w':round(total_w,2),'h':round(total_h,2),'rows':len(rows),'columns':len(cols),'rotation':round(rot,2)})
    audit['text_content'].extend(''.join(r['text'] for r in text_runs(body)) for _,body in cells)
    return True


def parse_root(raw):
    head=kwargs(raw)
    m=re.search(r'\bchildren\s*:\s*\[',raw)
    if not m:return head,''
    content,_=balanced(raw,m.end()-1,'[',']')
    return head,content


def export_one(content,slide,base,cache,audit,n,ox=0.0,oy=0.0,angle=0):
    # Top-level Place calls are consumed as entire nodes, so nested children are
    # not processed again. This preserves local group order and z-order.
    for raw,_,_ in find_calls(content,'Place'):
        a=kwargs(raw);x=ox+length(a.get('x'))+length(a.get('dx'));y=oy+length(a.get('y'))+length(a.get('dy'))
        rot=angle+length(a.get('orientation'));child=a.get('child','')
        group=re.match(r'GroupContent\(',child)
        if group:
            gs=kwargs(balanced(child,group.end()-1)[0]);ch=gs.get('children','')
            if ch.startswith('['):ch=balanced(ch,0,'[',']')[0]
            export_one(ch,slide,base,cache,audit,n,ox=x,oy=y,angle=rot)
            audit['groups_flattened']+=1
            continue
        if add_table(slide,child,x,y,rot,audit,n):continue
        if add_image(slide,child,base,x,y,rot,cache,audit,n):continue
        if add_text(slide,child,x,y,rot,audit,n):continue
        if add_native(slide,child,x,y,rot,audit,n):continue
        show_unsupported(audit,'child_construct',child,n)


def clear_theme_shadows(slide):
    """Disable PowerPoint theme effects on objects generated from Typst.

    python-pptx assigns auto-shapes an effect reference to theme effect style
    2 by default. Many Office themes define that style as a drop shadow, even
    when the Typst source has no shadow at all.
    """
    for shape in slide.shapes:
        style = shape._element.find(qn('p:style'))
        effect_ref = style.find(qn('a:effectRef')) if style is not None else None
        if effect_ref is not None:
            effect_ref.set('idx', '0')


def find_definitions(source, kind):
    """Read simple #let name = PptxMaster/Layout(...) declarations only."""
    pattern=re.compile(r"#let\s+([A-Za-z_][\w-]*)\s*=\s*"+re.escape(kind)+r"\s*\(")
    out=[]
    pos=0
    while True:
        match=pattern.search(source,pos)
        if not match:break
        raw,end=balanced(source,match.end()-1)
        out.append((match.group(1),raw,match.start(),end))
        pos=end
    return out


def find_macro_records(expression, kind):
    """Read a static array of named PptxPlaceholder records."""
    pattern=re.compile(r'(?<![\w-])#?'+re.escape(kind)+r'\s*\(')
    records=[];pos=0
    while True:
        match=pattern.search(expression or '',pos)
        if not match:break
        raw,end=balanced(expression,match.end()-1)
        records.append(kwargs(raw));pos=end
    return records


PLACEHOLDER_TYPES={'title':'title','centered-title':'ctrTitle','body':'body'}

def placeholder_type(record):
    value=txt_decode(record.get('type','"body"'))
    if value not in PLACEHOLDER_TYPES:
        raise RuntimeError(f'Unsupported native placeholder type {value!r}')
    return value


def placeholder_text(record):
    value=record.get('text','""')
    text=txt_decode(value)
    if text:return text
    paragraphs=find_macro_records(record.get('paragraphs','()'),'PptxPlaceholderParagraph')
    if paragraphs:return '\n'.join(txt_decode(item.get('text','""')) for item in paragraphs)
    content=record.get('content','')
    return ''.join(run['text'] for run in text_runs(content))


def merged_placeholder(slots, parent_slots, instance=None):
    """Resolve static render geometry while preserving local override fields."""
    instance=instance or {}
    slot_id=txt_decode(instance.get('id','""'))
    own=next((item for item in slots if txt_decode(item.get('id','""'))==slot_id),None)
    parent=next((item for item in parent_slots if txt_decode(item.get('id','""'))==slot_id),None)
    result={}
    for key in ('x','y','width','height','rotation','font','size','fill','background','weight','align','bullet',
                'anchor','margin-left','margin-right','margin-top','margin-bottom'):
        value=instance.get(key)
        if value is None and own is not None:value=own.get(key)
        if value is None and parent is not None:value=parent.get(key)
        if value is not None:result[key]=value
    result['id']=slot_id
    result['type']=instance.get('type',own.get('type') if own else parent.get('type') if parent else '"body"')
    result['idx']=instance.get('idx',own.get('idx') if own else parent.get('idx') if parent else '0')
    return result,own,parent


def add_native_placeholder(slide, record, resolved, *, keep_geometry=False, marker=None):
    """Create a real p:sp placeholder; omitted local transforms inherit from layout."""
    x=length(resolved.get('x'));y=length(resolved.get('y'))
    w=length(resolved.get('width'));h=length(resolved.get('height'))
    if w<=0 or h<=0:
        raise RuntimeError(f'Placeholder {txt_decode(record.get("id", "\"?\""))!r} has no resolvable layout geometry')
    shape=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h))
    if record.get('background'):
        shape.fill.solid()
        shape.fill.fore_color.rgb=RGBColor(*simple_color(record['background']))
    kind=placeholder_type(record)
    idx=int(length(record.get('idx'),0))
    shape.name=marker or f'PptxTypst {kind} placeholder {idx}'
    ph=OxmlElement('p:ph')
    ph.set('type',PLACEHOLDER_TYPES[kind]);ph.set('idx',str(idx))
    shape._element.nvSpPr.nvPr.append(ph)
    tf=shape.text_frame;tf.clear()
    body_pr=shape._element.txBody.bodyPr
    for attr in ('lIns','rIns','tIns','bIns','anchor','wrap'):
        body_pr.attrib.pop(attr,None)
    text=placeholder_text(record)
    paragraph_records=find_macro_records(record.get('paragraphs','()'),'PptxPlaceholderParagraph')
    paragraph_values=([(txt_decode(item.get('text','""')),item) for item in paragraph_records]
                      if paragraph_records else [(part,{}) for part in text.replace('\v','\n').split('\n')])
    for para_index,(part,paragraph_record) in enumerate(paragraph_values):
        para=tf.paragraphs[0] if para_index==0 else tf.add_paragraph()
        para.level=int(length(paragraph_record.get('level'),0))
        if paragraph_record.get('align'):
            alignment=txt_decode(paragraph_record['align'])
            if alignment in ('left','center','right'):
                para.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[alignment]
        if paragraph_record.get('bullet'):
            ppr=para._p.get_or_add_pPr()
            bullet=OxmlElement('a:buChar');bullet.set('char',txt_decode(paragraph_record['bullet']));ppr.append(bullet)
        run_records=find_macro_records(paragraph_record.get('runs','()'),'PptxPlaceholderRun')
        if run_records:
            for run_record in run_records:
                run=para.add_run();run.text=txt_decode(run_record.get('text','""'))
                if run_record.get('fill'):run.font.color.rgb=RGBColor(*simple_color(run_record['fill']))
                if run_record.get('font'):run.font.name=txt_decode(run_record['font'])
                if run_record.get('size'):run.font.size=Pt(length(run_record['size']))
                if run_record.get('weight'):run.font.bold=txt_decode(run_record['weight'])=='bold'
                if run_record.get('italic') is not None:run.font.italic=run_record['italic'].strip().lower()=='true'
        else:
            run=para.add_run();run.text=part
            if any(k in record for k in ('font','size','fill','weight')):
                if record.get('fill'):
                    run.font.color.rgb=RGBColor(*simple_color(record['fill']))
                if record.get('font'):run.font.name=txt_decode(record['font'])
                if record.get('size'):run.font.size=Pt(length(record['size']))
                if record.get('weight'):run.font.bold=txt_decode(record['weight'])=='bold'
    for key,attr in (('margin-left','margin_left'),('margin-right','margin_right'),
                     ('margin-top','margin_top'),('margin-bottom','margin_bottom')):
        if key in record:setattr(tf,attr,Pt(length(record[key])))
    if record.get('align'):
        alignment=txt_decode(record['align'])
        if alignment in ('left','center','right'):
            tf.paragraphs[0].alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[alignment]
    if record.get('anchor'):
        anchor=txt_decode(record['anchor'])
        tf.vertical_anchor={'top':MSO_ANCHOR.TOP,'center':MSO_ANCHOR.MIDDLE,'bottom':MSO_ANCHOR.BOTTOM}.get(anchor,MSO_ANCHOR.TOP)
    if not text and any(record.get(key) is not None for key in
                        ('align','size','weight','font','fill','bullet')):
        # Emit a local placeholder text style only when the declaration defines
        # one. An empty lvl1pPr blocks the parent master/layout text style.
        tx_body=shape._element.txBody
        lst=tx_body.find(qn('a:lstStyle'))
        if lst is None:
            lst=OxmlElement('a:lstStyle')
            tx_body.insert(1,lst)
        lvl=OxmlElement('a:lvl1pPr')
        if record.get('align'):
            alignment=txt_decode(record['align'])
            lvl.set('algn',{'left':'l','center':'ctr','right':'r','justify':'just'}.get(alignment,'l'))
        props=OxmlElement('a:defRPr')
        if record.get('size'):props.set('sz',str(round(length(record['size'])*100)))
        if record.get('weight'):props.set('b','1' if txt_decode(record['weight'])=='bold' else '0')
        if record.get('fill'):
            solid=OxmlElement('a:solidFill');color=OxmlElement('a:srgbClr')
            color.set('val','%02X%02X%02X'%simple_color(record['fill']));solid.append(color);props.append(solid)
        if record.get('font'):
            latin=OxmlElement('a:latin');latin.set('typeface',txt_decode(record['font']));props.append(latin)
        if record.get('bullet'):
            bullet=OxmlElement('a:buChar');bullet.set('char',txt_decode(record['bullet']));lvl.append(bullet)
        lvl.append(props);lst.append(lvl)
    elif not text:
        # python-pptx leaves an empty a:lstStyle after clearing a text frame;
        # remove it so the placeholder can inherit its parent's text style.
        tx_body=shape._element.txBody
        lst=tx_body.find(qn('a:lstStyle'))
        if lst is not None and len(lst)==0:
            tx_body.remove(lst)
    if resolved.get('rotation'):
        shape.rotation=length(resolved['rotation'])%360
    if not keep_geometry:
        xfrm=shape._element.spPr.find(qn('a:xfrm'))
        if xfrm is not None:shape._element.spPr.remove(xfrm)
    return shape


def background_asset(expression, base, audit, owner):
    expression=(expression or 'none').strip()
    if expression=='none':return None
    image=img_call(expression)
    if not image:
        show_unsupported(audit,'background_expression',expression,owner)
        return None
    root=base.resolve();path=(root/image['path']).resolve()
    if not path.is_relative_to(root) or not path.is_file():
        show_unsupported(audit,'missing_background_media',image['path'],owner)
        return None
    if path.suffix.lower() not in ('.png','.jpeg','.jpg','.bmp','.gif','.tif','.tiff'):
        show_unsupported(audit,'background_media_format',image['path'],owner)
        return None
    return path


def remove_spans(source, spans):
    for start,end in sorted(spans,reverse=True):
        source=source[:start]+source[end:]
    return source


def convert(main,output):
    main=Path(main).resolve();base=main.parent
    content=main.read_text(encoding='utf-8')
    content=re.sub(r'(?m)^\s*//[^\n]*$', '', content)
    master_defs=find_definitions(content,'PptxMaster')
    layout_defs=find_definitions(content,'PptxLayout')
    declaration_spans=[(v[2],v[3]) for v in master_defs+layout_defs]
    master_records=[]
    master_fills={}
    master_background_exprs={}
    master_placeholders={}
    for name,raw,start,end in master_defs:
        head,body=parse_root(raw)
        body=normalize_starter_aliases(body)
        master_records.append((name,body,start,end))
        master_fills[name]=head.get('fill','none')
        master_background_exprs[name]=head.get('background','none')
        master_placeholders[name]=find_macro_records(head.get('placeholders','[]'),'PptxPlaceholderSlot')
    layout_records={}
    layout_background_exprs={}
    layout_placeholders={}
    for name,raw,_,_ in layout_defs:
        head,body=parse_root(raw)
        body=normalize_starter_aliases(body)
        master_name=head.get('master','').strip()
        if not re.fullmatch(r'[A-Za-z_][\w-]*',master_name):
            raise RuntimeError(f'Layout {name!r} must reference a statically named PptxMaster')
        layout_records[name]=(master_name,body,head.get('show-master','true').strip().lower() not in ('false','0'),head.get('fill','none'))
        layout_background_exprs[name]=head.get('background','none')
        layout_placeholders[name]=find_macro_records(head.get('placeholders','[]'),'PptxPlaceholderSlot')
    content=remove_spans(content,declaration_spans)
    content=normalize_starter_aliases(content)
    slide_calls=list(find_calls(content,'Root'))
    if not slide_calls:raise RuntimeError('No explicit #Root/#PptxSlide slides; expected the supported Root/Place source subset')
    first=kwargs(slide_calls[0][0]);sw=length(first.get('width'),960);sh=length(first.get('height'),540)
    if sw<=0 or sh<=0:raise RuntimeError('Invalid logical slide size')
    prs=Presentation();prs.slide_width=Pt(sw);prs.slide_height=Pt(sh)
    cache=output.parent/'raster_cache';cache.mkdir(parents=True,exist_ok=True)
    audit={'scope':'Place-based generated Typst subset, not general Typst','slides':len(slide_calls),'slide_size_pt':[sw,sh],'pictures':0,'text_boxes':0,'native_shapes':0,'tables':0,'groups_flattened':0,'cropped_instance_fallbacks':0,'svg_assets_as_png':0,'text_content':[],'unsupported':[],'source_media':set(),'items':[],'masters':len(master_records),'layouts':len(layout_records)}
    master_backgrounds={name:path for name,expr in master_background_exprs.items()
                       if (path:=background_asset(expr,base,audit,f'master:{name}')) is not None}
    layout_backgrounds={name:path for name,expr in layout_background_exprs.items()
                        if (path:=background_asset(expr,base,audit,f'layout:{name}')) is not None}
    slide_assignments={}
    slide_visibility={}
    for n,(raw,_,_) in enumerate(slide_calls,1):
        head,body=parse_root(raw)
        w=length(head.get('width'),sw);h=length(head.get('height'),sh)
        if (w,h)!=(sw,sh):show_unsupported(audit,'mixed_slide_dimensions',str((w,h)),n)
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        fill=head.get('fill','none').strip()
        if fill!='none':
            bg=slide.background.fill;bg.solid();bg.fore_color.rgb=RGBColor(*simple_color(fill,(255,255,255)))
            if 'gradient.' in fill:show_unsupported(audit,'slide_background_gradient',fill,n)
        export_one(body,slide,base,cache,audit,n)
        slide_placeholder_records=find_macro_records(head.get('placeholders','[]'),'PptxPlaceholder')
        layout_name=head.get('layout','').strip()
        parent_slots=master_placeholders.get(layout_records.get(layout_name,('', '', False, ''))[0],[])
        local_slots=layout_placeholders.get(layout_name,[])
        for item in slide_placeholder_records:
            resolved,local,parent=merged_placeholder(local_slots,parent_slots,item)
            local_geometry=any(key in item for key in ('x','y','width','height','rotation'))
            shape=add_native_placeholder(
                slide,item,resolved,keep_geometry=local_geometry,
                marker='PptxTypst placeholder '+txt_decode(item.get('id','""')),
            )
            audit['items'].append({'slide':n,'type':'placeholder','placeholder_type':placeholder_type(item),
                                   'idx':int(length(item.get('idx'),0)),'id':txt_decode(item.get('id','""')),
                                   'inherited_from_layout':local is not None,'inherited_from_master':local is None and parent is not None,
                                   'geometry_override':local_geometry,'text':placeholder_text(item)})
            audit['text_content'].append(placeholder_text(item))
            audit.setdefault('placeholders',[]).append({'slide':n,'id':txt_decode(item.get('id','""')),
                'type':placeholder_type(item),'idx':int(length(item.get('idx'),0)),
                'layout':layout_name or None,'geometry_override':local_geometry,'native':True})
        clear_theme_shadows(slide)
        if layout_name:
            if layout_name not in layout_records:
                raise RuntimeError(f'Slide {n} references unknown PptxLayout {layout_name!r}')
            slide_assignments[str(slide.part.partname)]=layout_name
            slide_visibility[str(slide.part.partname)]=head.get('show-master','true').strip().lower() not in ('false','0')
        audit.setdefault('per_slide',[]).append({'slide':n,'objects':len(slide.shapes),'layout':layout_name or None,'master':layout_records[layout_name][0] if layout_name else None})

    # Stage each owner's supported native objects on a temporary slide. The
    # package writer moves these shape trees and their image rels into real parts.
    stage_parts={};master_parts={};layout_parts={};master_object_counts={};layout_object_counts={}
    for index,(name,body,_,_) in enumerate(master_records,1):
        stage=prs.slides.add_slide(prs.slide_layouts[6]);stage_parts[name]=str(stage.part.partname);master_parts[name]=str(stage.part.partname)
        export_one(body,stage,base,cache,audit,f'master:{name}')
        for item in master_placeholders.get(name,[]):
            resolved,_,_=merged_placeholder([],[],item)
            add_native_placeholder(stage,item,resolved,keep_geometry=any(k in item for k in ('x','y','width','height','rotation')))
        master_object_counts[name]=len(stage.shapes)
        clear_theme_shadows(stage)
    layout_visibility={};layout_fills={}
    for index,(name,(master_name,body,show_master,fill)) in enumerate(layout_records.items(),1):
        if not any(defn[0]==master_name for defn in master_records):
            raise RuntimeError(f'Layout {name!r} references undefined master {master_name!r}')
        stage=prs.slides.add_slide(prs.slide_layouts[6]);stage_parts[name]=str(stage.part.partname);layout_parts[name]=(master_name,str(stage.part.partname));layout_visibility[name]=show_master;layout_fills[name]=fill
        export_one(body,stage,base,cache,audit,f'layout:{name}')
        for item in layout_placeholders.get(name,[]):
            resolved,_,_=merged_placeholder(layout_placeholders.get(name,[]),master_placeholders.get(master_name,[]),item)
            add_native_placeholder(stage,item,resolved,keep_geometry=any(k in item for k in ('x','y','width','height','rotation')))
        layout_object_counts[name]=len(stage.shapes)
        clear_theme_shadows(stage)
    for name,fill in master_fills.items():
        if 'gradient.' in fill:show_unsupported(audit,'master_background_gradient',fill,f'master:{name}')
    for name,(_,_,_,fill) in layout_records.items():
        if 'gradient.' in fill:show_unsupported(audit,'layout_background_gradient',fill,f'layout:{name}')
    output.parent.mkdir(parents=True,exist_ok=True);prs.save(output)
    if master_parts and layout_parts and slide_assignments:
        from .ooxml_hierarchy import install_hierarchy
        install_hierarchy(output,masters=master_parts,layouts=layout_parts,slide_layouts=slide_assignments,stage_parts=stage_parts,layout_show_master=layout_visibility,slide_show_master=slide_visibility,master_fills=master_fills,layout_fills=layout_fills,master_backgrounds=master_backgrounds,layout_backgrounds=layout_backgrounds)
    audit['unique_source_media']=len(audit['source_media']);audit['source_media']=sorted(audit['source_media'])
    audit['hierarchy_written']=bool(master_parts and layout_parts and slide_assignments)
    audit['hierarchy']={'masters':[{'id':name,'shared_objects':master_object_counts.get(name,0)} for name in master_parts], 'layouts':[{'id':name,'master':layout_parts[name][0],'shared_objects':layout_object_counts.get(name,0),'show_master':layout_visibility.get(name,True),'fill':layout_fills.get(name,'none')} for name in layout_parts], 'slides':[{'part':part,'layout':layout_name} for part,layout_name in slide_assignments.items()]}
    (output.parent/'conversion_report.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
    return audit


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('main_typ');p.add_argument('--output',default='converted.pptx');a=p.parse_args()
    report=convert(a.main_typ,Path(a.output).resolve())
    print(json.dumps({k:v for k,v in report.items() if k in ('slides','pictures','text_boxes','native_shapes','tables','groups_flattened','unique_source_media','cropped_instance_fallbacks','svg_assets_as_png')},indent=2))
if __name__=='__main__':main()
