from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from PIL import Image,ImageDraw,ImageFont
from pathlib import Path
from lxml import etree
P=Path(__file__).parent; P.mkdir(exist_ok=True)
im=Image.new('RGB',(1600,900),(220,233,241));d=ImageDraw.Draw(im)
colors=['#E75A39','#F0A13A','#FFDF50','#51B9BC','#3D75A9','#906AC0','#74AA3D']
for x in range(0,1600,100):
 d.rectangle((x,0,x+99,899),fill=colors[x//100%len(colors)])
 d.text((x+12,110),str(x//100),fill=(10,15,20))
for y in range(0,900,100):
 d.line([(0,y),(1600,y)],fill=(20,20,20),width=6)
d.text((520,500),'CROP AND MASK TEST', fill=(0,0,0),stroke_fill=(255,255,255),stroke_width=3)
im.save(P/'source_grid.png')
prs=Presentation();prs.slide_width=Inches(13.3333);prs.slide_height=Inches(7.5)
C=[
('plain ellipse', 'ellipse',0,0,0, None,None),
('crop ellipse','ellipse',0,0,0,{'l':16000,'t':11000,'r':18000,'b':9000},None),
('crop & fill ellipse','ellipse',0,0,0,{'l':16000,'t':11000,'r':18000,'b':9000},{'l':-6000,'t':-5000,'r':-12000,'b':0}),
('rotated ellipse','ellipse',32,0,0,{'l':15000,'t':13000,'r':21000,'b':8000},None),
('horizontal flip ellipse','ellipse',0,1,0,{'l':15000,'t':13000,'r':21000,'b':8000},None),
('vertical flip ellipse','ellipse',0,0,1,{'l':9000,'t':23000,'r':14000,'b':7000},None),
('rotated + flipped ellipse','ellipse',38,1,0,{'l':9000,'t':23000,'r':14000,'b':7000},None),
('roundRect crop','roundRect',0,0,0,{'l':12000,'t':12000,'r':15000,'b':6000},None),
('roundRect rotated','roundRect',28,0,0,{'l':10000,'t':22000,'r':8000,'b':5000},None),
('diamond crop','diamond',0,0,0,{'l':10000,'t':14000,'r':20000,'b':9000},None),
('complex freeform crop','custom',0,0,0,{'l':14000,'t':6000,'r':6000,'b':16000},None),
]
for i,(name,shape,angle,flip_h,flip_v,crop,fill) in enumerate(C,1):
 sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor(248,248,244)
 tb=sl.shapes.add_textbox(Inches(.6),Inches(.35), Inches(10),Inches(.7));tb.text=f'{i:02d} {name}'
 for pa in tb.text_frame.paragraphs:
  for run in pa.runs:run.font.size=Pt(32)
 sh=sl.shapes.add_picture(str(P/'source_grid.png'),Inches(3.1),Inches(1.6),Inches(7.1),Inches(4.1))
 sh.rotation=angle
 xf=sh._element.xpath('./p:spPr/a:xfrm')[0]
 if flip_h:xf.set('flipH','1')
 if flip_v:xf.set('flipV','1')
 g=sh._element.xpath('./p:spPr/a:prstGeom')[0]
 if shape=='custom':
  sp=g.getparent();idx=sp.index(g);sp.remove(g)
  cg=OxmlElement('a:custGeom');cg.append(OxmlElement('a:avLst'));pl=OxmlElement('a:pathLst')
  p=OxmlElement('a:path');p.set('w','100000');p.set('h','100000')
  def command(cmd,coords):
   elem=OxmlElement('a:'+cmd)
   for x,y in coords:
    t=OxmlElement('a:pt');t.set('x',str(x));t.set('y',str(y));elem.append(t)
   p.append(elem)
  command('moveTo',[(50000,0)]);command('cubicBezTo',[(98000,0),(110000,83000),(50000,100000)]);command('cubicBezTo',[(-12000,82000),(1200,0),(50000,0)]);p.append(OxmlElement('a:close'))
  pl.append(p);cg.append(pl);sp.insert(idx,cg)
 else:g.set('prst',shape)
 if crop:
  blip=sh._element.xpath('./p:blipFill')[0]
  src=OxmlElement('a:srcRect')
  for k,v in crop.items():src.set(k,str(v))
  blip.insert(1,src)
 if fill:
  rect=sh._element.xpath('./p:blipFill/a:stretch')[0]
  fr=rect.find(qn('a:fillRect'))
  if fr is None:fr=OxmlElement('a:fillRect');rect.append(fr)
  for k,v in fill.items():fr.set(k,str(v))
prs.save(P/'transform_cases.pptx')
print('saved',len(C))
