#!/usr/bin/env python3
"""Produce focused structural/rendered text-placement diagnostics for PPTX samples."""
from __future__ import annotations

import argparse
from collections import Counter
from difflib import SequenceMatcher
import json
import math
import os
from pathlib import Path
import subprocess
import sys
import tempfile

import numpy as np
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

try:
    from .run_text_fidelity_benchmark import (
        ROOT, export_powerpoint_pdf, get_typst_font_names, match_text_blocks,
        match_text_blocks_spatial, page_image, pdf_text_blocks, pptx_stats, WORD_RE,
    )
except ImportError:
    from run_text_fidelity_benchmark import (
        ROOT, export_powerpoint_pdf, get_typst_font_names, match_text_blocks,
        match_text_blocks_spatial, page_image, pdf_text_blocks, pptx_stats, WORD_RE,
    )


EMU = 12700.0


def trans(x, y):
    return np.array([[1., 0., x], [0., 1., y], [0., 0., 1.]])


def scale(x, y):
    return np.array([[x, 0., 0.], [0., y, 0.], [0., 0., 1.]])


def rotate(degrees):
    angle = math.radians(degrees)
    c, s = math.cos(angle), math.sin(angle)
    return np.array([[c, -s, 0.], [s, c, 0.], [0., 0., 1.]])


def source_shape_matrix(shape):
    element = shape._element
    xfrm = element.xpath("./p:xfrm") if element.tag.endswith("graphicFrame") else element.xpath("./p:spPr/a:xfrm")
    if xfrm:
        xf = xfrm[0]
        angle = int(xf.get("rot", "0")) / 60000
        flip_h = xf.get("flipH") in ("1", "true")
        flip_v = xf.get("flipV") in ("1", "true")
    else:
        angle, flip_h, flip_v = 0., False, False
    x, y = shape.left / EMU, shape.top / EMU
    w, h = shape.width / EMU, shape.height / EMU
    return (trans(x + w / 2, y + h / 2) @ rotate(angle)
            @ scale(-1 if flip_h else 1, -1 if flip_v else 1)
            @ scale(w, h) @ trans(-.5, -.5))


def source_group_matrix(shape):
    xf = shape._element.grpSpPr.xfrm
    x, y = xf.off.x / EMU, xf.off.y / EMU
    w, h = xf.ext.cx / EMU, xf.ext.cy / EMU
    child_x, child_y = xf.chOff.x / EMU, xf.chOff.y / EMU
    child_w, child_h = xf.chExt.cx / EMU, xf.chExt.cy / EMU
    angle = int(xf.get("rot", "0")) / 60000
    flip_h = xf.get("flipH") in ("1", "true")
    flip_v = xf.get("flipV") in ("1", "true")
    return (trans(x + w / 2, y + h / 2) @ rotate(angle)
            @ scale(-1 if flip_h else 1, -1 if flip_v else 1)
            @ trans(-w / 2, -h / 2)
            @ scale(w / child_w if child_w else 1, h / child_h if child_h else 1)
            @ trans(-child_x, -child_y))


def source_placement(shape, parent):
    matrix = parent @ source_shape_matrix(shape)
    center = matrix @ np.array([.5, .5, 1.])
    vx, vy = matrix[:2, :2] @ np.array([1., 0.]), matrix[:2, :2] @ np.array([0., 1.])
    width, height = float(np.linalg.norm(vx)), float(np.linalg.norm(vy))
    return {
        "x": float(center[0] - width / 2), "y": float(center[1] - height / 2),
        "w": width, "h": height, "rot": math.degrees(math.atan2(vx[1], vx[0])),
    }


def source_text_shapes(deck):
    found = {}

    def walk(shapes, slide_number, parent, grouped):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes, slide_number, parent @ source_group_matrix(shape), True)
                continue
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                found[(slide_number, shape.shape_id)] = {
                    "text": shape.text,
                    "geometry": source_placement(shape, parent),
                    "grouped": grouped,
                }

    for slide_number, slide in enumerate(deck.slides, 1):
        walk(slide.shapes, slide_number, np.identity(3), False)
    return found


def geometry_metrics(deck, manifest):
    source = source_text_shapes(deck)
    output = {
        (slide["slide"], item["source_shape_id"]): item
        for slide in manifest["slides"] for item in slide["placements"]
        if item.get("type") == "text"
    }
    diffs = {key: [] for key in ("x", "y", "w", "h", "rot")}
    content_matches = missing = grouped = 0
    for key, original in source.items():
        converted = output.get(key)
        if converted is None:
            missing += 1
            continue
        if original["grouped"]:
            grouped += 1
        for field in diffs:
            delta = abs(original["geometry"][field] - float(converted[field]))
            if field == "rot":
                delta = min(delta % 360, 360 - (delta % 360))
            diffs[field].append(delta)
        content_matches += original["text"][:200] == converted.get("text", "")
    return {
        "source_text_shapes": len(source), "output_text_shapes": len(output),
        "matched_by_slide_and_shape_id": len(source) - missing,
        "missing_output_text_shapes": missing,
        "output_only_text_shapes": len(output.keys() - source.keys()),
        "grouped_text_shapes": grouped,
        "text_preview_exact_matches": content_matches,
        "geometry_abs_delta_pt": {
            key: {"mean": float(np.mean(values)) if values else None,
                 "max": float(np.max(values)) if values else None}
            for key, values in diffs.items()
        },
    }


def block_frame_evidence(block, shape):
    x, y, width, height = (float(shape["geometry"][key]) for key in ("x", "y", "w", "h"))
    x0, y0, x1, y1 = block["bbox"]
    intersection = max(0., min(x1, x + width) - max(x0, x)) * max(0., min(y1, y + height) - max(y0, y))
    area = max(1., (x1 - x0) * (y1 - y0))
    overlap = intersection / area
    cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
    dx = max(x - cx, 0., cx - (x + width))
    dy = max(y - cy, 0., cy - (y + height))
    return overlap, math.hypot(dx, dy)


def assign_pdf_blocks_to_shapes(blocks, shapes, *, minimum_text_evidence=0.58,
                                maximum_frame_distance_pt=72., ambiguity_distance_pt=3.):
    """Attach rendered PDF blocks to known PPTX/Typst text objects by text and frame."""
    assigned = {shape_id: [] for shape_id in shapes}
    ambiguous = unmatched = 0
    pieces = [span for block in blocks for span in block.get("spans", [])]
    if not pieces:
        pieces = blocks
    for block in pieces:
        block_words = Counter(WORD_RE.findall(block["text"].casefold()))
        if not block_words:
            unmatched += 1
            continue
        if len(block_words) == 1:
            only_word = next(iter(block_words))
            # Page numbers and single-character bullets routinely occur in long
            # text frames by coincidence; they are not reliable object identity.
            if only_word.isdecimal() or len(only_word) <= 1:
                unmatched += 1
                continue
        candidates = []
        for shape_id, shape in shapes.items():
            shape_words = Counter(WORD_RE.findall(shape["text"].casefold()))
            if not shape_words:
                continue
            common = sum((block_words & shape_words).values())
            coverage = common / sum(block_words.values())
            sequence = SequenceMatcher(None,
                                       " ".join(block_words.elements()),
                                       " ".join(shape_words.elements()),
                                       autojunk=False).ratio()
            evidence = max(coverage, sequence)
            if evidence < minimum_text_evidence:
                continue
            overlap, distance = block_frame_evidence(block, shape)
            if distance > maximum_frame_distance_pt:
                continue
            rank = evidence + .08 * overlap - .001 * distance
            candidates.append((rank, evidence, overlap, distance, shape_id))
        candidates.sort(reverse=True)
        if not candidates:
            unmatched += 1
            continue
        if len(candidates) > 1:
            best, second = candidates[:2]
            if (abs(best[1] - second[1]) <= .02
                    and abs(best[2] - second[2]) <= .08
                    and abs(best[3] - second[3]) <= ambiguity_distance_pt):
                ambiguous += 1
                continue
        assigned[candidates[0][4]].append(block)
    return assigned, ambiguous, unmatched


def combine_shape_blocks(blocks):
    if not blocks:
        return None
    first = min(blocks, key=lambda block: (block["bbox"][1], block["bbox"][0]))
    return {
        "start": first.get("start", (first["bbox"][0], first["bbox"][1])),
        "baseline": first.get("baseline", first.get("start", first["bbox"][:2])),
        "bbox": (min(b["bbox"][0] for b in blocks), min(b["bbox"][1] for b in blocks),
                 max(b["bbox"][2] for b in blocks), max(b["bbox"][3] for b in blocks)),
        "baselines": sorted({round(float(block.get("baseline", block.get("start", block["bbox"][:2]))[1]), 1)
                              for block in blocks}),
    }


def object_linked_render_metrics(source_blocks, output_blocks, source_shapes, output_shapes):
    source_assigned, source_ambiguous, source_unmatched = assign_pdf_blocks_to_shapes(source_blocks, source_shapes)
    output_assigned, output_ambiguous, output_unmatched = assign_pdf_blocks_to_shapes(output_blocks, output_shapes)
    source_positions = {key: combine_shape_blocks(blocks) for key, blocks in source_assigned.items()}
    output_positions = {key: combine_shape_blocks(blocks) for key, blocks in output_assigned.items()}
    source_positions = {key: value for key, value in source_positions.items() if value is not None}
    output_positions = {key: value for key, value in output_positions.items() if value is not None}
    common = source_positions.keys() & output_positions.keys()
    displacements, baseline_deltas, bbox_deltas, line_spacing_deltas = [], [], [], []
    source_overflows, typst_overflows = [], []
    samples = []
    for shape_id in sorted(common):
        source, output = source_positions[shape_id], output_positions[shape_id]
        dx = source["start"][0] - output["start"][0]
        dy = source["start"][1] - output["start"][1]
        distance = math.hypot(dx, dy)
        displacements.append(distance)
        baseline_deltas.append(abs(source["baseline"][1] - output["baseline"][1]))
        bbox_deltas.append(float(np.mean(np.abs(np.array(source["bbox"]) - np.array(output["bbox"])))) )
        if len(source["baselines"]) == len(output["baselines"]) and len(source["baselines"]) > 1:
            line_spacing_deltas.extend(abs(
                (source["baselines"][i] - source["baselines"][i - 1])
                - (output["baselines"][i] - output["baselines"][i - 1])
            ) for i in range(1, len(source["baselines"])))
        shape = source_shapes[shape_id]
        x, y, width, height, angle = (float(shape["geometry"][key]) for key in ("x", "y", "w", "h", "rot"))
        cx, cy = x + width / 2, y + height / 2
        corners = []
        for px, py in ((x, y), (x + width, y), (x + width, y + height), (x, y + height)):
            matrix = rotate(angle) @ np.array([px - cx, py - cy, 1.])
            corners.append((matrix[0] + cx, matrix[1] + cy))
        frame = (min(p[0] for p in corners), min(p[1] for p in corners),
                 max(p[0] for p in corners), max(p[1] for p in corners))
        source_overflow = max(0., frame[0] - source["bbox"][0], frame[1] - source["bbox"][1],
                              source["bbox"][2] - frame[2], source["bbox"][3] - frame[3])
        output_overflow = max(0., frame[0] - output["bbox"][0], frame[1] - output["bbox"][1],
                              output["bbox"][2] - frame[2], output["bbox"][3] - frame[3])
        source_overflows.append(source_overflow)
        typst_overflows.append(output_overflow)
        src_pitch = np.diff(source["baselines"])
        out_pitch = np.diff(output["baselines"])
        samples.append({"shape_id": shape_id, "distance_pt": round(distance, 2),
                        "dx_pt": round(dx, 2), "dy_pt": round(dy, 2),
                        "text": shape["text"][:100],
                        "source_glyph_bbox": [round(v, 1) for v in source["bbox"]],
                        "typst_glyph_bbox": [round(v, 1) for v in output["bbox"]],
                        "frame_bbox": [round(v, 1) for v in frame],
                        "source_overflow_pt": round(source_overflow, 2),
                        "typst_overflow_pt": round(output_overflow, 2),
                        "source_line_count": len(source["baselines"]),
                        "typst_line_count": len(output["baselines"]),
                        "source_mean_baseline_pitch_pt": float(np.mean(src_pitch)) if len(src_pitch) else None,
                        "typst_mean_baseline_pitch_pt": float(np.mean(out_pitch)) if len(out_pitch) else None})
    return {
        "source_shape_count": len(source_shapes), "typst_shape_count": len(output_shapes),
        "objects_with_glyphs_in_both_pdfs": len(common),
        "objects_missing_source_glyphs": len(source_shapes) - len(source_positions),
        "objects_missing_typst_glyphs": len(output_shapes) - len(output_positions),
        "ambiguous_source_pdf_spans": source_ambiguous,
        "ambiguous_typst_pdf_spans": output_ambiguous,
        "unassigned_source_pdf_spans": source_unmatched,
        "unassigned_typst_pdf_spans": output_unmatched,
        "mean_first_glyph_displacement_pt": float(np.mean(displacements)) if displacements else None,
        "median_first_glyph_displacement_pt": float(np.median(displacements)) if displacements else None,
        "within_12pt_share": float(np.mean(np.array(displacements) <= 12)) if displacements else None,
        "mean_first_glyph_baseline_delta_pt": float(np.mean(baseline_deltas)) if baseline_deltas else None,
        "mean_text_bbox_edge_delta_pt": float(np.mean(bbox_deltas)) if bbox_deltas else None,
        "mean_line_spacing_delta_pt_when_line_count_matches": float(np.mean(line_spacing_deltas)) if line_spacing_deltas else None,
        "objects_overflowing_frame_by_more_than_2pt": {
            "powerpoint": int(sum(value > 2 for value in source_overflows)),
            "typst": int(sum(value > 2 for value in typst_overflows)),
        },
        "mean_overflow_beyond_frame_pt": {
            "powerpoint": float(np.mean(source_overflows)) if source_overflows else None,
            "typst": float(np.mean(typst_overflows)) if typst_overflows else None,
        },
        "largest_displacements": sorted(samples, key=lambda item: item["distance_pt"], reverse=True)[:8],
        "_displacements": displacements,
        "_baseline_deltas": baseline_deltas,
        "_bbox_deltas": bbox_deltas,
        "_line_spacing_deltas": line_spacing_deltas,
        "_source_overflows": source_overflows,
        "_typst_overflows": typst_overflows,
    }


def collect_render_metrics(source_pdf, typst_pdf, supported_text_slides, source_shapes, output_placements):
    old_errors, spatial_matches = [], []
    source_blocks_total = output_blocks_total = 0
    ambiguous_source = ambiguous_output = unmatched_source = unmatched_output = 0
    wrap_equal = wrap_total = 0
    clipped_source = clipped_output = 0
    slides = []
    object_displacements = []
    object_baselines = []
    object_bboxes = []
    object_line_spacing = []
    object_source_overflows = []
    object_typst_overflows = []
    object_ambiguous_source = object_ambiguous_output = 0
    object_unassigned_source = object_unassigned_output = 0
    object_missing_source = object_missing_output = 0
    object_matched_count = object_source_count = object_output_count = 0
    for page_index in sorted(supported_text_slides):
        if page_index >= min(len(source_pdf), len(typst_pdf)):
            continue
        source_page, output_page = source_pdf[page_index], typst_pdf[page_index]
        source_blocks, output_blocks = pdf_text_blocks(source_page), pdf_text_blocks(output_page)
        source_blocks_total += len(source_blocks)
        output_blocks_total += len(output_blocks)
        old_pairs = match_text_blocks(source_blocks, output_blocks)
        old_errors.extend(math.dist(a["start"], b["start"]) for a, b in old_pairs)
        result = match_text_blocks_spatial(source_blocks, output_blocks)
        matches = result["matches"]
        spatial_matches.extend(matches)
        ambiguous_source += result["ambiguous_source_count"]
        ambiguous_output += result["ambiguous_output_count"]
        unmatched_source += result["unmatched_source_count"]
        unmatched_output += result["unmatched_output_count"]
        for match in matches:
            wrap_total += 1
            wrap_equal += match["source"]["lines"] == match["output"]["lines"]
        bounds = (0., 0., source_page.rect.width, source_page.rect.height)
        clipped_source += sum(
            b[0] < bounds[0] or b[1] < bounds[1] or b[2] > bounds[2] or b[3] > bounds[3]
            for block in source_blocks for b in block["boxes"]
        )
        clipped_output += sum(
            b[0] < bounds[0] or b[1] < bounds[1] or b[2] > bounds[2] or b[3] > bounds[3]
            for block in output_blocks for b in block["boxes"]
        )
        errors = [match["distance"] for match in matches]
        source_objects = {str(shape_id): item for (slide_no, shape_id), item in source_shapes.items()
                          if slide_no == page_index + 1}
        output_objects = {}
        for item in output_placements.get(page_index + 1, []):
            source_shape = source_objects.get(str(item["source_shape_id"]))
            # Compare each source object only with its converter object. Extra
            # text placements (for example inherited placeholders) are counted
            # structurally, but cannot be allowed to steal a PDF block match.
            if source_shape is None:
                continue
            output_objects[str(item["source_shape_id"])] = {
                "text": source_shape["text"],
                "geometry": {key: float(item[key]) for key in ("x", "y", "w", "h", "rot")},
                "grouped": bool(item.get("group_path")),
            }
        object_result = object_linked_render_metrics(source_blocks, output_blocks, source_objects, output_objects)
        object_displacements.extend(object_result.pop("_displacements"))
        object_baselines.extend(object_result.pop("_baseline_deltas"))
        object_bboxes.extend(object_result.pop("_bbox_deltas"))
        object_line_spacing.extend(object_result.pop("_line_spacing_deltas"))
        object_source_overflows.extend(object_result.pop("_source_overflows"))
        object_typst_overflows.extend(object_result.pop("_typst_overflows"))
        object_ambiguous_source += object_result["ambiguous_source_pdf_spans"]
        object_ambiguous_output += object_result["ambiguous_typst_pdf_spans"]
        object_unassigned_source += object_result["unassigned_source_pdf_spans"]
        object_unassigned_output += object_result["unassigned_typst_pdf_spans"]
        object_missing_source += object_result["objects_missing_source_glyphs"]
        object_missing_output += object_result["objects_missing_typst_glyphs"]
        object_matched_count += object_result["objects_with_glyphs_in_both_pdfs"]
        object_source_count += object_result["source_shape_count"]
        object_output_count += object_result["typst_shape_count"]
        for private_key in ("_displacements", "_baseline_deltas", "_bbox_deltas", "_line_spacing_deltas",
                            "_source_overflows", "_typst_overflows"):
            object_result.pop(private_key, None)
        slides.append({"slide": page_index + 1, "source_blocks": len(source_blocks),
                       "output_blocks": len(output_blocks), "spatial_matches": len(matches),
                       "mean_start_displacement_pt": float(np.mean(errors)) if errors else None,
                       "object_linked_position": object_result})
    errors = [match["distance"] for match in spatial_matches]
    baseline_errors = [abs(match["source"]["baseline"][1] - match["output"]["baseline"][1])
                       for match in spatial_matches]
    bbox_errors = [float(np.mean(np.abs(np.array(match["source"]["bbox"])
                                         - np.array(match["output"]["bbox"]))))
                   for match in spatial_matches]
    line_spacing_errors = []
    for match in spatial_matches:
        a, b = match["source"]["baselines"], match["output"]["baselines"]
        if len(a) == len(b) and len(a) > 1:
            line_spacing_errors.extend(abs((a[i] - a[i - 1]) - (b[i] - b[i - 1]))
                                       for i in range(1, len(a)))
    return {
        "source_pdf_text_blocks": source_blocks_total,
        "typst_pdf_text_blocks": output_blocks_total,
        "historical_text_only_greedy_match": {
            "matched": len(old_errors),
            "mean_start_displacement_pt": float(np.mean(old_errors)) if old_errors else None,
            "within_12pt_share": float(np.mean(np.array(old_errors) <= 12)) if old_errors else None,
        },
        "spatially_constrained_match": {
            "matched": len(errors),
            "ambiguous_source_blocks": ambiguous_source,
            "ambiguous_output_blocks": ambiguous_output,
            "unmatched_source_blocks": unmatched_source,
            "unmatched_output_blocks": unmatched_output,
            "mean_first_glyph_displacement_pt": float(np.mean(errors)) if errors else None,
            "median_first_glyph_displacement_pt": float(np.median(errors)) if errors else None,
            "within_12pt_share": float(np.mean(np.array(errors) <= 12)) if errors else None,
            "mean_first_glyph_baseline_delta_pt": float(np.mean(baseline_errors)) if baseline_errors else None,
            "mean_text_bbox_edge_delta_pt": float(np.mean(bbox_errors)) if bbox_errors else None,
            "mean_line_spacing_delta_pt_when_wrap_count_matches": float(np.mean(line_spacing_errors)) if line_spacing_errors else None,
            "wrap_line_count_match": wrap_equal / wrap_total if wrap_total else None,
            "largest_displacements": [
                {"distance_pt": round(match["distance"], 2),
                 "dx_pt": round(match["dx"], 2), "dy_pt": round(match["dy"], 2),
                 "similarity": round(match["similarity"], 3),
                 "source_text": match["source"]["text"][:110],
                 "typst_text": match["output"]["text"][:110]}
                for match in sorted(spatial_matches, key=lambda item: item["distance"], reverse=True)[:8]
            ],
        },
        "object_linked_rendered_position": {
            "source_text_objects": object_source_count,
            "typst_text_objects": object_output_count,
            "objects_with_glyphs_in_both_pdfs": object_matched_count,
            "objects_missing_source_glyphs": object_missing_source,
            "objects_missing_typst_glyphs": object_missing_output,
            "ambiguous_source_pdf_spans": object_ambiguous_source,
            "ambiguous_typst_pdf_spans": object_ambiguous_output,
            "unassigned_source_pdf_spans": object_unassigned_source,
            "unassigned_typst_pdf_spans": object_unassigned_output,
            "mean_first_glyph_displacement_pt": float(np.mean(object_displacements)) if object_displacements else None,
            "median_first_glyph_displacement_pt": float(np.median(object_displacements)) if object_displacements else None,
            "within_12pt_share": float(np.mean(np.array(object_displacements) <= 12)) if object_displacements else None,
            "mean_first_glyph_baseline_delta_pt": float(np.mean(object_baselines)) if object_baselines else None,
            "mean_text_bbox_edge_delta_pt": float(np.mean(object_bboxes)) if object_bboxes else None,
            "mean_line_spacing_delta_pt_when_line_count_matches": float(np.mean(object_line_spacing)) if object_line_spacing else None,
            "objects_overflowing_frame_by_more_than_2pt": {
                "powerpoint": int(sum(value > 2 for value in object_source_overflows)),
                "typst": int(sum(value > 2 for value in object_typst_overflows)),
            },
            "mean_overflow_beyond_frame_pt": {
                "powerpoint": float(np.mean(object_source_overflows)) if object_source_overflows else None,
                "typst": float(np.mean(object_typst_overflows)) if object_typst_overflows else None,
            },
        },
        "glyph_boxes_outside_slide": {"powerpoint": clipped_source, "typst": clipped_output},
        "slides": slides,
        "_matches": spatial_matches,
    }


def visual_pair(source_pdf, typst_pdf, page_index, output_path, title):
    source = page_image(source_pdf[page_index], 1.2)
    converted = page_image(typst_pdf[page_index], 1.2)
    if converted.size != source.size:
        converted = converted.resize(source.size, Image.Resampling.LANCZOS)
    width, height = source.size
    header = 52
    canvas = Image.new("RGB", (width * 2 + 24, height + header + 24), "#eeeeee")
    draw = ImageDraw.Draw(canvas)
    draw.text((12, 12), f"{title}  |  PowerPoint", fill="#111111")
    draw.text((width + 24, 12), "Typst", fill="#111111")
    canvas.paste(source, (0, header))
    canvas.paste(converted, (width + 24, header))
    canvas.save(output_path, optimize=True)


def analyze_one(app, source: Path, output_dir: Path, installed_fonts, text_layout="anchors",
                comparison_slide=None):
    source = source.resolve()
    deck, _, _, slide_fonts = pptx_stats(source, installed_fonts)
    font_supported = set()
    for index, families in enumerate(slide_fonts):
        missing = {family for family in families
                   if family != "(inherited)" and family.casefold() not in installed_fonts}
        if not missing:
            font_supported.add(index)
    text_source = {index for index, families in enumerate(slide_fonts) if families}
    supported_text = text_source & font_supported
    output_dir.mkdir(parents=True, exist_ok=True)
    key = source.stem[:64]
    with tempfile.TemporaryDirectory(prefix="pptx-typst-placement-") as temp:
        temp = Path(temp)
        project = temp / "typst"
        env = os.environ.copy()
        env["PYTHONPATH"] = str(ROOT / "src") + os.pathsep + env.get("PYTHONPATH", "")
        converted = subprocess.run(
            [sys.executable, "-m", "pptx_typst.pptx_to_typst", str(source), "--output", str(project),
             "--text-layout", text_layout],
            cwd=ROOT, env=env, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=900,
        )
        if converted.returncode:
            raise RuntimeError(converted.stderr[-2000:] or converted.stdout[-2000:])
        manifest = json.loads((project / "asset_manifest.json").read_text(encoding="utf-8"))
        typst_result = subprocess.run(
            ["typst", "compile", "main.typ", str(temp / "typst.pdf")], cwd=project,
            capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=900,
        )
        if typst_result.returncode:
            raise RuntimeError(typst_result.stderr[-2000:] or typst_result.stdout[-2000:])
        export_powerpoint_pdf(app, source, temp / "powerpoint.pdf")
        import pymupdf as fitz
        source_pdf, typst_pdf = fitz.open(temp / "powerpoint.pdf"), fitz.open(temp / "typst.pdf")
        try:
            source_shapes = source_text_shapes(deck)
            output_placements = {
                slide["slide"]: [item for item in slide["placements"] if item.get("type") == "text"]
                for slide in manifest["slides"]
            }
            rendered = collect_render_metrics(source_pdf, typst_pdf, supported_text,
                                              source_shapes, output_placements)
            matches = rendered.pop("_matches")
            eligible = [slide for slide in rendered["slides"]
                        if slide["object_linked_position"]["objects_with_glyphs_in_both_pdfs"]]
            if eligible:
                selected = next((slide for slide in eligible if slide["slide"] == comparison_slide), None)
                if selected is None:
                    selected = max(eligible, key=lambda item: (
                        item["object_linked_position"]["mean_first_glyph_displacement_pt"] or 0
                    ))
                visual_path = output_dir / f"{key}_slide_{selected['slide']:03d}.png"
                visual_pair(source_pdf, typst_pdf, selected["slide"] - 1, visual_path, source.stem)
                rendered["comparison_image"] = visual_path.name
                rendered["comparison_slide"] = selected["slide"]
            record = {
                "source": str(source.relative_to(ROOT)) if source.is_relative_to(ROOT) else source.name,
                "source_text_slides": len(text_source),
                "font_supported_text_slides": len(supported_text),
                "font_unsupported_text_slides": len(text_source - supported_text),
                "structural": geometry_metrics(deck, manifest),
                "rendered": rendered,
            }
            return record
        finally:
            source_pdf.close()
            typst_pdf.close()


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("inputs", nargs="+", type=Path, help="representative PPTX files")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "benchmarks/text_fidelity/diagnostics")
    parser.add_argument("--text-layout", choices=("anchors", "anchors-spacing"), default="anchors-spacing",
                        help="paragraph separation mode to measure (anchors-spacing is the converter default)")
    parser.add_argument("--comparison-slide", type=int,
                        help="prefer this slide for the visual comparison when it has matched text")
    args = parser.parse_args()
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    app = win32com.client.DispatchEx("PowerPoint.Application")
    app.Visible = 1
    app.DisplayAlerts = 0
    app.WindowState = 2
    try:
        installed_fonts = get_typst_font_names()
        results = [analyze_one(app, path, args.output_dir, installed_fonts, args.text_layout,
                               args.comparison_slide)
                   for path in args.inputs]
    finally:
        app.Quit()
        pythoncom.CoUninitialize()
    report = {"sample_decks": results}
    target = args.output_dir / "text_placement_diagnostics.json"
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(target)
    for item in results:
        print(item["source"], "structural max pt", item["structural"]["geometry_abs_delta_pt"],
              "rendered", item["rendered"]["historical_text_only_greedy_match"],
              item["rendered"]["spatially_constrained_match"])


if __name__ == "__main__":
    main()
