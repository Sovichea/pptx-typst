#!/usr/bin/env python3
"""Render slide QA previews and inspect Typst-PDF / exported-PPTX artifacts.

This is a triage tool, not a PPTX parser or a visual-fidelity benchmark.
Requires: PyMuPDF, Pillow. Optionally python-pptx for --pptx.

Examples:
    python inspect_deck.py --pdf deck/preview.pdf --out deck/qa --expect-pages 10
    python inspect_deck.py --pdf deck/preview.pdf --pptx deck/export/deck.pptx \
        --report deck/export/conversion_report.json --out deck/qa
    python inspect_deck.py --pdf deck/preview.pdf \
        --roundtrip-pdf deck/export/powerpoint-render.pdf --out deck/qa
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import json
from pathlib import Path
import re
import sys

try:
    import fitz
    from PIL import Image, ImageChops, ImageDraw, ImageStat
except ImportError as exc:
    raise SystemExit("Install dependencies: python -m pip install pymupdf pillow") from exc


def _render(page, width: int) -> Image.Image:
    scale = width / float(page.rect.width)
    pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
    return Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)


def _contact_sheet(paths: list[Path], destination: Path, columns: int, thumb_width: int) -> None:
    gap, label_height, padding = 16, 24, 16
    thumbs: list[Image.Image] = []
    for path in paths:
        with Image.open(path) as image:
            w = thumb_width
            h = max(1, round(image.height * (w / image.width)))
            thumbs.append(image.convert("RGB").resize((w, h), Image.Resampling.LANCZOS))
    if not thumbs:
        return
    rows = (len(thumbs) + columns - 1) // columns
    cell_height = max(im.height for im in thumbs) + label_height
    width = padding * 2 + columns * thumb_width + (columns - 1) * gap
    height = padding * 2 + rows * cell_height + (rows - 1) * gap
    output = Image.new("RGB", (width, height), (239, 240, 242))
    d = ImageDraw.Draw(output)
    for i, (im, path) in enumerate(zip(thumbs, paths)):
        x = padding + (i % columns) * (thumb_width + gap)
        y = padding + (i // columns) * (cell_height + gap)
        output.paste(im, (x, y + label_height))
        d.text((x + 4, y + 4), path.stem, fill=(30, 32, 38))
        d.rectangle((x, y + label_height, x + im.width - 1, y + label_height + im.height - 1), outline=(162, 165, 170), width=1)
    destination.parent.mkdir(parents=True, exist_ok=True)
    output.save(destination)


def _inspect_pdf(path: Path, out: Path, width: int, columns: int) -> dict:
    slides_dir = out / "slides"
    slides_dir.mkdir(parents=True, exist_ok=True)
    slide_files: list[Path] = []
    entries = []
    with fitz.open(path) as document:
        for index, page in enumerate(document):
            image = _render(page, width)
            file = slides_dir / f"slide-{index + 1:03d}.png"
            image.save(file)
            slide_files.append(file)
            page_rect = page.rect
            overflow = []
            text_blocks = [b for b in page.get_text("blocks") if len(b) >= 7 and b[6] == 0]
            for b in text_blocks:
                bbox = fitz.Rect(b[:4])
                if (
                    bbox.x0 < page_rect.x0 - 2 or bbox.y0 < page_rect.y0 - 2
                    or bbox.x1 > page_rect.x1 + 2 or bbox.y1 > page_rect.y1 + 2
                ):
                    overflow.append({
                        "bounds_pt": [round(v, 2) for v in bbox],
                        "sample": str(b[4]).strip().replace("\n", " ")[:90],
                    })
            entries.append({
                "page": index + 1,
                "width_pt": round(page_rect.width, 2),
                "height_pt": round(page_rect.height, 2),
                "text_block_count": len(text_blocks),
                "image_resource_count": len(page.get_images(full=True)),
                "out_of_page_text_blocks": overflow,
            })
    chunk_size = 30
    for start in range(0, len(slide_files), chunk_size):
        chunk = slide_files[start:start+chunk_size]
        name = "contact-sheet.png" if len(slide_files) <= chunk_size else f"contact-sheet-{start//chunk_size+1:03d}.png"
        _contact_sheet(chunk, out / name, columns, max(200, width // 2))
    sizes = {(entry["width_pt"], entry["height_pt"]) for entry in entries}
    return {"file": str(path), "pages": len(entries), "page_sizes_pt": [list(p) for p in sorted(sizes)], "slides": entries}


def _inspect_pptx(path: Path) -> dict:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise SystemExit("To check --pptx, install python-pptx") from exc
    p = Presentation(path)
    emu = 12700.0
    result = {"file": str(path), "slides": len(p.slides), "size_pt": [round(p.slide_width / emu, 2), round(p.slide_height / emu, 2)], "shapes": []}
    for num, slide in enumerate(p.slides, 1):
        stat = Counter()
        out_of_canvas = []
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                stat["pictures"] += 1
            if shape.has_text_frame:
                stat["text_frames"] += 1
            if shape.has_table:
                stat["tables"] += 1
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                stat["groups"] += 1
            stat["total"] += 1
            try:
                x,y,w,h = shape.left/emu, shape.top/emu, shape.width/emu, shape.height/emu
            except (TypeError, AttributeError):
                continue
            W,H=p.slide_width/emu, p.slide_height/emu
            if min(x,y)<-2 or x+w>W+2 or y+h>H+2:
                out_of_canvas.append({"id":shape.shape_id,"name":shape.name,"bounds_pt":[round(x,2),round(y,2),round(w,2),round(h,2)]})
        result["shapes"].append({"slide":num, "counts":dict(stat), "out_of_canvas_shapes":out_of_canvas})
    return result


def _inspect_conversion_report(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    warnings = data.get("warnings", data.get("unsupported", []))
    if not isinstance(warnings, list):
        warnings=[]
    counts = Counter()
    by_slide: dict[str, int] = {}
    material = []
    for warning in warnings:
        if not isinstance(warning, dict):
            continue
        issue = warning.get("issue", warning.get("kind", "unknown"))
        counts[str(issue)] += 1
        slide = str(warning.get("slide", "unknown"))
        by_slide[slide] = by_slide.get(slide, 0) + 1
        if warning.get("outcome") in ("unsupported", "failed") or issue in (
            "unsupported_image_mask", "unsupported_media_format", "placeholder_layering_risk"
        ):
            material.append(warning)
    return {"file": str(path), "status": data.get("status"), "warnings": len(warnings),
            "issue_counts": dict(counts.most_common()), "warnings_by_slide": by_slide,
            "potentially_material_warnings": material}


def _compare_pdfs(pdf_path: Path, pptx_pdf: Path, out: Path, width: int) -> dict:
    folder = out / "comparison"
    folder.mkdir(parents=True,exist_ok=True)
    measurements = []
    with fitz.open(pdf_path) as src, fitz.open(pptx_pdf) as rendered:
        total = min(len(src),len(rendered))
        for index in range(total):
            left, right = _render(src[index],width),_render(rendered[index],width)
            if right.size != left.size:
                right = right.resize(left.size, Image.Resampling.BICUBIC)
            diff = ImageChops.difference(left,right)
            mean_abs_pixel_difference = sum(ImageStat.Stat(diff).mean) / (3*255)
            paired = Image.new("RGB",(left.width*2+16,left.height),(255,255,255))
            paired.paste(left,(0,0));paired.paste(right,(left.width+16,0))
            paired.save(folder/f"page-{index+1:03d}.png")
            measurements.append({"page":index+1,"normalized_mae":round(mean_abs_pixel_difference,5)})
        return {"file":str(pptx_pdf),"pages":len(rendered),"compared_pages":total,
                "mean_normalized_mae":round(sum(row["normalized_mae"] for row in measurements)/total,5) if total else None,
                "per_page":measurements,
                "note":"Coarse raster comparison only; font and renderer differences affect values. Visual and native-editability inspection remain necessary."}


def _component_measurements(pdf_path: Path, components: list[dict], max_x: float, max_y: float) -> dict:
    """Measure searchable, single-line component labels against declared IR bounds."""
    results = []
    with fitz.open(pdf_path) as document:
        for component in components:
            record = {"id": component.get("id"), "slide": component.get("slide"),
                      "container_bounds_pt": component.get("container_bounds_pt"),
                      "text_frame_bounds_pt": component.get("text_frame_bounds_pt"),
                      "horizontal": component.get("horizontal", "center"),
                      "vertical": component.get("vertical", "center"),
                      "expected_font": component.get("font"),
                      "expected_text_lines": component.get("text_lines", [component.get("text", "")])}
            page_number = component.get("slide")
            if not isinstance(page_number, int) or page_number < 1 or page_number > len(document):
                record.update({"status": "unavailable", "reason": "slide_out_of_range"})
                results.append(record)
                continue
            text_lines = record["expected_text_lines"]
            if not isinstance(text_lines, list) or not text_lines or any(
                not isinstance(line, str) or not line or "\n" in line or "\r" in line
                for line in text_lines
            ):
                record.update({"status": "unavailable", "reason": "empty_or_wrapped_search_text"})
                results.append(record)
                continue
            bounds = component.get("text_frame_bounds_pt")
            padding = component.get("padding_pt", {})
            try:
                if not isinstance(bounds, list) or len(bounds) != 4:
                    raise ValueError
                x, y, width, height = (float(v) for v in bounds)
                frame = fitz.Rect(x, y, x + width, y + height)
                left = float(padding.get("left", 0)); right = float(padding.get("right", 0))
                top = float(padding.get("top", 0)); bottom = float(padding.get("bottom", 0))
                content = fitz.Rect(frame.x0 + left, frame.y0 + top,
                                    frame.x1 - right, frame.y1 - bottom)
                if content.is_empty or not content.is_valid:
                    raise ValueError
            except (TypeError, ValueError, AttributeError):
                record.update({"status": "unavailable", "reason": "invalid_component_bounds"})
                results.append(record)
                continue
            page = document[page_number - 1]
            matches = []
            match_error = None
            for line in text_lines:
                found = [fitz.Rect(rect) for rect in page.search_for(line)
                         if fitz.Rect(rect).intersects(frame)]
                if len(found) != 1:
                    match_error = "text_not_found" if not found else "ambiguous_text_match"
                    break
                matches.extend(found)
            if match_error:
                record.update({"status": "unavailable", "reason": match_error,
                               "expected_line_count": len(text_lines), "matched_line_count": len(matches)})
                results.append(record)
                continue
            glyphs = matches[0]
            for rect in matches[1:]:
                glyphs |= rect
            expected_terms = {term.casefold() for line in text_lines
                              for term in re.findall(r"\w+", line, flags=re.UNICODE)}
            resolved_fonts = set()
            for block in page.get_text("dict").get("blocks", []):
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        span_text = str(span.get("text", ""))
                        span_terms = {term.casefold() for term in re.findall(r"\w+", span_text, flags=re.UNICODE)}
                        if expected_terms.intersection(span_terms) and fitz.Rect(span["bbox"]).intersects(frame):
                            resolved_fonts.add(str(span.get("font", "")))
            expected_font = str(record.get("expected_font") or "").casefold()
            if expected_font:
                record["resolved_pdf_fonts"] = sorted(resolved_fonts)
                record["font_status"] = ("match" if any(expected_font in font.casefold() for font in resolved_fonts)
                                         else "mismatch_or_unavailable")
            horizontal = record["horizontal"]
            vertical = record["vertical"]
            if horizontal == "center":
                x_error = abs((glyphs.x0 + glyphs.x1 - content.x0 - content.x1) / 2)
            elif horizontal == "left":
                x_error = abs(glyphs.x0 - content.x0)
            elif horizontal == "right":
                x_error = abs(content.x1 - glyphs.x1)
            else:
                record.update({"status": "unavailable", "reason": "unsupported_horizontal_alignment"})
                results.append(record)
                continue
            if vertical == "center":
                y_error = abs((glyphs.y0 + glyphs.y1 - content.y0 - content.y1) / 2)
            elif vertical == "top":
                y_error = abs(glyphs.y0 - content.y0)
            elif vertical == "bottom":
                y_error = abs(content.y1 - glyphs.y1)
            else:
                record.update({"status": "unavailable", "reason": "unsupported_vertical_alignment"})
                results.append(record)
                continue
            tolerance = 1.0
            overflow = (glyphs.x0 < frame.x0 - tolerance or glyphs.y0 < frame.y0 - tolerance
                        or glyphs.x1 > frame.x1 + tolerance or glyphs.y1 > frame.y1 + tolerance)
            passed = x_error <= max_x and y_error <= max_y and not overflow
            record.update({"status": "pass" if passed else "misaligned_or_overflow",
                           "glyph_bounds_pt": [round(v, 2) for v in glyphs],
                           "horizontal_alignment_error_pt": round(x_error, 2),
                           "vertical_alignment_error_pt": round(y_error, 2),
                           "overflow": overflow})
            results.append(record)
    return {"file": str(pdf_path), "components": results,
            "matched": sum(row["status"] != "unavailable" for row in results),
            "unavailable": sum(row["status"] == "unavailable" for row in results),
            "passed": sum(row["status"] == "pass" for row in results),
            "tolerances_pt": {"horizontal": max_x, "vertical": max_y}}


def _compare_component_measurements(source: dict, rendered: dict | None, renderer: str) -> dict:
    if rendered is None:
        return {"renderer": None, "components": [
            {"id": row["id"], "classification": "source_only",
             "authoring": row, "conversion": None} for row in source["components"]]}
    converted_by_id = {row.get("id"): row for row in rendered["components"]}
    comparisons = []
    for authored in source["components"]:
        converted = converted_by_id.get(authored.get("id"))
        if authored["status"] == "unavailable":
            classification = "authoring_measurement_unavailable"
        elif authored["status"] != "pass":
            classification = "authoring_error"
        elif converted is None or converted["status"] == "unavailable":
            classification = "conversion_measurement_unavailable"
        elif converted["status"] != "pass":
            font_changed = (authored.get("font_status") == "mismatch_or_unavailable"
                            or converted.get("font_status") == "mismatch_or_unavailable"
                            or authored.get("resolved_pdf_fonts") != converted.get("resolved_pdf_fonts"))
            if font_changed:
                classification = "font_substitution_or_rendering_difference"
            else:
                classification = ("converter_error" if renderer.lower() in ("powerpoint", "microsoft powerpoint")
                                  else "conversion_or_renderer_difference")
        else:
            classification = "pass"
        comparisons.append({"id": authored.get("id"), "classification": classification,
                            "authoring": authored, "conversion": converted})
    return {"renderer": renderer, "components": comparisons,
            "counts": dict(Counter(row["classification"] for row in comparisons)),
            "note": "Each unique, unwrapped paragraph label is searched inside its declared frame. Font names are recorded when PDF extraction exposes them. A non-PowerPoint backend is reported as a conversion-or-renderer difference, not proof of a converter defect."}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pdf", type=Path, required=True, help="Typst-compiled preview PDF")
    parser.add_argument("--out", type=Path, required=True, help="Directory for contact sheet and QA report")
    parser.add_argument("--pptx", type=Path, help="Optional exported PowerPoint for structural checks")
    parser.add_argument("--report", type=Path, help="Optional conversion_report.json")
    parser.add_argument("--roundtrip-pdf", type=Path, help="PPTX rendered by PowerPoint or another viewer to PDF")
    parser.add_argument("--components", type=Path, help="JSON manifest emitted by lower_components.py")
    parser.add_argument("--roundtrip-renderer", default="unspecified", help="Backend used for --roundtrip-pdf, e.g. PowerPoint or LibreOffice")
    parser.add_argument("--max-horizontal-error", type=float, default=2.0, help="Allowed component label alignment error in pt")
    parser.add_argument("--max-vertical-error", type=float, default=3.0, help="Allowed component label alignment error in pt")
    parser.add_argument("--expect-pages", type=int, help="Expected page/slide count")
    parser.add_argument("--width", type=int, default=960, help="Rendered slide pixel width")
    parser.add_argument("--columns", type=int, default=3, help="Contact sheet columns")
    parser.add_argument("--strict", action="store_true", help="Exit nonzero on unexpected slide counts")
    args=parser.parse_args()
    if args.width < 200 or args.columns < 1:
        parser.error("--width must be at least 200 and --columns must be positive")
    for required in [args.pdf,args.pptx,args.report,args.roundtrip_pdf,args.components]:
        if required is not None and not required.exists():
            parser.error(f"File not found: {required}")
    args.out.mkdir(parents=True,exist_ok=True)
    result={"timestamp_utc":datetime.now(timezone.utc).isoformat(),"purpose":"visual QA triage, not a fidelity guarantee"}
    result["preview"]=_inspect_pdf(args.pdf,args.out,args.width,args.columns)
    issues=[]
    if args.expect_pages is not None and result["preview"]["pages"] != args.expect_pages:
        issues.append(f"Expected {args.expect_pages} PDF pages but got {result['preview']['pages']}")
    if args.pptx:
        result["pptx"]=_inspect_pptx(args.pptx)
        if result["pptx"]["slides"] != result["preview"]["pages"]:
            issues.append(f"PPTX has {result['pptx']['slides']} slides, preview PDF has {result['preview']['pages']}")
    if args.report:
        result["conversion_report"]=_inspect_conversion_report(args.report)
    if args.roundtrip_pdf:
        result["roundtrip_comparison"]=_compare_pdfs(args.pdf,args.roundtrip_pdf,args.out,args.width)
        if result["roundtrip_comparison"]["pages"] != result["preview"]["pages"]:
            issues.append("PowerPoint-rendered PDF page count differs from Typst preview")
    if args.components:
        try:
            manifest = json.loads(args.components.read_text(encoding="utf-8"))
            components = manifest.get("components", [])
            if not isinstance(components, list):
                raise ValueError("manifest components must be an array")
        except (json.JSONDecodeError, OSError, ValueError) as exc:
            parser.error(f"Invalid --components manifest: {exc}")
        authored = _component_measurements(args.pdf, components,
                                           args.max_horizontal_error, args.max_vertical_error)
        converted = (_component_measurements(args.roundtrip_pdf, components,
                                              args.max_horizontal_error, args.max_vertical_error)
                     if args.roundtrip_pdf else None)
        result["component_alignment"] = {
            "authoring_pdf": authored,
            "roundtrip_pdf": converted,
            "comparison": _compare_component_measurements(authored, converted, args.roundtrip_renderer),
        }
    result["gate_issues"]=issues
    report=args.out/"qa.json"
    report.write_text(json.dumps(result,indent=2,ensure_ascii=False),encoding="utf-8")
    print(json.dumps({"pdf_pages":result["preview"]["pages"],"pptx_slides":result.get("pptx",{}).get("slides"),
                      "conversion_warnings":result.get("conversion_report",{}).get("warnings"),
                      "gate_issues":issues,"qa_report":str(report)},indent=2))
    return 1 if args.strict and issues else 0

if __name__=="__main__":
    sys.exit(main())
